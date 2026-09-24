"""
ابزار ایمپورت و خروجی Excel برای دانش‌آموزان
"""

import os
import sys

import jdatetime

from utils.logger import get_logger

logger = get_logger(__name__)

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.warning("⚠️ openpyxl نصب نیست. pip install openpyxl")

from dal.academic_year_dal import AcademicYearDAL
from dal.student_academic_profile_dal import StudentAcademicProfileDAL
from dal.student_dal import StudentDAL
from models.student import Student
from models.student_academic_profile import StudentAcademicProfile
from utils.logger import get_logger
from utils.persian_date import PersianDate, normalize_digits, to_db_date
from utils.time_utils import utc_now


class ExcelImporter:
    """ایمپورت و خروجی Excel برای دانش‌آموزان"""
    
    def __init__(self):
        self.student_dal = StudentDAL()
        self.profile_dal = StudentAcademicProfileDAL()
        self.academic_year_dal = AcademicYearDAL()
        self.logger = get_logger(self.__class__.__name__)
    
    def export_students_to_excel(self, students, file_path, academic_year=None,
                                 deleted_view=False):
        """
        خروجی لیست دانش‌آموزان به Excel
        
        Args:
            students: لیست دانش‌آموزان
            file_path: مسیر ذخیره فایل
            academic_year: سال تحصیلی (برای نمایش در هدر)
            deleted_view: اگر True باشد، فایل صریحاً «رکوردهای حذف‌شده»
                          معرفی می‌شود (دور هفدهم — بند ۱۵). چرا لازم است؟
                          چون صفحه می‌تواند در حالت «نمایش حذف‌شده‌ها» باشد
                          و خروجی گرفتن از همان فهرست با عنوان «لیست
                          دانش‌آموزان» یعنی فایلی که خواننده نمی‌فهمد
                          داده‌های حذف‌شده‌اند. پیش‌فرض (False) رفتار
                          قبلی را برای هر فراخوان دیگر حفظ می‌کند.
            
        Returns:
            tuple: (success, message)
        """
        if not OPENPYXL_AVAILABLE:
            return False, "کتابخانه openpyxl نصب نیست. pip install openpyxl"
        
        if not students:
            return False, "هیچ دانش‌آموزی برای خروجی وجود ندارد."
        
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "دانش‌آموزان"
            
            # استایل هدر
            header_font = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
            header_alignment = Alignment(horizontal='center', vertical='center')
            
            # ===== عنوان =====
            export_title = ("لیست دانش‌آموزان حذف‌شده (سطل بازیافت)"
                            if deleted_view else "لیست دانش‌آموزان")
            ws.merge_cells('A1:I1')
            title_cell = ws.cell(row=1, column=1, value=export_title)
            title_cell.font = Font(name='B Nazanin', size=16, bold=True)
            title_cell.alignment = Alignment(horizontal='center')
            
            # ===== اطلاعات مدرسه =====
            row = 3
            ws.cell(row=row, column=1, value="تعداد دانش‌آموزان:").font = Font(name='B Nazanin', size=11, bold=True)
            ws.cell(row=row, column=2, value=str(len(students))).font = Font(name='B Nazanin', size=11)
            
            if academic_year:
                row += 1
                ws.cell(row=row, column=1, value="سال تحصیلی:").font = Font(name='B Nazanin', size=11, bold=True)
                ws.cell(row=row, column=2, value=academic_year.title if hasattr(academic_year, 'title') else str(academic_year)).font = Font(name='B Nazanin', size=11)

            if deleted_view:
                row += 1
                ws.cell(row=row, column=1, value="وضعیت رکوردها:").font = Font(name='B Nazanin', size=11, bold=True)
                note = ws.cell(row=row, column=2,
                               value="حذف‌شده (این ردیف‌ها در فهرست فعال دانش‌آموزان نیستند)")
                note.font = Font(name='B Nazanin', size=11, color='C62828')
            
            # ===== هدر جدول =====
            row += 2
            headers = ["ردیف", "نام", "نام خانوادگی", "کد ملی", "تاریخ تولد", "نام پدر", "نام ولی", "تلفن ولی", "آدرس"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            
            # ===== داده‌ها =====
            for idx, student in enumerate(students, 1):
                row += 1
                ws.cell(row=row, column=1, value=idx)
                ws.cell(row=row, column=2, value=student.first_name or "")
                ws.cell(row=row, column=3, value=student.last_name or "")
                ws.cell(row=row, column=4, value=student.national_code or "")
                ws.cell(row=row, column=5, value=student.birth_date or "")
                ws.cell(row=row, column=6, value=student.father_name or "")
                ws.cell(row=row, column=7, value=student.guardian_name or "")
                ws.cell(row=row, column=8, value=student.guardian_phone or "")
                ws.cell(row=row, column=9, value=student.address or "")
            
            # ===== تنظیم عرض ستون‌ها =====
            column_widths = [8, 18, 18, 15, 15, 15, 18, 15, 35]
            for col, width in enumerate(column_widths, 1):
                ws.column_dimensions[chr(64 + col) if col <= 26 else f"A{chr(64 + col - 26)}"].width = width
            
            # ===== فوتر =====
            # ===== اصلاح =====
            # قبلاً «utc_now().strftime('%Y/%m/%d %H:%M')» نوشته می‌شد
            # یعنی تاریخ «میلادی» با فرمت «شمسی» - کاربر 2026/09/18 را
            # به‌عنوان «1405/09/18» می‌خواند: سه ماه جلوتر! حالا مثل
            # بقیهٔ گزارش‌های پروژه از تقویم جلالی استفاده می‌شود.
            row += 2
            try:
                j_now = jdatetime.datetime.now()
                jalali_stamp = (
                    f"{j_now.year:04d}/{j_now.month:02d}/{j_now.day:02d} "
                    f"{j_now.hour:02d}:{j_now.minute:02d}"
                )
            except Exception:
                jalali_stamp = utc_now().strftime('%Y/%m/%d %H:%M')
            footer_cell = ws.cell(row=row, column=1, value=f"تاریخ خروجی: {jalali_stamp}")
            footer_cell.font = Font(name='B Nazanin', size=10, italic=True)
            footer_cell.alignment = Alignment(horizontal='left')
            
            wb.save(file_path)
            return True, f"✅ {len(students)} دانش‌آموز با موفقیت در {file_path} ذخیره شدند."
            
        except Exception as e:
            self.logger.error(f"خطا در خروجی Excel: {e}")
            return False, f"خطا در خروجی Excel: {e!s}"
    
    def import_students_from_excel(self, file_path, academic_year_id=None):
        """
        ایمپورت دانش‌آموزان از فایل Excel
        
        Args:
            file_path: مسیر فایل Excel
            academic_year_id: شناسه سال تحصیلی (در صورت عدم وجود، سال فعال استفاده می‌شود)
            
        Returns:
            tuple: (success, message, imported_count, errors)
        """
        if not OPENPYXL_AVAILABLE:
            return False, "کتابخانه openpyxl نصب نیست. pip install openpyxl", 0, []
        
        if not os.path.exists(file_path):
            return False, "فایل مورد نظر وجود ندارد.", 0, []
        
        try:
            wb = load_workbook(file_path)
            ws = wb.active
            
            # دریافت سال تحصیلی
            if academic_year_id:
                academic_year = self.academic_year_dal.get_by_id(academic_year_id)
            else:
                academic_year = self.academic_year_dal.get_active()
            
            if not academic_year:
                return False, "سال تحصیلی فعالی وجود ندارد. لطفاً یک سال تحصیلی انتخاب کنید.", 0, []
            
            # پیدا کردن هدرها
            headers = []
            header_row = None
            for row in range(1, min(10, ws.max_row + 1)):
                row_data = [cell.value for cell in ws[row]]
                if any(cell and str(cell).strip() in ["نام", "نام خانوادگی", "کد ملی"] for cell in row_data):
                    headers = row_data
                    header_row = row
                    break
            
            if not headers or header_row is None:
                return False, "هدرهای فایل شناسایی نشدند. لطفاً فایل استاندارد را بررسی کنید.", 0, []
            
            # پیدا کردن ستون‌ها
            column_map = {}
            header_map = {
                'نام': 'first_name',
                'نام کوچک': 'first_name',
                'نام خانوادگی': 'last_name',
                'کد ملی': 'national_code',
                'تاریخ تولد': 'birth_date',
                'نام پدر': 'father_name',
                'نام ولی': 'guardian_name',
                'تلفن ولی': 'guardian_phone',
                'آدرس': 'address',
                'پایه': 'grade',
                'کلاس': 'class_name',
            }
            
            # (بازرسی شانزدهم) نگاشت ستون‌ها: نسخهٔ قبلی با «زیررشته» تطبیق
            # می‌داد و چون «نام» زیررشتهٔ «نام خانوادگی»، «نام پدر» و «نام ولی»
            # است، همهٔ این ستون‌ها روی first_name می‌افتادند و last_name هرگز
            # پیدا نمی‌شد؛ حتی فایل نمونهٔ خودِ برنامه با «ستون‌های ضروری یافت
            # نشدند: last_name» رد می‌شد. حالا: اول تطبیق دقیق، بعد طولانی‌ترین
            # کلید؛ هر فیلد فقط به اولین ستون هم‌خوان نگاشت می‌شود.
            keys_longest_first = sorted(header_map, key=len, reverse=True)
            for col, header in enumerate(headers, 1):
                if not header:
                    continue
                header_str = self._normalize_header(header)
                field = header_map.get(header_str)
                if field is None:
                    for key in keys_longest_first:
                        if key in header_str:
                            field = header_map[key]
                            break
                if field and field not in column_map:
                    column_map[field] = col
            
            # بررسی وجود ستون‌های ضروری
            required_fields = ['first_name', 'last_name']
            missing = [f for f in required_fields if f not in column_map]
            if missing:
                return False, f"ستون‌های ضروری یافت نشدند: {', '.join(missing)}", 0, []
            
            # خواندن داده‌ها
            imported = 0
            errors = []
            students = []
            
            for row in range(header_row + 1, ws.max_row + 1):
                try:
                    # خواندن داده‌ها
                    # (دور هفدهم — بند ۱۴) مقدار «خام» سلول‌ها هم نگه داشته
                    # می‌شود: برای ستون تاریخ تولد باید بدانیم کاربر متن
                    # نوشته یا سلول واقعاً از نوع «تاریخ/عدد» اکسل است.
                    student_data = {}
                    raw_values = {}
                    for field, col in column_map.items():
                        cell_value = ws.cell(row=row, column=col).value
                        raw_values[field] = cell_value
                        if cell_value is not None:
                            student_data[field] = str(cell_value).strip()
                        else:
                            student_data[field] = ""

                    # بررسی داده‌های ضروری
                    if not student_data.get('first_name') or not student_data.get('last_name'):
                        errors.append(f"ردیف {row}: نام یا نام خانوادگی خالی است.")
                        continue

                    # ===== (دور هفدهم — بند ۱۴) یکدست‌سازی مقدارها =====
                    # مسیر فرم/سرویس، کد ملی و تلفن را با ارقام ASCII و تاریخ
                    # را با قالب «yyyy/MM/dd» ذخیره می‌کند. ایمپورت این
                    # یکدست‌سازی را انجام نمی‌داد؛ نتیجهٔ عملی:
                    #   • اعتبارسنجی کد ملی «۱۲۳۴۵۶۷۸۹۰» را قبول می‌کرد
                    #     (چون str.isdigit برای ارقام فارسی هم True است) و
                    #     همان مقدار در دیتابیس می‌ماند؛ بعد تکراری‌یابی و
                    #     جست‌وجو با «1234567890» هرگز آن رکورد را پیدا
                    #     نمی‌کرد — یعنی امکان ثبت دوبارهٔ همان شخص.
                    #   • تاریخ «۱۳۹۵/۰۱/۰۱» یا سلول تاریخِ اکسل
                    #     («2026-09-23 00:00:00») عیناً ذخیره می‌شد و
                    #     مقایسه/مرتب‌سازی رشته‌ای تاریخ‌ها می‌شکست.
                    for field in ('national_code', 'guardian_phone'):
                        if student_data.get(field):
                            student_data[field] = normalize_digits(
                                student_data[field]).strip()

                    if student_data.get('birth_date'):
                        raw_birth = raw_values.get('birth_date')
                        if not isinstance(raw_birth, str):
                            errors.append(
                                f"ردیف {row}: سلول «تاریخ تولد» باید متن باشد "
                                "(قالب yyyy/MM/dd)؛ سلول تاریخ/عددِ اکسل "
                                "قابل تفسیر نیست.")
                            continue
                        normalized_birth = to_db_date(raw_birth)
                        if not normalized_birth or not PersianDate.is_valid_persian_date(
                                normalized_birth):
                            errors.append(
                                f"ردیف {row}: تاریخ تولد «{student_data['birth_date']}» "
                                "معتبر نیست. قالب درست: yyyy/MM/dd (مثلاً 1395/03/05)")
                            continue
                        student_data['birth_date'] = normalized_birth

                    # بررسی کد ملی (اختیاری)
                    if student_data.get('national_code'):
                        # بررسی تکراری نبودن کد ملی
                        existing = self.student_dal.get_by_national_code(student_data['national_code'])
                        if existing:
                            errors.append(f"ردیف {row}: کد ملی '{student_data['national_code']}' تکراری است.")
                            continue
                    
                    # ایجاد دانش‌آموز
                    student = Student()
                    student.first_name = student_data.get('first_name', '')
                    student.last_name = student_data.get('last_name', '')
                    student.national_code = student_data.get('national_code', '')
                    student.birth_date = student_data.get('birth_date', '')
                    student.father_name = student_data.get('father_name', '')
                    student.guardian_name = student_data.get('guardian_name', '')
                    student.guardian_phone = student_data.get('guardian_phone', '')
                    student.address = student_data.get('address', '')
                    student.is_active = 1
                    
                    # اعتبارسنجی
                    errors_list = student.validate()
                    if errors_list:
                        errors.append(f"ردیف {row}: {', '.join(errors_list)}")
                        continue
                    
                    # پایه (پیش از هر نوشتن، تا خطای آن ردیف را رد کند نه اینکه بی‌صدا ۱ شود)
                    grade = 1
                    if 'grade' in column_map and student_data.get('grade'):
                        grade_text = student_data['grade'].strip()
                        try:
                            grade = int(float(grade_text))
                        except (TypeError, ValueError):
                            errors.append(f"ردیف {row}: مقدار پایه «{grade_text}» عدد نیست.")
                            continue

                    # ذخیرهٔ دانش‌آموز + پروندهٔ سالانه در «یک» تراکنش (بازرسی شانزدهم):
                    # قبلاً دانش‌آموز commit می‌شد و اگر ساخت پرونده شکست می‌خورد،
                    # دانش‌آموزِ بدون پرونده می‌ماند و در شمارش «موفق» هم حساب می‌شد.
                    db = self.student_dal.db
                    db.begin_transaction()
                    try:
                        created = self.student_dal.create(student)

                        profile = StudentAcademicProfile()
                        profile.student_id = created.id
                        profile.academic_year_id = academic_year.id
                        profile.grade = grade
                        profile.class_name = student_data.get('class_name', '')
                        profile.status = StudentAcademicProfile.STATUS_ACTIVE
                        self.profile_dal.create(profile)
                        db.commit_transaction()
                    except Exception:
                        db.rollback_transaction()
                        raise

                    students.append(created)
                    imported += 1

                except Exception as e:
                    self.logger.warning(f"ردیف {row} ایمپورت نشد: {e}")
                    errors.append(f"ردیف {row}: {e!s}")
                    continue
            
            # ===== خلاصه =====
            # (دور هفدهم — بند ۱۴) صفر رکوردِ واردشده «موفقیت» نیست:
            # نسخهٔ قبلی «✅ 0 دانش‌آموز با موفقیت ایمپورت شدند» برمی‌گرداند
            # و صفحه هم آن را در کادر «اطلاع» نشان می‌داد — یعنی وقتی همهٔ
            # ردیف‌ها رد شده بودند، کاربر پیام موفقیت می‌دید.
            if imported == 0:
                message = ("هیچ دانش‌آموزی ایمپورت نشد؛ همهٔ ردیف‌ها رد "
                           f"شدند ({len(errors)} خطا).")
                return False, message, imported, errors

            message = f"✅ {imported} دانش‌آموز با موفقیت ایمپورت شدند."
            if errors:
                message += f"\n⚠️ {len(errors)} خطا رخ داده است."
            
            return True, message, imported, errors
            
        except Exception as e:
            self.logger.error(f"خطا در ایمپورت Excel: {e}")
            return False, f"خطا در ایمپورت: {e!s}", 0, []
    
    @staticmethod
    def _normalize_header(header):
        """یکدست‌سازی عنوان ستون: حذف فاصله‌های اضافی/نیم‌فاصله و یکسان‌سازی ی/ک عربی"""
        text = str(header).replace('\u200c', ' ').replace('ي', 'ی').replace('ك', 'ک')
        return ' '.join(text.split())

    def create_sample_excel(self, file_path):
        """
        ایجاد فایل Excel نمونه برای ایمپورت
        
        Args:
            file_path: مسیر ذخیره فایل نمونه
            
        Returns:
            tuple: (success, message)
        """
        if not OPENPYXL_AVAILABLE:
            return False, "کتابخانه openpyxl نصب نیست. pip install openpyxl"
        
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "دانش‌آموزان"
            
            # هدرها
            headers = ["نام", "نام خانوادگی", "کد ملی", "تاریخ تولد", "نام پدر", "نام ولی", "تلفن ولی", "آدرس", "پایه", "کلاس"]
            for col, header in enumerate(headers, 1):
                ws.cell(row=1, column=col, value=header)
            
            # نمونه داده
            sample_data = [
                ["علی", "محمدی", "1234567890", "1390/05/15", "رضا", "رضا محمدی", "09123456789", "تهران", "3", "الف"],
                ["سارا", "احمدی", "0987654321", "1391/08/20", "حسن", "حسن احمدی", "09123456788", "تهران", "2", "ب"],
            ]
            
            for row, data in enumerate(sample_data, 2):
                for col, value in enumerate(data, 1):
                    ws.cell(row=row, column=col, value=value)
            
            # تنظیم عرض ستون‌ها
            for col in range(1, 11):
                ws.column_dimensions[chr(64 + col)].width = 18
            
            wb.save(file_path)
            return True, f"فایل نمونه با موفقیت در {file_path} ایجاد شد."
            
        except Exception as e:
            return False, f"خطا در ایجاد فایل نمونه: {e!s}"