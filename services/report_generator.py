"""
سرویس تولید گزارش‌های دانش‌آموزی - نسخه اصلاح شده با PDF بدون Emoji
"""

import os
import sys
from typing import ClassVar

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from dal.academic_year_dal import AcademicYearDAL
from dal.competency_dal import CompetencyDAL
from dal.family_context_dal import FamilyContextDAL
from dal.followup_dal import FollowUpDAL
from dal.intervention_dal import InterventionDAL
from dal.observation_dal import ObservationDAL
from dal.parent_interview_dal import ParentInterviewDAL
from dal.professional_interpretation_dal import ProfessionalInterpretationDAL
from dal.screening_dal import ScreeningDAL
from dal.screening_result_dal import ScreeningResultDAL
from dal.staff_dal import StaffDAL
from dal.student_academic_profile_dal import StudentAcademicProfileDAL
from dal.student_dal import StudentDAL
from database.connection import DatabaseConnection
from services.case_timeline_service import CaseTimelineService
from utils.behavior_analysis import (
    MEANINGFUL_SHARE_CHANGE,
    PATTERN_MIXED,
    PATTERN_NEEDS_ATTENTION,
    PATTERN_STRENGTH,
    classify_pattern,
    count_behaviors,
    growth_direction,
    observations_volume_note,
    pattern_label,
    shares,
    summarize_by_competency,
)
from utils.logger import get_logger
from utils.persian_pdf import PersianPDF
from utils.time_utils import utc_now

logger = get_logger(__name__)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.warning("⚠️ کتابخانه openpyxl نصب نیست. برای نصب: pip install openpyxl")


class ReportGenerator:
    """تولید کننده گزارش‌های دانش‌آموزی با قابلیت ردیابی و خروجی PDF"""
    
    def __init__(self):
        self.student_dal = StudentDAL()
        self.profile_dal = StudentAcademicProfileDAL()
        self.observation_dal = ObservationDAL()
        self.intervention_dal = InterventionDAL()
        self.followup_dal = FollowUpDAL()
        self.academic_year_dal = AcademicYearDAL()
        self.competency_dal = CompetencyDAL()
        self.staff_dal = StaffDAL()
        # لایه‌های اطلاعاتی جدا از مشاهدهٔ رفتاری (بازرسی یازدهم)
        self.family_dal = FamilyContextDAL()
        self.parent_interview_dal = ParentInterviewDAL()
        self.screening_dal = ScreeningDAL()
        self.screening_result_dal = ScreeningResultDAL()
        self.interpretation_dal = ProfessionalInterpretationDAL()
        self.timeline_service = CaseTimelineService()
        self.db = DatabaseConnection()
        self.logger = get_logger(self.__class__.__name__)
    
    def generate_student_report(self, profile_id):
        """
        تولید گزارش کامل برای یک پرونده سالانه دانش‌آموز
        
        Returns:
            dict: شامل تمام اطلاعات گزارش با قابلیت ردیابی
        """
        profile = self.profile_dal.get_by_id(profile_id)
        if not profile:
            return None
        
        student = self.student_dal.get_by_id(profile.student_id)
        if not student:
            return None
        
        academic_year = self.academic_year_dal.get_by_id(profile.academic_year_id)
        
        # دریافت داده‌ها با استفاده از Timeline Service
        timeline = self.timeline_service.get_timeline(profile_id)
        
        observations = self.observation_dal.get_by_student_profile(profile_id)
        interventions = self.intervention_dal.get_by_student_profile(profile_id)
        
        # دریافت پیگیری‌ها
        all_followups = []
        for inter in interventions:
            all_followups.extend(self.followup_dal.get_by_intervention(inter.id))
        
        # تحلیل شایستگی‌ها
        competency_stats = self.calculate_competency_stats(observations)
        # الگوهای رفتاری: قوت/ضعف بر اساس «نوع رفتار ثبت‌شده» (بازرسی یازدهم)
        # (بازرسی دوازدهم) نام‌ها با یک کوئری دسته‌ای؛ summarize هم دیکشنری
        # می‌پذیرد، پس خروجی دقیقاً یکسان است.
        _comp_titles = self.competency_dal.get_titles_by_ids(
            [o.competency_id for o in observations])
        behavior_patterns = summarize_by_competency(
            observations, _comp_titles, min_count=2)
        strengths = behavior_patterns['strengths']
        weaknesses = behavior_patterns['needs_attention']
        recommendations = self.generate_recommendations(
            competency_stats, observations, interventions)
        
        # روند و آمار نیمسال (ترکیب رفتارها، نه تعداد مشاهدات)
        trend_data = self.calculate_trend(observations)
        semester_stats = self.calculate_semester_stats(observations)
        # اثربخشی مداخلات: مداخله ← پیگیری ← نتیجه
        intervention_effectiveness = self.build_intervention_effectiveness(
            interventions, all_followups)
        # تفکیک لایه‌های اطلاعاتی: مشاهده / غربالگری / تفسیر حرفه‌ای
        information_layers = self.build_information_layers(profile_id, observations)
        # زمینهٔ خانوادگی: اطلاعات زمینه‌ای، نه قضاوت
        family_background = self.build_family_background(profile_id)
        # روایت چندسالهٔ رشد (بازرسی دوازدهم): گزارش معمول هم باید مسیر
        # رشد را نشان بدهد، نه فقط خروجی PDF. خطا در ساخت آن نباید گزارش
        # سالانه را بشکند.
        try:
            growth_narrative = self.generate_growth_narrative(student.id)
        except Exception as e:
            self.logger.warning(f"ساخت روایت رشد چندساله ممکن نشد: {e}")
            growth_narrative = {'has_data': False, 'years': [],
                                'narrative': '', 'synthesis': {}}
        
        # داده‌های ردیابی
        traceability = self._build_traceability(observations, interventions, all_followups)
        
        report = {
            'student': student,
            'profile': profile,
            'academic_year': academic_year,
            'observations_count': len(observations),
            'interventions_count': len(interventions),
            'followups_count': len(all_followups),
            'competency_stats': competency_stats,
            'behavior_patterns': behavior_patterns,
            'strengths': strengths,
            'weaknesses': weaknesses,
            'recommendations': recommendations,
            'trend_data': trend_data,
            'semester_stats': semester_stats,
            'intervention_effectiveness': intervention_effectiveness,
            'information_layers': information_layers,
            'family_background': family_background,
            'growth_narrative': growth_narrative,
            'summary': self.generate_summary(
                observations, interventions, all_followups, competency_stats,
                behavior_patterns=behavior_patterns,
            ),
            'observations': observations,
            'interventions': interventions,
            'followups': all_followups,
            'timeline': timeline,
            'has_data': len(observations) > 0 or len(interventions) > 0,
            'traceability': traceability,
        }
        
        return report
    
    def generate_parent_report(self, profile_id):
        """
        تولید گزارش مختصر و غیرمحرمانه برای والدین

        ===== منبع واحد (بازرسی پانزدهم) =====
        پیش از این، این متد نسخهٔ مستقلی از گزارش والدین می‌ساخت و
        ``services/parent_report_service.py`` نسخهٔ دیگری داشت که هیچ
        صفحه‌ای از آن استفاده نمی‌کرد. اکنون منطق فقط در
        ``ParentReportService.generate_parent_report_data`` است (که خودش
        روی ``generate_student_report`` همین کلاس ساخته می‌شود) و این
        متد صرفاً برای سازگاری با فراخوانی‌های قبلی به آن واگذار می‌کند.

        Returns:
            dict | None: گزارش مناسب برای والدین (همان دادهٔ واحد)
        """
        from services.parent_report_service import ParentReportService
        return ParentReportService(report_generator=self).generate_parent_report_data(
            profile_id)

    # ============================================================
    # متدهای کمکی (بدون تغییر)
    # ============================================================
    
    def calculate_competency_stats(self, observations):
        """محاسبه آمار شایستگی‌ها بر اساس مشاهدات"""
        if not observations:
            return {}

        # (بازرسی دوازدهم) عنوان‌ها با یک کوئری دسته‌ای، نه یکی برای هر
        # مشاهده (رفع N+1؛ خروجی یکسان).
        _titles = self.competency_dal.get_titles_by_ids(
            [o.competency_id for o in observations])

        stats = {}
        for obs in observations:
            if obs.competency_id:
                title = _titles.get(obs.competency_id)
                if title:
                    key = title
                    if key not in stats:
                        stats[key] = {
                            'count': 0,
                            'total_severity': 0,
                            'positive': 0,
                            'negative': 0,
                            'avg_severity': 0,
                            'competency_id': obs.competency_id,
                            'observation_ids': []
                        }
                    
                    stats[key]['count'] += 1
                    stats[key]['total_severity'] += obs.severity or 1
                    stats[key]['observation_ids'].append(obs.id)
                    
                    if obs.behavior_type == "مثبت":
                        stats[key]['positive'] += 1
                    elif obs.behavior_type == "منفی":
                        stats[key]['negative'] += 1
        
        for key in stats:
            if stats[key]['count'] > 0:
                stats[key]['avg_severity'] = round(stats[key]['total_severity'] / stats[key]['count'], 1)
        
        return stats
    
    def analyze_strengths_weaknesses(self, competency_stats, observations=None):
        """
        تحلیل توانمندی‌ها و زمینه‌های نیازمند توجه — **بر پایهٔ نوع رفتار**

        بازرسی یازدهم: پیش از این، قوت/ضعف با «میانگین شدت» تعیین می‌شد.
        این منطق محتوایی نادرست بود؛ چون شدت به‌تنهایی نمی‌گوید یک رفتار
        مثبت است یا منفی. اکنون:

          • نقطهٔ قوت  = الگوی تکرارشوندهٔ رفتارهای مثبت یک شایستگی
          • نیازمند توجه = الگوی تکرارشوندهٔ رفتارهای منفی یک شایستگی
          • رفتار خنثی به‌خودی‌خود نه قوت است نه ضعف
          • یک مشاهدهٔ منفرد مبنای نتیجه‌گیری نیست (حداقل ۲ مشاهده)
          • شدت فقط به‌عنوان «اطلاع اطلاعاتی» در خروجی می‌ماند

        Args:
            competency_stats: آمار شایستگی‌ها (شامل positive/negative)
            observations: در صورت وجود، الگوها از خود مشاهدات ساخته می‌شوند

        Returns:
            tuple: (strengths, weaknesses) — هر آیتم قابل ردیابی به
            شناسهٔ مشاهدات ثبت‌شده است.
        """
        if observations:
            _titles = self.competency_dal.get_titles_by_ids(
                [o.competency_id for o in observations])
            summary = summarize_by_competency(
                observations, _titles, min_count=2)
            return summary['strengths'], summary['needs_attention']

        # مسیر پشتیبان: اگر مشاهدات در دست نبود، از آمار شایستگی‌ها
        strengths, weaknesses = [], []
        for competency, stats in (competency_stats or {}).items():
            positive = stats.get('positive', 0) or 0
            negative = stats.get('negative', 0) or 0
            total = stats.get('count', 0) or 0
            kind = classify_pattern(positive, negative,
                                    max(total - positive - negative, 0), total)
            entry = {
                'competency': competency,
                'pattern': kind,
                'pattern_label': pattern_label(kind),
                'positive': positive,
                'negative': negative,
                'neutral': max(total - positive - negative, 0),
                'count': total,
                'positive_share': shares({'positive': positive, 'negative': negative,
                                          'neutral': max(total - positive - negative, 0),
                                          'total': total})['positive'],
                'negative_share': shares({'positive': positive, 'negative': negative,
                                          'neutral': max(total - positive - negative, 0),
                                          'total': total})['negative'],
                'avg_severity': stats.get('avg_severity', 0),
                'severity_is_auxiliary': True,
                'competency_id': stats.get('competency_id'),
                'observation_ids': stats.get('observation_ids', []),
                'examples': [],
            }
            if kind == PATTERN_STRENGTH:
                strengths.append(entry)
            elif kind == PATTERN_NEEDS_ATTENTION:
                weaknesses.append(entry)

        strengths.sort(key=lambda x: (x['positive'], x['count']), reverse=True)
        weaknesses.sort(key=lambda x: (x['negative'], x['count']), reverse=True)
        return strengths, weaknesses

    def _competency_name(self, competency_id):
        """نام شایستگی از شناسه (برای گزارش‌های الگومحور)."""
        comp = self.competency_dal.get_by_id(competency_id)
        return comp.title if comp else None
    
    def generate_recommendations(self, competency_stats, observations, interventions):
        """
        تولید پیشنهادهای محتاطانه و قابل ردیابی

        بازرسی یازدهم: پیشنهادها از «دادهٔ ثبت‌شده ← الگوی مشاهده‌شده ←
        پیشنهاد برای بررسی/اقدام» ساخته می‌شوند؛ نه «تشخیص قطعی ← دستور».
        هر پیشنهاد به شناسهٔ مشاهدات مرتبط اشاره می‌کند و هیچ‌گاه
        برچسب روان‌شناختی به دانش‌آموز نمی‌زند.
        """
        recommendations = {'teacher': [], 'parents': [], 'counselor': []}
        _titles = self.competency_dal.get_titles_by_ids(
            [o.competency_id for o in (observations or [])])
        patterns = summarize_by_competency(observations or [],
                                           _titles, min_count=2) \
            if observations else {'strengths': [], 'needs_attention': [], 'mixed': []}

        def _ids_text(entry, limit=6):
            ids = [str(i) for i in entry.get('observation_ids', []) if i is not None]
            return ", ".join(ids[:limit]) if ids else "-"

        def _has_intervention(entry):
            obs_ids = set(entry.get('observation_ids', []))
            return any(getattr(inter, 'observation_id', None) in obs_ids
                       for inter in (interventions or []))

        # ۱) الگوهای تکرارشوندهٔ رفتار منفی → بررسی و اقدام
        for entry in patterns['needs_attention']:
            name = entry['competency']
            trace = _ids_text(entry)
            if _has_intervention(entry):
                recommendations['teacher'].append(
                    f"در زمینهٔ «{name}» الگوی تکرارشوندهٔ رفتار منفی مشاهده شده "
                    f"({entry['negative']} مورد از {entry['count']} مشاهدهٔ ثبت‌شده؛ "
                    f"شناسهٔ مشاهدات: {trace}). یک مداخله ثبت شده است؛ "
                    "بررسی اثر آن در پیگیری بعدی پیشنهاد می‌شود."
                )
            else:
                recommendations['teacher'].append(
                    f"در زمینهٔ «{name}» الگوی تکرارشوندهٔ رفتار منفی مشاهده شده "
                    f"({entry['negative']} مورد از {entry['count']} مشاهدهٔ ثبت‌شده؛ "
                    f"شناسهٔ مشاهدات: {trace}). پیشنهاد می‌شود ابتدا با ثبت "
                    "مشاهدهٔ بیشتر، الگو را تأیید و سپس مداخلهٔ هدفمند ثبت کنید."
                )
            recommendations['counselor'].append(
                f"در زمینهٔ «{name}» الگوی تکرارشوندهٔ رفتار منفی ثبت شده است "
                f"(شناسهٔ مشاهدات: {trace}). بررسی و گفت‌وگوی حمایتی پیشنهاد "
                "می‌شود؛ این یک پیشنهاد بررسی است، نه تشخیص."
            )

        # ۲) الگوهای تکرارشوندهٔ رفتار مثبت → تقویت
        for entry in patterns['strengths']:
            name = entry['competency']
            recommendations['parents'].append(
                f"در زمینهٔ «{name}» الگوی تکرارشوندهٔ رفتار مثبت مشاهده شده است "
                f"({entry['positive']} مورد از {entry['count']} مشاهدهٔ ثبت‌شده؛ "
                f"شناسهٔ مشاهدات: {_ids_text(entry)}). تقویت و تشویق این رفتار "
                "پیشنهاد می‌شود."
            )

        # ۳) زمینه‌های ترکیبی/دادهٔ ناکافی → ثبت مشاهدهٔ بیشتر
        for entry in patterns['mixed']:
            recommendations['teacher'].append(
                f"در زمینهٔ «{entry['competency']}» رفتارهای ثبت‌شده ترکیبی است "
                f"({entry['positive']} مثبت و {entry['negative']} منفی از "
                f"{entry['count']} مشاهده). برای جمع‌بندی دقیق‌تر، ثبت مشاهدهٔ "
                "بیشتر پیشنهاد می‌شود."
            )

        # ۴) شایستگی‌های بدون مشاهده
        for competency, stats in (competency_stats or {}).items():
            if (stats.get('count', 0) or 0) == 0:
                recommendations['teacher'].append(
                    f"برای زمینهٔ «{competency}» هیچ مشاهده‌ای ثبت نشده است. "
                    "ثبت مشاهده، مبنای هر تحلیل بعدی است."
                )

        if not recommendations['teacher']:
            recommendations['teacher'].append(
                "برای این پرونده الگوی تکرارشونده‌ای ثبت نشده است یا داده کافی "
                "نیست. ثبت مشاهدهٔ بیشتر به تحلیل دقیق‌تر کمک می‌کند."
            )
        if not recommendations['parents']:
            recommendations['parents'].append(
                "برای این پرونده الگوی تکرارشوندهٔ رفتار مثبت ثبت نشده است. "
                "ثبت مشاهده‌های مثبت، تصویر کامل‌تری از توانمندی‌ها می‌دهد."
            )
        if not recommendations['counselor']:
            recommendations['counselor'].append(
                "بر پایهٔ داده‌های ثبت‌شده، الگوی نیازمند بررسی ویژه‌ای مشاهده "
                "نشده است."
            )

        return recommendations

    def calculate_trend(self, observations):
        """
        محاسبهٔ روند تغییرات بر پایهٔ **ترکیب رفتارها** در طول زمان

        بازرسی یازدهم: تعداد مشاهدات ثبت‌شده در یک ماه، شاخص رشد دانش‌آموز
        نیست (کم شدن ثبت ≠ بهبود؛ زیاد شدن ثبت ≠ بدتر شدن). پس هر بازه با
        سهم رفتارهای مثبت/منفی توصیف می‌شود و «جهت تغییر» از همین سهم‌ها
        به دست می‌آید. تعداد مشاهدات فقط «حجم ثبت و پایش» است.
        """
        if len(observations) < 3:
            return None

        monthly_data = {}
        for obs in observations:
            if obs.observation_date and len(obs.observation_date) >= 7:
                month_key = obs.observation_date[:7]
                if month_key not in monthly_data:
                    monthly_data[month_key] = {
                        'count': 0,
                        'positive': 0,
                        'negative': 0,
                        'neutral': 0,
                        'total_severity': 0,
                        'observation_ids': []
                    }
                bucket = monthly_data[month_key]
                bucket['count'] += 1
                bucket['total_severity'] += obs.severity or 1
                bucket['observation_ids'].append(obs.id)
                if obs.behavior_type == "مثبت":
                    bucket['positive'] += 1
                elif obs.behavior_type == "منفی":
                    bucket['negative'] += 1
                else:
                    bucket['neutral'] += 1

        months = sorted(monthly_data.keys())
        if len(months) < 2:
            return None

        trend_data = []
        for month in months:
            data = monthly_data[month]
            count = data['count']
            share = shares({'positive': data['positive'], 'negative': data['negative'],
                            'neutral': data['neutral'], 'total': count})
            trend_data.append({
                'month': month,
                'label': month,
                'count': count,                       # حجم ثبت و پایش
                'total': count,                       # همان مقدار؛ کلید موردنیاز growth_direction
                'positive': data['positive'],
                'negative': data['negative'],
                'neutral': data['neutral'],
                'positive_share': share['positive'],
                'negative_share': share['negative'],
                # شدت: فقط اطلاعات تکمیلی
                'avg_severity': round(data['total_severity'] / count, 1) if count else 0,
                'severity_is_auxiliary': True,
                'observation_ids': data['observation_ids'],
            })
        return trend_data

    def calculate_trend_direction(self, trend_data):
        """
        جهت تغییر رفتار در طول زمان (فقط از ترکیب رفتارها)

        ===== اصلاح (بازرسی پانزدهم) =====
        ``calculate_trend`` حجم هر ماه را با کلید ``count`` می‌داد ولی
        ``growth_direction`` بازه‌ها را با کلید ``total`` می‌شناسد؛ در
        نتیجه همهٔ بازه‌ها «بدون داده» شمرده می‌شدند و «جهت تغییر» در
        گزارش سالانهٔ برنامه و PDF همیشه «دادهٔ کافی برای تحلیل روند»
        گزارش می‌شد، حتی با چند ماه مشاهده. اکنون هر دو کلید پذیرفته
        می‌شوند (و ``calculate_trend`` هم ``total`` را می‌دهد).
        """
        periods = [
            {'label': p.get('label') or p.get('month'),
             'positive': p.get('positive', 0) or 0,
             'negative': p.get('negative', 0) or 0,
             'neutral': p.get('neutral', 0) or 0,
             'total': p.get('total', p.get('count', 0)) or 0}
            for p in (trend_data or [])
        ]
        return growth_direction(periods)
    
    def calculate_semester_stats(self, observations):
        """محاسبه آمار نیمسال اول و دوم"""
        if not observations:
            return None
        
        first_semester = []
        second_semester = []
        
        for obs in observations:
            if obs.observation_date and len(obs.observation_date) >= 7:
                month = int(obs.observation_date[5:7])
                if 1 <= month <= 6:
                    second_semester.append(obs)
                else:
                    first_semester.append(obs)
        
        stats = {
            'first': {
                'count': len(first_semester),
                'positive': sum(1 for o in first_semester if o.behavior_type == "مثبت"),
                'negative': sum(1 for o in first_semester if o.behavior_type == "منفی"),
                'avg_severity': round(sum(o.severity or 1 for o in first_semester) / len(first_semester), 1) if first_semester else 0,
                'observation_ids': [o.id for o in first_semester]
            },
            'second': {
                'count': len(second_semester),
                'positive': sum(1 for o in second_semester if o.behavior_type == "مثبت"),
                'negative': sum(1 for o in second_semester if o.behavior_type == "منفی"),
                'avg_severity': round(sum(o.severity or 1 for o in second_semester) / len(second_semester), 1) if second_semester else 0,
                'observation_ids': [o.id for o in second_semester]
            }
        }
        
        # بازرسی یازدهم: مقایسهٔ نیمسال‌ها بر پایهٔ «ترکیب رفتارها» است،
        # نه میانگین شدت. شدت تنها به‌عنوان اطلاعات تکمیلی گزارش می‌شود.
        for key in ('first', 'second'):
            share = shares({'positive': stats[key]['positive'],
                            'negative': stats[key]['negative'],
                            'neutral': stats[key]['count'] - stats[key]['positive']
                                       - stats[key]['negative'],
                            'total': stats[key]['count']})
            stats[key]['positive_share'] = share['positive']
            stats[key]['negative_share'] = share['negative']

        direction = growth_direction([
            {'label': 'نیمسال اول', 'positive': stats['first']['positive'],
             'negative': stats['first']['negative'],
             'neutral': max(stats['first']['count'] - stats['first']['positive']
                            - stats['first']['negative'], 0),
             'total': stats['first']['count']},
            {'label': 'نیمسال دوم', 'positive': stats['second']['positive'],
             'negative': stats['second']['negative'],
             'neutral': max(stats['second']['count'] - stats['second']['positive']
                            - stats['second']['negative'], 0),
             'total': stats['second']['count']},
        ])
        stats['direction'] = direction
        stats['trend'] = direction['label']
        stats['trend_note'] = direction['message']
        stats['volume_note'] = direction['volume_note']

        if stats['first']['count'] > 0 and stats['second']['count'] > 0:
            # اختلاف شدت فقط برای اطلاع (تکمیلی)
            stats['severity_change'] = round(
                stats['second']['avg_severity'] - stats['first']['avg_severity'], 1)
        elif stats['first']['count'] > 0 and stats['second']['count'] == 0:
            stats['trend'] = "داده ناکافی در نیمسال دوم"
        elif stats['first']['count'] == 0 and stats['second']['count'] > 0:
            stats['trend'] = "داده ناکافی در نیمسال اول"
        else:
            stats['trend'] = "داده ناکافی"
        
        return stats
    
    def generate_summary(self, observations, interventions, followups,
                         competency_stats, behavior_patterns=None):
        """
        جمع‌بندی سالانهٔ رشد (بازرسی یازدهم)

        بر پایهٔ اطلاعات موجود نشان می‌دهد: توانمندی‌ها، زمینه‌های نیازمند
        توجه، اقدامات انجام‌شده، نتیجهٔ پیگیری‌ها و روند کلی تغییر رفتار
        در همان سال. تصویر دانش‌آموز فقط «مجموعه‌ای از مشکلات» نیست.
        """
        counts = count_behaviors(observations)
        if not observations:
            return ("هنوز مشاهده‌ای برای این پرونده ثبت نشده است. "
                    "ثبت مشاهده، مبنای تحلیل رشد است.")

        if behavior_patterns is None:
            _titles = self.competency_dal.get_titles_by_ids(
                [o.competency_id for o in observations])
            patterns = summarize_by_competency(
                observations, _titles, min_count=2)
        else:
            patterns = behavior_patterns
        share = shares(counts)

        strengths = "؛ ".join(
            f"{e['competency']} ({e['positive']} رفتار مثبت از {e['count']} مشاهده)"
            for e in patterns['strengths'][:3]) or "الگوی تکرارشوندهٔ رفتار مثبت ثبت نشده است"
        needs = "؛ ".join(
            f"{e['competency']} ({e['negative']} رفتار منفی از {e['count']} مشاهده)"
            for e in patterns['needs_attention'][:3]) or "الگوی تکرارشوندهٔ رفتار منفی ثبت نشده است"

        # نتیجهٔ پیگیری‌ها (مداخله ← پیگیری ← نتیجه)
        effectiveness = self.build_intervention_effectiveness(interventions, followups)
        outcomes = effectiveness['outcome_counts']
        outcome_text = (
            f"بهبود مشاهده‌شده: {outcomes.get('improved', 0)}؛ "
            f"بدون تغییر قابل مشاهده: {outcomes.get('no_change', 0)}؛ "
            f"تداوم وضعیت: {outcomes.get('continued', 0)}؛ "
            f"نیازمند پیگیری بیشتر: {outcomes.get('needs_more', 0)}؛ "
            f"بدون پیگیری ثبت‌شده: {outcomes.get('no_followup', 0)}"
        )

        direction = growth_direction([
            {'label': p['month'], 'positive': p['positive'], 'negative': p['negative'],
             'neutral': p['neutral'], 'total': p['count']}
            for p in (self.calculate_trend(observations) or [])
        ])

        summary = f"""
جمع‌بندی سالانهٔ رشد دانش‌آموز:

۱) توانمندی‌ها (الگوی رفتار مثبت): {strengths}
۲) زمینه‌های نیازمند توجه (الگوی رفتار منفی): {needs}
۳) ترکیب رفتارهای ثبت‌شده: {counts['positive']} مثبت، {counts['negative']} منفی و {counts['neutral']} خنثی
   از مجموع {counts['total']} مشاهده (سهم مثبت {share['positive']}٪، سهم منفی {share['negative']}٪)
۴) اقدامات و نتایج: {len(interventions)} مداخله و {len(followups)} پیگیری ثبت شده است.
   نتیجهٔ پیگیری‌ها → {outcome_text}
۵) روند تغییر رفتار در سال: {direction['label']} — {direction['message']}
   مسیر تغییر بین بازه‌ها: {direction.get('path_text') or 'دادهٔ کافی برای ترسیم مسیر وجود ندارد'}

یادداشت‌های محتوایی:
• {observations_volume_note(counts)}
• {direction['volume_note']}
• این جمع‌بندی بر پایهٔ رفتارهای ثبت‌شده است و هیچ تشخیص روان‌شناختی
  یا برچسب قطعی دربارهٔ دانش‌آموز ارائه نمی‌کند.
• مقایسه فقط با خود دانش‌آموز در طول زمان انجام می‌شود، نه با دیگران.
"""
        return summary

    def build_intervention_effectiveness(self, interventions, followups):
        """
        اتصال «مشکل/نیاز مشاهده‌شده ← مداخله ← پیگیری ← نتیجه»

        Returns:
            dict: {'items': [...], 'outcome_counts': {...}, 'summary': str}
        """
        by_intervention = {}
        for follow in followups or []:
            inter_id = getattr(follow, 'intervention_id', None)
            by_intervention.setdefault(inter_id, []).append(follow)

        items = []
        outcome_counts = {'improved': 0, 'no_change': 0, 'continued': 0,
                          'new_status': 0, 'insufficient': 0, 'needs_more': 0,
                          'no_followup': 0}

        for inter in interventions or []:
            related = sorted(by_intervention.get(inter.id, []),
                             key=lambda f: f.date or '')
            followup_items = []
            for follow in related:
                key = getattr(follow, 'result_type', None) or 'insufficient'
                outcome_counts[key if key in outcome_counts else 'insufficient'] += 1
                followup_items.append({
                    'date': follow.date,
                    'method': follow.method or 'پیگیری',
                    'status': follow.status_display,
                    'result_type': key,
                    'result_label': follow.result_type_display,
                    'result_description': follow.result_description or '',
                    'observation_ids': [follow.id] if follow.id else [],
                })
            if not related:
                outcome_counts['no_followup'] += 1
            items.append({
                'intervention_id': inter.id,
                'type': inter.type_display,
                'date': inter.date,
                'goal': inter.goal or '',
                'description': inter.description or '',
                'status': inter.status_display,
                'observation_id': inter.observation_id,
                'followups': followup_items,
                'outcome': (followup_items[-1]['result_label'] if followup_items
                            else 'پیگیری ثبت نشده است'),
                'outcome_known': bool(followup_items),
            })

        known = sum(1 for i in items if i['outcome_known'])
        summary = (f"از {len(items)} مداخلهٔ ثبت‌شده، برای {known} مورد نتیجهٔ "
                   "پیگیری ثبت شده است.")
        return {'items': items, 'outcome_counts': outcome_counts, 'summary': summary}

    def build_information_layers(self, profile_id, observations):
        """
        تفکیک سه لایهٔ اطلاعاتی (بازرسی یازدهم)

          • مشاهدهٔ رفتاری: آنچه کارکنان مدرسه ثبت کرده‌اند.
          • غربالگری: نتیجهٔ یک ابزار/آزمون مشخص (تشخیص نیست).
          • تفسیر حرفه‌ای: ثبت نظر متخصص بر پایهٔ اطلاعات موجود.

        این لایه‌ها با هم مخلوط نمی‌شوند و نتیجهٔ غربالگری هرگز به‌عنوان
        تشخیص یا برچسب قطعی نمایش داده نمی‌شود.
        """
        counts = count_behaviors(observations)
        layers = {
            'observation': {
                'title': 'مشاهدهٔ رفتاری (ثبت کارکنان مدرسه)',
                'counts': counts,
                'note': observations_volume_note(counts),
                'items': [{
                    'date': getattr(o, 'observation_date', None),
                    'behavior': getattr(o, 'behavior', '') or '',
                    'type': getattr(o, 'behavior_type', '') or '',
                    'severity': getattr(o, 'severity', None),
                } for o in (observations or [])[:20]],
            },
            'screening': {'title': 'غربالگری (نتیجهٔ ابزار/آزمون)', 'items': [],
                          'note': ('نتیجهٔ غربالگری، «یافتهٔ ابزار» است و '
                                   'به‌معنای تشخیص روان‌شناختی یا برچسب قطعی '
                                   'برای دانش‌آموز نیست.')},
            'interpretation': {'title': 'تفسیر حرفه‌ای (نظر متخصص)',
                               'items': [],
                               'note': ('تفسیر حرفه‌ای، ثبت نظر متخصص بر پایهٔ '
                                        'اطلاعات موجود است و جای تشخیص بالینی '
                                        'را نمی‌گیرد.')},
            'chain_note': ('مسیر اطلاعات: مشاهدهٔ رفتاری ← غربالگری ← '
                           'تفسیر حرفه‌ای؛ هر لایه معنای مستقل خود را دارد.'),
        }

        try:
            for screening in self.screening_dal.get_by_student_profile(profile_id) or []:
                layers['screening']['items'].append({
                    'tool': getattr(screening, 'tool_display_name', None)
                            or getattr(screening, 'tool_name', '') or 'ابزار غربالگری',
                    'date': getattr(screening, 'execution_date', None),
                    'status': screening.status_display,
                    'domain': getattr(screening, 'domain', '') or '',
                    'total_score': getattr(screening, 'total_score', None),
                    'note': 'یافتهٔ ابزار — نه تشخیص',
                })
        except Exception as e:
            self.logger.warning(f"خواندن لایهٔ غربالگری ممکن نشد: {e}")

        try:
            for interp in self.interpretation_dal.get_by_student_profile(profile_id) or []:
                layers['interpretation']['items'].append({
                    'level': interp.level_display,
                    'title': interp.title or '',
                    'summary': interp.summary or '',
                    'status': interp.status_display,
                    'next_steps': interp.next_steps or '',
                })
        except Exception as e:
            self.logger.warning(f"خواندن لایهٔ تفسیر حرفه‌ای ممکن نشد: {e}")

        return layers

    def build_family_background(self, profile_id):
        """
        زمینهٔ خانوادگی و گفت‌وگوهای والدین (بازرسی یازدهم)

        این اطلاعات **زمینه‌ای** است برای درک بهتر وضعیت دانش‌آموز و
        هرگز مبنای قضاوت دربارهٔ خانواده یا برچسب‌گذاری نیست.
        """
        background = {
            'note': ('اطلاعات خانواده، زمینه‌ای است؛ مبنای قضاوت دربارهٔ خانواده '
                     'یا دانش‌آموز نیست و فقط برای درک بهتر وضعیت رشد او '
                     'ثبت می‌شود.'),
            'contexts': [],
            'interviews': [],
        }
        def _as_list(value):
            """DAL ممکن است یک رکورد یا فهرست برگرداند؛ هر دو حالت را بپذیر."""
            if value is None:
                return []
            if isinstance(value, (list, tuple)):
                return list(value)
            return [value]

        try:
            for ctx in _as_list(self.family_dal.get_by_student_profile(profile_id)):
                background['contexts'].append({
                    'guardian_status': getattr(ctx, 'guardian_status_display', None)
                    or getattr(ctx, 'guardian_status', '') or '-',
                    'parental_support': getattr(ctx, 'parental_support', None) or '-',
                    'economic_status': getattr(ctx, 'economic_status', None) or '-',
                    'family_stress': getattr(ctx, 'family_stress', None) or '-',
                    'study_space': ('دارد' if getattr(ctx, 'has_study_space', 0) else 'ندارد'),
                    'notes': getattr(ctx, 'educational_notes', '') or '',
                })
        except Exception as e:
            self.logger.warning(f"خواندن زمینهٔ خانوادگی ممکن نشد: {e}")

        try:
            for iv in _as_list(self.parent_interview_dal.get_by_student_profile(profile_id)):
                background['interviews'].append({
                    'date': getattr(iv, 'interview_date', None),
                    'method': getattr(iv, 'method_display', None)
                    or getattr(iv, 'interview_method', '') or '-',
                    'topic': getattr(iv, 'topic', '') or '',
                    'status': getattr(iv, 'status_display', None)
                    or getattr(iv, 'status', '') or '-',
                    'summary': getattr(iv, 'summary', '') or '',
                    'result': getattr(iv, 'result', '') or '',
                    'next_action': getattr(iv, 'next_action', '') or '',
                })
        except Exception as e:
            self.logger.warning(f"خواندن مصاحبه‌های والدین ممکن نشد: {e}")

        return background

    def generate_growth_narrative(self, student_id):
        """
        روایت رشد چندسالهٔ دانش‌آموز (بازرسی یازدهم → تکمیل در دوازدهم و سیزدهم)

        برای گزارش پایان دورهٔ ابتدایی: مسیر رشد دانش‌آموز در سال‌های
        مختلف با هم نشان داده می‌شود — توانمندی‌ها، زمینه‌های نیازمند
        توجه، اقدامات، نتایج و جهت تغییر. مقایسه فقط با خودِ دانش‌آموز
        در طول زمان است.

        ===== تکمیل (بازرسی سیزدهم) =====
        خروجی دیگر «فهرست جداگانهٔ سال‌ها» نیست؛ علاوه بر اطلاعات هر سال،
        یک **روایت منسجم مسیر رشد** ساخته می‌شود که در همهٔ خروجی‌ها
        (گزارش معمول برنامه، PDF و Excel) با یک قالب مشترک رندر می‌شود:

          ۱) مسیر کلی تغییر رفتار سال‌به‌سال (با لحاظ سال‌های میانی)
          ۲) الگوهایی که در طول سال‌ها ادامه داشته‌اند (و آیا هنوز
             در آخرین سال هم دیده می‌شوند)
          ۳) زمینه‌هایی که تغییر کرده‌اند — با جهت تغییر (بهبود / افت /
             تازه پدیدآمده / دیگر به‌صورت الگو ثبت نشده) و (بازرسی پانزدهم)
             مسیر سال‌به‌سال هر زمینه به همراه «تغییر پس از مداخله»: اگر در
             سالی برای زمینه‌ای مداخله ثبت شده باشد، وضعیت همان زمینه در
             سال بعد (کمتر شده / ادامه یافته / برطرف شده / بیشتر شده / بدون
             دادهٔ کافی) صریحاً گفته می‌شود
          ۴) کدام نوع مداخله‌ها در پیگیری نتیجهٔ بهتری داشته‌اند (و در
             کدام زمینه‌ها)، کدام‌ها بدون بهبود ثبت‌شده و کدام‌ها بدون
             پیگیری بوده‌اند
          ۵) جمع‌بندی مسیر رشد و اولویت‌های ادامهٔ مسیر

        هیچ مقایسه‌ای با دانش‌آموزان دیگر انجام نمی‌شود و هیچ
        تشخیص/برچسبی تولید نمی‌شود.
        """
        empty = {'has_data': False, 'years': [], 'narrative': '',
                 'synthesis': {}, 'direction': None,
                 'message': 'برای این دانش‌آموز پرونده‌ای ثبت نشده است.'}
        profiles = self.profile_dal.get_all_profiles_for_student(student_id)
        if not profiles:
            return empty

        years = []
        timeline_periods = []
        # الگوهای هر سال (نام شایستگی → نوع الگو) برای تحلیل تداوم/تغییر
        yearly_patterns = []
        for profile in profiles:
            observations = self.observation_dal.get_by_student_profile(profile.id)
            interventions = self.intervention_dal.get_by_student_profile(profile.id)
            followups = []
            for inter in interventions:
                followups.extend(self.followup_dal.get_by_intervention(inter.id))

            counts = count_behaviors(observations)
            _titles = self.competency_dal.get_titles_by_ids(
                [o.competency_id for o in observations])
            patterns = summarize_by_competency(observations, _titles, 2)
            effectiveness = self.build_intervention_effectiveness(interventions, followups)
            year_title = getattr(profile, 'academic_year_title', 'نامشخص')

            # پیوند مداخله ← مشاهدهٔ مبنا ← زمینه (شایستگی)
            observation_area = {
                getattr(o, 'id', None): _titles.get(getattr(o, 'competency_id', None))
                for o in observations
            }
            intervention_items = []
            for item in effectiveness['items']:
                last_result = (item['followups'][-1]['result_type']
                               if item['followups'] else None)
                intervention_items.append({
                    'intervention_id': item['intervention_id'],
                    'type': item['type'],
                    'date': item['date'],
                    'competency': observation_area.get(item.get('observation_id')),
                    'result_type': last_result,
                    'result_label': item['outcome'],
                    'followups_count': len(item['followups']),
                })

            year_shares = shares(counts)
            years.append({
                'year': year_title,
                'grade': profile.grade_display,
                'counts': counts,
                'shares': year_shares,
                'positive_share': year_shares['positive'],
                'negative_share': year_shares['negative'],
                'strengths': [{'competency': e['competency'],
                               'positive': e['positive'], 'count': e['count']}
                              for e in patterns['strengths'][:3]],
                'needs_attention': [{'competency': e['competency'],
                                     'negative': e['negative'], 'count': e['count']}
                                    for e in patterns['needs_attention'][:3]],
                'interventions_count': len(interventions),
                'followups_count': len(followups),
                'outcomes': effectiveness['outcome_counts'],
                'intervention_items': intervention_items,
                'volume_note': observations_volume_note(counts),
                'has_data': counts['total'] > 0,
            })
            yearly_patterns.append({
                'year': year_title,
                'grade': profile.grade_display,
                'has_data': counts['total'] > 0,
                'strengths': [e['competency'] for e in patterns['strengths']],
                'needs': [e['competency'] for e in patterns['needs_attention']],
                'mixed': [e['competency'] for e in patterns['mixed']],
                'observed': [e['competency'] for e in patterns['all']],
                # شمارش هر زمینه در این سال (برای مسیر زمینه و تغییر پس از
                # مداخله — بازرسی پانزدهم)
                'areas': {
                    e['competency']: {
                        'pattern': e['pattern'],
                        'positive': e['positive'],
                        'negative': e['negative'],
                        'neutral': e.get('neutral', 0),
                        'count': e['count'],
                        'positive_share': e.get('positive_share'),
                        'negative_share': e.get('negative_share'),
                    }
                    for e in patterns['all']
                },
            })
            timeline_periods.append({
                'label': year_title,
                'positive': counts['positive'], 'negative': counts['negative'],
                'neutral': counts['neutral'], 'total': counts['total'],
            })

        direction = growth_direction(timeline_periods)
        synthesis = self._build_growth_synthesis(years, yearly_patterns, direction)
        years_with_data = [y for y in years if y['has_data']]
        total_obs = sum(y['counts']['total'] for y in years)
        total_inter = sum(y['interventions_count'] for y in years)
        narrative = (
            f"مسیر رشد دانش‌آموز در {len(years)} سال تحصیلی ثبت‌شده بررسی شد؛ در "
            f"{len(years_with_data)} سال، مشاهدهٔ رفتاری ثبت شده است (در مجموع "
            f"{total_obs} مشاهده و {total_inter} مداخله). "
            f"جهت کلی تغییر رفتار در طول سال‌ها: {direction['label']}. "
            "این روایت بر پایهٔ رفتارهای ثبت‌شده است، فقط دانش‌آموز را با خودش در "
            "طول زمان مقایسه می‌کند و شامل تشخیص روان‌شناختی یا مقایسه با "
            "دانش‌آموزان دیگر نیست. جزئیات در پنج بخش زیر آمده است: مسیر "
            "سال‌به‌سال، الگوهای ادامه‌دار، زمینه‌های تغییریافته (همراه با مسیر هر "
            "زمینه و تغییر پس از مداخله)، نتیجهٔ مداخلات و جمع‌بندی."
        )
        return {'has_data': bool(years_with_data), 'years': years,
                'direction': direction, 'narrative': narrative,
                'synthesis': synthesis}

    # ----- کمکی‌های روایت چندساله ---------------------------------------

    @staticmethod
    def _years_text(years, limit=3):
        years = [str(y) for y in (years or []) if y is not None]
        if not years:
            return ''
        text = ' و '.join(years[:limit])
        if len(years) > limit:
            text += ' و …'
        return text

    @staticmethod
    def _share_text(value):
        if value is None:
            return '-'
        value = float(value)
        return str(int(value)) if value.is_integer() else str(value)

    def _classify_area_changes(self, yearly_patterns):
        """
        تغییر الگوی هر زمینه در طول سال‌ها — فقط مقایسهٔ دانش‌آموز با خودش.

        برای هر زمینه، الگوی سال‌های قبلی با الگوی «آخرین سال دارای داده»
        مقایسه می‌شود:
          • need_to_strength : قبلاً نیازمند توجه، در آخرین سال توانمندی (بهبود)
          • strength_to_need : قبلاً توانمندی، در آخرین سال نیازمند توجه (افت)
          • need_resolved    : قبلاً نیازمند توجه، در آخرین سال دیگر الگوی
                               تکرارشونده‌ای ثبت نشده (با احتیاط تفسیر شود)
          • strength_faded   : قبلاً توانمندی، در آخرین سال دیگر الگو نیست
          • need_emerged     : فقط در آخرین سال به‌عنوان نیاز پدید آمده
          • strength_emerged : فقط در آخرین سال به‌عنوان توانمندی پدید آمده
        """
        data_years = [yp for yp in (yearly_patterns or []) if yp.get('has_data')]
        if len(data_years) < 2:
            return []
        latest = data_years[-1]
        earlier = data_years[:-1]

        strength_years, need_years = {}, {}
        for entry in earlier:
            for name in entry.get('strengths', []):
                strength_years.setdefault(name, []).append(entry['year'])
            for name in entry.get('needs', []):
                need_years.setdefault(name, []).append(entry['year'])

        latest_strengths = set(latest.get('strengths', []))
        latest_needs = set(latest.get('needs', []))
        latest_observed = set(latest.get('observed', []))
        latest_mixed = set(latest.get('mixed', []))

        notes = {
            'need_to_strength': ('این زمینه در سال(های) قبل «نیازمند توجه» بود و در '
                                 'آخرین سال، الگوی تکرارشوندهٔ رفتار مثبت دارد؛ یعنی '
                                 'تغییر به سمت توانمندی ثبت شده است.'),
            'strength_to_need': ('این زمینه در سال(های) قبل «توانمندی» بود و در آخرین '
                                 'سال، الگوی تکرارشوندهٔ رفتار منفی دارد؛ بررسی علت این '
                                 'تغییر پیشنهاد می‌شود (این یک پیشنهاد بررسی است، نه '
                                 'تشخیص).'),
            'need_resolved': ('این زمینه در سال(های) قبل «نیازمند توجه» بود اما در '
                              'آخرین سال، دیگر الگوی تکرارشوندهٔ رفتار منفی ثبت نشده '
                              'است. این می‌تواند نشانهٔ تغییر باشد، ولی چون نبودِ الگو '
                              'ممکن است ناشی از ثبت کمتر باشد، تأیید آن به مشاهدهٔ '
                              'بیشتر نیاز دارد.'),
            'strength_faded': ('این زمینه در سال(های) قبل «توانمندی» بود اما در آخرین '
                               'سال، الگوی تکرارشوندهٔ رفتار مثبت برای آن ثبت نشده است؛ '
                               'ثبت مشاهدهٔ بیشتر روشن می‌کند که تغییر رخ داده یا فقط '
                               'کمتر ثبت شده است.'),
            'need_emerged': ('این زمینه در سال‌های قبل الگوی تکرارشونده‌ای نداشت و در '
                             'آخرین سال به‌عنوان «نیازمند توجه» پدید آمده است؛ پیگیری '
                             'در ادامهٔ مسیر پیشنهاد می‌شود.'),
            'strength_emerged': ('این زمینه در سال‌های قبل الگوی تکرارشونده‌ای نداشت و در '
                                 'آخرین سال به‌عنوان «توانمندی» پدید آمده است.'),
        }
        labels = {
            'need_to_strength': 'بهبود: از نیازمند توجه به توانمندی',
            'strength_to_need': 'افت: از توانمندی به نیازمند توجه',
            'need_resolved': 'نیاز پیشین که دیگر به‌صورت الگو ثبت نشده',
            'strength_faded': 'توانمندی پیشین که دیگر به‌صورت الگو ثبت نشده',
            'need_emerged': 'نیاز تازه‌پدیدآمده در آخرین سال',
            'strength_emerged': 'توانمندی تازه‌پدیدآمده در آخرین سال',
        }

        changes = []

        def _add(name, change, extra=None):
            entry = {
                'competency': name,
                'change': change,
                'change_label': labels[change],
                'strength_years': list(strength_years.get(name, [])),
                'need_years': list(need_years.get(name, [])),
                'latest_year': latest['year'],
                'note': notes[change],
            }
            if extra:
                entry.update(extra)
            changes.append(entry)

        for name in sorted(set(strength_years) | set(need_years)):
            was_strength = name in strength_years
            was_need = name in need_years
            if name in latest_strengths:
                if was_need:
                    _add(name, 'need_to_strength')
                # توانمندی که توانمندی مانده → ماندگار (در بخش دیگر گزارش می‌شود)
            elif name in latest_needs:
                if was_strength:
                    _add(name, 'strength_to_need')
                # نیازی که نیاز مانده → ماندگار
            else:
                observed_note = (
                    'در آخرین سال مشاهده‌ای برای این زمینه ثبت نشده است.'
                    if name not in latest_observed else
                    ('در آخرین سال رفتارهای ثبت‌شدهٔ این زمینه ترکیبی است.'
                     if name in latest_mixed else
                     'در آخرین سال تعداد رفتارهای ثبت‌شدهٔ این زمینه برای الگو کافی نیست.')
                )
                if was_need:
                    _add(name, 'need_resolved', {'latest_status': observed_note})
                elif was_strength:
                    _add(name, 'strength_faded', {'latest_status': observed_note})

        earlier_pattern_names = set(strength_years) | set(need_years)
        for name in sorted(latest_needs - earlier_pattern_names):
            _add(name, 'need_emerged')
        for name in sorted(latest_strengths - earlier_pattern_names):
            _add(name, 'strength_emerged')

        return changes

    def _summarize_interventions_by_type(self, years):
        """
        اثربخشی مداخلات در طول سال‌ها — به تفکیک «نوع مداخله» و زمینه.

        منبع: زنجیرهٔ ثبت‌شدهٔ مداخله ← پیگیری ← نتیجه در همهٔ سال‌ها.
        «مؤثرتر» یعنی در پیگیری‌های ثبت‌شده، «بهبود مشاهده‌شده» بیشتری
        داشته است؛ نه قضاوت دربارهٔ خود روش. مداخلهٔ بدون پیگیری، قابل
        ارزیابی نیست و جداگانه گزارش می‌شود.
        """
        from collections import Counter

        by_type = {}
        area_map = {}
        for year in years or []:
            for item in year.get('intervention_items') or []:
                bucket = by_type.setdefault(item['type'], {
                    'type': item['type'], 'total': 0, 'with_followup': 0,
                    'improved': 0, 'no_change': 0, 'continued': 0,
                    'new_status': 0, 'insufficient': 0, 'needs_more': 0,
                    'no_followup': 0, 'years': [], 'competencies': Counter(),
                })
                bucket['total'] += 1
                if year['year'] not in bucket['years']:
                    bucket['years'].append(year['year'])
                if item.get('competency'):
                    bucket['competencies'][item['competency']] += 1
                result = item.get('result_type')
                if not result:
                    bucket['no_followup'] += 1
                else:
                    bucket['with_followup'] += 1
                    key = result if result in bucket else 'insufficient'
                    bucket[key] += 1
                if item.get('competency'):
                    area = area_map.setdefault(item['competency'], {})
                    stats = area.setdefault(item['type'], {'total': 0, 'improved': 0,
                                                           'with_followup': 0})
                    stats['total'] += 1
                    if result:
                        stats['with_followup'] += 1
                    if result == 'improved':
                        stats['improved'] += 1

        summary = []
        for bucket in by_type.values():
            share = (round(bucket['improved'] / bucket['with_followup'] * 100)
                     if bucket['with_followup'] else None)
            bucket['improved_share'] = share
            bucket['competencies'] = [name for name, _ in bucket['competencies'].most_common(3)]
            if bucket['with_followup'] == 0:
                bucket['verdict'] = 'not_evaluable'
                bucket['verdict_label'] = 'بدون پیگیری ثبت‌شده — قابل ارزیابی نیست'
            elif bucket['improved'] > 0 and share >= 50:
                bucket['verdict'] = 'improved'
                bucket['verdict_label'] = 'در پیگیری، بهبود ثبت شده است'
            elif bucket['improved'] > 0:
                bucket['verdict'] = 'partial'
                bucket['verdict_label'] = 'نتیجهٔ پیگیری ترکیبی (بهبود فقط در بخشی از موارد)'
            else:
                bucket['verdict'] = 'no_improvement'
                bucket['verdict_label'] = 'در پیگیری‌های ثبت‌شده، بهبودی ثبت نشده است'
            summary.append(bucket)

        summary.sort(key=lambda b: (b['improved'], b['improved_share'] or 0,
                                    b['with_followup'], b['total']), reverse=True)
        return summary, area_map

    # وضعیت یک زمینه در یک سال (مسیر زمینه — بازرسی پانزدهم)
    _AREA_STATUS_LABELS: ClassVar[dict] = {
        'need': 'نیازمند توجه',
        'strength': 'توانمندی',
        'mixed': 'ترکیبی',
        'insufficient': 'ثبت کم (زیر آستانهٔ الگو)',
        'not_observed': 'مشاهده‌ای ثبت نشده',
    }
    # وضعیت زمینه در «سال بعد از مداخله»
    _POST_INTERVENTION_LABELS: ClassVar[dict] = {
        'resolved': 'برطرف شده (در سال بعد رفتار منفی ثبت نشده؛ نیازمند تأیید با مشاهدهٔ بیشتر)',
        'reduced': 'کمتر شده',
        'continued': 'ادامه یافته',
        'increased': 'بیشتر شده',
        'no_data': 'در سال بعد مشاهده‌ای برای این زمینه ثبت نشده؛ قابل جمع‌بندی نیست',
        'pending': 'سال بعدی هنوز ثبت نشده؛ نتیجه در ادامهٔ مسیر مشخص می‌شود',
    }

    @staticmethod
    def _post_intervention_status(step, next_step):
        """
        وضعیت زمینه در سال بعد از مداخله — فقط مقایسهٔ دانش‌آموز با خودش.

          • سال بعدی وجود ندارد → pending
          • در سال بعد مشاهده‌ای برای زمینه نیست → no_data (نه «برطرف شده»)
          • در سال بعد رفتار منفی ثبت نشده → resolved (با احتیاط)
          • هنوز الگوی منفی است: سهم منفی دست‌کم ۱۰ واحد کم شده → reduced
            (ولی هنوز الگو)، وگرنه continued
          • دیگر الگوی منفی نیست: اگر پیش از مداخله الگوی منفی بود یا سهم
            منفی معنادار کم شده → reduced؛ اگر سهم منفی معنادار زیاد شده →
            increased؛ وگرنه continued
        """
        if next_step is None:
            return 'pending'
        if next_step['status'] == 'not_observed' or not next_step.get('count'):
            return 'no_data'
        if not next_step.get('negative'):
            return 'resolved'
        before = step.get('negative_share')
        after = next_step.get('negative_share')
        before = float(before) if before is not None and step.get('count') else None
        after = float(after) if after is not None else 0.0
        dropped = before is not None and after <= before - MEANINGFUL_SHARE_CHANGE
        raised = before is not None and after >= before + MEANINGFUL_SHARE_CHANGE
        if next_step['status'] == 'need':
            return 'reduced' if dropped else 'continued'
        if dropped or step['status'] == 'need':
            return 'reduced'
        if raised:
            return 'increased'
        return 'continued'

    def _build_area_paths(self, years, yearly_patterns):
        """
        مسیر هر زمینه در طول سال‌ها + تغییر پس از مداخله (بازرسی پانزدهم)

        داوری مدیر پروژه: روایت چندساله نباید فقط فهرست سال‌ها باشد؛ اگر
        در سال اول مشکلی ثبت شده و برای آن مداخله انجام شده، در سال بعد
        باید مشخص شود که آن مشکل «کمتر شده»، «ادامه یافته» یا «برطرف شده»
        است. این متد برای هر زمینه‌ای که در یکی از سال‌ها «نیازمند توجه»
        بوده یا مداخله‌ای به آن پیوند خورده، مسیر سال‌به‌سال را می‌سازد و
        برای هر سالِ دارای مداخله، وضعیت همان زمینه در سال بعد را با
        ``_post_intervention_status`` تعیین می‌کند. مقایسه فقط دانش‌آموز
        با خودش است؛ نبودِ مشاهده «برطرف شدن» تلقی نمی‌شود.

        Returns:
            list[dict]: {competency, path[], post_intervention[], text}
        """
        data_years = [yp for yp in (yearly_patterns or []) if yp.get('has_data')]
        if not data_years:
            return []
        items_by_year = {y['year']: (y.get('intervention_items') or [])
                         for y in (years or [])}

        candidates = []
        for yp in data_years:
            for name in yp.get('needs', []):
                if name not in candidates:
                    candidates.append(name)
        for year in years or []:
            for item in year.get('intervention_items') or []:
                name = item.get('competency')
                if name and name not in candidates:
                    candidates.append(name)

        kind_map = {PATTERN_NEEDS_ATTENTION: 'need', PATTERN_STRENGTH: 'strength',
                    PATTERN_MIXED: 'mixed'}
        paths = []
        for name in candidates:
            path = []
            for yp in data_years:
                area = (yp.get('areas') or {}).get(name) or {}
                if not area.get('count'):
                    status = 'not_observed'
                else:
                    status = kind_map.get(area.get('pattern'), 'insufficient')
                interventions = [
                    {'type': item.get('type'),
                     'date': item.get('date'),
                     'result_type': item.get('result_type'),
                     'result_label': (item.get('result_label')
                                      if item.get('result_type') else 'بدون پیگیری')}
                    for item in items_by_year.get(yp['year'], [])
                    if item.get('competency') == name
                ]
                path.append({
                    'year': yp['year'],
                    'grade': yp.get('grade'),
                    'status': status,
                    'label': self._AREA_STATUS_LABELS[status],
                    'positive': area.get('positive', 0) or 0,
                    'negative': area.get('negative', 0) or 0,
                    'count': area.get('count', 0) or 0,
                    'negative_share': area.get('negative_share'),
                    'interventions': interventions,
                })

            post = []
            for index, step in enumerate(path):
                if not step['interventions']:
                    continue
                next_step = path[index + 1] if index + 1 < len(path) else None
                status = self._post_intervention_status(step, next_step)
                post.append({
                    'year': step['year'],
                    'next_year': next_step['year'] if next_step else None,
                    'types': [i['type'] for i in step['interventions']],
                    'results': [i['result_label'] for i in step['interventions']],
                    'status': status,
                    'label': self._POST_INTERVENTION_LABELS[status],
                    'negative_before': step['negative'],
                    'negative_after': next_step['negative'] if next_step else None,
                    'count_before': step['count'],
                    'count_after': next_step['count'] if next_step else None,
                })

            paths.append({
                'competency': name,
                'path': path,
                'post_intervention': post,
                'text': self._area_path_text(name, path, post),
            })
        return paths

    def _area_path_text(self, name, path, post):
        """یک خط خوانا برای مسیر زمینه: سال‌به‌سال ← … + تغییر پس از مداخله."""
        steps = []
        for step in path:
            if step['status'] == 'not_observed':
                detail = step['label']
            else:
                detail = (f"{step['label']} ({step['negative']} منفی / "
                          f"{step['positive']} مثبت از {step['count']} مشاهده)")
            if step['interventions']:
                detail += " + مداخله: " + "، ".join(
                    f"{i['type']} ("
                    + (f"پیگیری: {i['result_label']}" if i['result_type'] else 'بدون پیگیری')
                    + ")"
                    for i in step['interventions'])
            steps.append(f"سال {step['year']}: {detail}")
        text = f"مسیر «{name}»: " + " ← ".join(steps) + "."
        if post:
            text += " تغییر پس از مداخله: " + "؛ ".join(
                (f"پس از مداخلهٔ سال {p['year']}"
                 + (f" در سال {p['next_year']}" if p['next_year'] else '')
                 + f": {p['label']}"
                 + (f" ({p['negative_before']} منفی از {p['count_before']} ← "
                    f"{p['negative_after']} منفی از {p['count_after']})"
                    if p['status'] in ('reduced', 'continued', 'increased', 'resolved')
                    else ''))
                for p in post) + "."
        return text

    def _build_growth_synthesis(self, years, yearly_patterns, direction):
        """
        جمع‌بندی منسجم مسیر رشد چندساله (بازرسی دوازدهم → تکمیل در سیزدهم)

        فقط از مقایسهٔ دانش‌آموز با خودش در طول زمان ساخته می‌شود.

        Returns:
            dict با کلیدهای:
              text (متن فشرده)، sections (بخش‌های روایت برای رندر یکسان در
              همهٔ خروجی‌ها)، trajectory، persistent_strengths،
              persistent_needs، changed_areas، effective_interventions،
              area_interventions، priorities، caveats، effective_years،
              totals (برای سازگاری با نسخهٔ قبل).
        """
        from collections import Counter

        data_patterns = [yp for yp in (yearly_patterns or []) if yp.get('has_data')]
        latest = data_patterns[-1] if data_patterns else None
        latest_year = latest['year'] if latest else None

        # ---------- ۱) الگوهای ماندگار ----------
        strength_years, need_years = {}, {}
        for entry in data_patterns:
            for name in entry.get('strengths', []):
                strength_years.setdefault(name, []).append(entry['year'])
            for name in entry.get('needs', []):
                need_years.setdefault(name, []).append(entry['year'])

        def _persistent(mapping, latest_names):
            items = [
                {'competency': name, 'years': yrs, 'count': len(yrs),
                 'still_present': name in latest_names}
                for name, yrs in mapping.items() if len(yrs) >= 2
            ]
            items.sort(key=lambda x: (x['still_present'], x['count'], x['competency']),
                       reverse=True)
            return items

        latest_strength_names = set(latest.get('strengths', [])) if latest else set()
        latest_need_names = set(latest.get('needs', [])) if latest else set()
        persistent_strengths = _persistent(strength_years, latest_strength_names)
        persistent_needs = _persistent(need_years, latest_need_names)

        # ---------- ۲) زمینه‌های تغییریافته (با جهت تغییر) ----------
        changed_areas = self._classify_area_changes(yearly_patterns)
        # مسیر هر زمینه + تغییر پس از مداخله (بازرسی پانزدهم)
        area_paths = self._build_area_paths(years, yearly_patterns)
        post_counts = Counter()
        for area in area_paths:
            for item in area['post_intervention']:
                post_counts[item['status']] += 1

        # ---------- ۳) مداخلات به تفکیک نوع و زمینه ----------
        effective_interventions, area_interventions = \
            self._summarize_interventions_by_type(years)

        # سازگاری با نسخهٔ قبل: سال‌هایی که «بهبود مشاهده‌شده» بیشتری داشتند
        totals = Counter()
        improved_by_year = {}
        for year in years or []:
            outcomes = year.get('outcomes', {}) or {}
            for key, value in outcomes.items():
                totals[key] += value or 0
            improved = outcomes.get('improved', 0) or 0
            if improved > 0:
                improved_by_year[year.get('year')] = improved
        effective_years = [
            {'year': year, 'improved': count}
            for year, count in sorted(improved_by_year.items(),
                                      key=lambda kv: kv[1], reverse=True)
        ]

        # ---------- ۴) مسیر سال‌به‌سال ----------
        step_by_to = {s['to']: s for s in (direction.get('steps') or [])}
        patterns_by_year = {yp['year']: yp for yp in data_patterns}
        per_year = []
        for year in years or []:
            if not year.get('has_data'):
                continue
            step = step_by_to.get(year['year'])
            year_patterns = patterns_by_year.get(year['year'], {})
            per_year.append({
                'year': year['year'],
                'grade': year.get('grade'),
                'positive_share': year.get('positive_share'),
                'negative_share': year.get('negative_share'),
                'total': year['counts']['total'],
                'strengths_count': len(year_patterns.get('strengths', [])),
                'needs_count': len(year_patterns.get('needs', [])),
                'step_label': step['label'] if step else None,
                'step_status': step['status'] if step else None,
            })
        trajectory = {
            'status': direction['status'],
            'label': direction['label'],
            'message': direction['message'],
            'path_text': direction.get('path_text', ''),
            'turning_points': direction.get('turning_points', []),
            'per_year': per_year,
            'overall_status': direction.get('overall_status'),
            'volume_note': direction.get('volume_note', ''),
        }

        # ---------- ۵) اولویت‌های ادامهٔ مسیر ----------
        priorities = []
        for item in persistent_needs:
            if item['still_present']:
                tried = area_interventions.get(item['competency'], {})
                helpful = [f"{t} ({s['improved']} از {s['with_followup']})"
                           for t, s in tried.items() if s['improved'] > 0]
                priorities.append({
                    'kind': 'persistent_need',
                    'competency': item['competency'],
                    'text': (f"«{item['competency']}» در {item['count']} سال (از جمله آخرین "
                             "سال) نیازمند توجه بوده است"
                             + (f"؛ مداخلاتی که پیگیری آن‌ها در این زمینه بهبود نشان داده: "
                                f"{'، '.join(helpful)}" if helpful else
                                "؛ برای این زمینه هنوز مداخله‌ای با نتیجهٔ «بهبود» در "
                                "پیگیری ثبت نشده است")
                             + "."),
                })
        for change in changed_areas:
            if change['change'] == 'need_emerged':
                priorities.append({
                    'kind': 'new_need', 'competency': change['competency'],
                    'text': (f"«{change['competency']}» در آخرین سال به‌عنوان زمینهٔ "
                             "نیازمند توجه پدید آمده و برای ادامهٔ مسیر باید پایش شود."),
                })
            elif change['change'] == 'strength_to_need':
                priorities.append({
                    'kind': 'reversed', 'competency': change['competency'],
                    'text': (f"«{change['competency']}» از توانمندی به نیازمند توجه تغییر "
                             "کرده است؛ بررسی زمینهٔ این تغییر پیشنهاد می‌شود."),
                })
        build_on = [p['competency'] for p in persistent_strengths if p['still_present']]
        build_on += [c['competency'] for c in changed_areas
                     if c['change'] in ('need_to_strength', 'strength_emerged')]
        if build_on:
            priorities.append({
                'kind': 'build_on', 'competency': None,
                'text': ("توانمندی‌هایی که در آخرین سال هم دیده می‌شوند و می‌توان بر "
                         f"آن‌ها تکیه کرد: {'، '.join(dict.fromkeys(build_on))}."),
            })

        # ---------- ۶) بخش‌های روایت ----------
        sections = []

        # ۱) مسیر کلی
        lines = []
        if per_year:
            for entry in per_year:
                line = (f"سال {entry['year']} (پایهٔ {entry['grade']}): "
                        f"{self._share_text(entry['positive_share'])}٪ مثبت / "
                        f"{self._share_text(entry['negative_share'])}٪ منفی از "
                        f"{entry['total']} مشاهده؛ {entry['strengths_count']} توانمندی و "
                        f"{entry['needs_count']} زمینهٔ نیازمند توجه")
                if entry['step_label']:
                    line += f" — نسبت به سال قبل: {entry['step_label']}"
                lines.append(line)
        if direction['status'] == 'insufficient':
            lines.append("فقط یک سال با مشاهدهٔ ثبت‌شده وجود دارد؛ مقایسهٔ سال‌به‌سال "
                         "هنوز ممکن نیست و مسیر با ثبت سال‌های بعد کامل می‌شود.")
        else:
            lines.append(f"جهت کلی در طول سال‌ها: {direction['label']} — "
                         f"{direction['message']}")
            if direction.get('turning_points'):
                lines.append("نقطهٔ برگشت جهت: سال "
                             f"{self._years_text(direction['turning_points'])}")
        lines.append(direction.get('volume_note', ''))
        sections.append({'key': 'trajectory',
                         'title': '۱) مسیر کلی تغییر رفتار در طول سال‌ها',
                         'lines': [x for x in lines if x]})

        # ۲) الگوهای ماندگار
        lines = []
        for item in persistent_strengths:
            lines.append(
                f"توانمندی ماندگار: {item['competency']} — در سال‌های "
                f"{self._years_text(item['years'])}"
                + (" (در آخرین سال هم دیده می‌شود)" if item['still_present']
                   else " (در آخرین سال به‌صورت الگو ثبت نشده)"))
        for item in persistent_needs:
            lines.append(
                f"نیازمند توجه ماندگار: {item['competency']} — در سال‌های "
                f"{self._years_text(item['years'])}"
                + (" (هنوز در آخرین سال هم ادامه دارد)" if item['still_present']
                   else " (در آخرین سال به‌صورت الگو ثبت نشده)"))
        if not lines:
            lines.append("الگوی تکرارشوندهٔ مشترکی بین سال‌ها دیده نشد؛ هر سال الگوی "
                         "خودش را دارد و نتیجه‌گیری دربارهٔ تداوم، نیازمند مشاهدهٔ "
                         "بیشتر است.")
        sections.append({'key': 'persistent',
                         'title': '۲) الگوهایی که در طول سال‌ها ادامه داشته‌اند',
                         'lines': lines})

        # ۳) زمینه‌های تغییریافته
        lines = []
        order = ['need_to_strength', 'strength_emerged', 'need_resolved',
                 'strength_to_need', 'need_emerged', 'strength_faded']
        for change_kind in order:
            for change in changed_areas:
                if change['change'] != change_kind:
                    continue
                detail = ''
                if change['change'] in ('need_to_strength', 'need_resolved'):
                    detail = f" (نیازمند توجه در سال {self._years_text(change['need_years'])})"
                elif change['change'] in ('strength_to_need', 'strength_faded'):
                    detail = f" (توانمندی در سال {self._years_text(change['strength_years'])})"
                line = f"{change['change_label']}: {change['competency']}{detail}"
                if change.get('latest_status'):
                    line += f" — {change['latest_status']}"
                lines.append(line)
        if not lines:
            if len(data_patterns) < 2:
                lines.append("برای شناسایی تغییر الگو، دست‌کم دو سال با مشاهدهٔ ثبت‌شده "
                             "لازم است.")
            else:
                lines.append("زمینه‌ای که الگوی آن بین سال‌ها عوض شده باشد، ثبت نشده است.")
        # مسیر هر زمینه و تغییر پس از مداخله (بازرسی پانزدهم): برای هر
        # زمینهٔ نیازمند توجه یا دارای مداخله، سال‌به‌سال گفته می‌شود که
        # پس از مداخله «کمتر شده / ادامه یافته / برطرف شده» است.
        for area in area_paths[:8]:
            lines.append(area['text'])
        if len(area_paths) > 8:
            lines.append(f"و {len(area_paths) - 8} زمینهٔ دیگر (در جدول سال‌ها).")
        sections.append({'key': 'changes',
                         'title': '۳) زمینه‌هایی که در طول سال‌ها تغییر کرده‌اند '
                                  '(با مسیر هر زمینه و تغییر پس از مداخله)',
                         'lines': lines})

        # ۴) مداخلات
        lines = []
        improved_types = [b for b in effective_interventions if b['verdict'] == 'improved']
        partial_types = [b for b in effective_interventions if b['verdict'] == 'partial']
        flat_types = [b for b in effective_interventions if b['verdict'] == 'no_improvement']
        unknown_types = [b for b in effective_interventions if b['verdict'] == 'not_evaluable']
        for bucket in improved_types:
            areas = f" در زمینهٔ {'، '.join(bucket['competencies'])}" if bucket['competencies'] else ''
            lines.append(
                f"مؤثرتر بر اساس پیگیری: {bucket['type']}{areas} — "
                f"{bucket['improved']} «بهبود مشاهده‌شده» از {bucket['with_followup']} "
                f"مداخلهٔ پیگیری‌شده (از مجموع {bucket['total']} مورد در سال‌های "
                f"{self._years_text(bucket['years'])})")
        for bucket in partial_types:
            areas = f" در زمینهٔ {'، '.join(bucket['competencies'])}" if bucket['competencies'] else ''
            lines.append(
                f"نتیجهٔ ترکیبی: {bucket['type']}{areas} — فقط {bucket['improved']} «بهبود "
                f"مشاهده‌شده» از {bucket['with_followup']} مداخلهٔ پیگیری‌شده (بدون تغییر: "
                f"{bucket['no_change']}، تداوم وضعیت: {bucket['continued']}، نیازمند پیگیری "
                f"بیشتر: {bucket['needs_more']})")
        for bucket in flat_types:
            lines.append(
                f"بدون بهبود ثبت‌شده: {bucket['type']} — {bucket['with_followup']} مداخلهٔ "
                "پیگیری‌شده، بدون «بهبود مشاهده‌شده» (بدون تغییر: "
                f"{bucket['no_change']}، تداوم وضعیت: {bucket['continued']}، نیازمند "
                f"پیگیری بیشتر: {bucket['needs_more']})")
        if unknown_types:
            lines.append(
                "قابل ارزیابی نیست (پیگیری ثبت نشده): "
                + "، ".join(f"{b['type']} ({b['total']} مورد)" for b in unknown_types))
        if effective_interventions:
            lines.append(
                f"در مجموع {totals['improved']} «بهبود مشاهده‌شده» در پیگیری‌ها ثبت شده است"
                + (f"؛ بیشترین آن در سال {effective_years[0]['year']}." if effective_years else ".")
                + " «مؤثرتر» یعنی نتیجهٔ پیگیریِ ثبت‌شده بهتر بوده است، نه قضاوت دربارهٔ "
                  "خود روش؛ مداخلات بدون پیگیری قابل ارزیابی نیستند.")
        else:
            lines.append("در هیچ‌یک از سال‌ها مداخله‌ای ثبت نشده است؛ ارزیابی اثربخشی "
                         "ممکن نیست.")
        sections.append({'key': 'interventions',
                         'title': '۴) مداخلات و نتیجهٔ پیگیری آن‌ها در طول سال‌ها',
                         'lines': lines})

        # ۵) جمع‌بندی و اولویت‌ها
        lines = [self._overall_course_text(direction, per_year, persistent_strengths,
                                           persistent_needs, changed_areas)]
        post_summary = self._post_intervention_summary_text(post_counts)
        if post_summary:
            lines.append(post_summary)
        for item in priorities:
            lines.append(item['text'])
        if latest_year and not priorities:
            lines.append("اولویت خاصی از الگوهای چندساله برنمی‌آید؛ ادامهٔ ثبت مشاهده و "
                         "پیگیری مداخلات پیشنهاد می‌شود.")
        sections.append({'key': 'conclusion',
                         'title': '۵) جمع‌بندی مسیر رشد و اولویت‌های ادامهٔ مسیر',
                         'lines': lines})

        caveats = [
            "این جمع‌بندی فقط دانش‌آموز را با خودش در طول زمان مقایسه می‌کند؛ هیچ "
            "مقایسه‌ای با دانش‌آموزان دیگر انجام نمی‌شود.",
            "همهٔ نتیجه‌ها از رفتارهای ثبت‌شده و شناسهٔ آن‌ها قابل ردیابی است و هیچ "
            "تشخیص روان‌شناختی یا برچسب قطعی ارائه نمی‌شود.",
            "تعداد مشاهدات هر سال فقط «حجم ثبت و پایش» است؛ کم یا زیاد شدن آن به‌تنهایی "
            "بهبود یا افت نیست و نبودِ الگو در یک سال می‌تواند ناشی از ثبت کمتر باشد.",
            "نتایج غربالگری و تفسیر حرفه‌ای لایه‌های جداگانه‌اند و در این جمع‌بندی "
            "رفتاری با مشاهده‌ها مخلوط نشده‌اند.",
        ]
        sections.append({'key': 'caveats', 'title': 'یادداشت‌های محتوایی',
                         'lines': caveats})

        # ---------- متن فشرده (سازگار با نسخهٔ قبل) ----------
        parts = [f"مسیر کلی رشد: {direction['label']}."]
        if persistent_strengths:
            parts.append("توانمندی‌های ماندگار (تکرار در چند سال): "
                         + "، ".join(p['competency'] for p in persistent_strengths[:5]) + ".")
        if persistent_needs:
            parts.append(
                "زمینه‌های نیازمند توجه ماندگار: "
                + "؛ ".join(f"{p['competency']} (در سال‌های {self._years_text(p['years'])})"
                            for p in persistent_needs[:5])
                + ". این موارد اولویت پیگیری در ادامهٔ مسیرند؛ نه برچسبی دربارهٔ "
                  "دانش‌آموز.")
        if changed_areas:
            parts.append("زمینه‌های تغییریافته در طول سال‌ها: "
                         + "؛ ".join(f"{c['competency']} ({c['change_label']})"
                                     for c in changed_areas[:5]) + ".")
        if not persistent_strengths and not persistent_needs and not changed_areas:
            parts.append("الگوی تکرارشوندهٔ مشترکی بین سال‌ها دیده نشد؛ هر سال الگوی "
                         "خودش را دارد و نتیجه‌گیری دربارهٔ تداوم، نیازمند مشاهدهٔ "
                         "بیشتر است.")
        if improved_types or partial_types:
            parts.append("مداخلاتی که پیگیری آن‌ها بهبود نشان داده: "
                         + "، ".join(f"{b['type']} ({b['improved']} از {b['with_followup']})"
                                     for b in (improved_types + partial_types)[:3])
                         + "؛ اقدامات با نتیجهٔ بهبود، در ادامهٔ مسیر قابل تکرارند.")
        elif totals['improved'] == 0 and any(
                (y.get('interventions_count') or 0) > 0 for y in years or []):
            parts.append("برای مداخلات ثبت‌شده، «بهبود مشاهده‌شده»‌ای در پیگیری‌ها ثبت "
                         "نشده است؛ ثبت دقیق نتیجهٔ پیگیری، ارزیابی اثربخشی مداخلات را "
                         "ممکن می‌کند.")

        return {
            'text': " ".join(parts),
            'sections': sections,
            'trajectory': trajectory,
            'persistent_strengths': persistent_strengths,
            'persistent_needs': persistent_needs,
            'changed_areas': changed_areas,
            'area_paths': area_paths,
            'post_intervention_counts': dict(post_counts),
            'post_intervention_summary': post_summary,
            'effective_interventions': effective_interventions,
            'area_interventions': area_interventions,
            'priorities': priorities,
            'caveats': caveats,
            'effective_years': effective_years,
            'totals': dict(totals),
            'latest_year': latest_year,
        }

    @staticmethod
    def _post_intervention_summary_text(post_counts):
        """یک جملهٔ جمع‌بندی دربارهٔ تغییر پس از مداخله (بر پایهٔ سال بعد)."""
        if not post_counts:
            return ''
        order = [('resolved', 'برطرف‌شده (با احتیاط)'), ('reduced', 'کمتر شده'),
                 ('continued', 'ادامه‌یافته'), ('increased', 'بیشتر شده'),
                 ('no_data', 'بدون مشاهده در سال بعد'),
                 ('pending', 'سال بعدی هنوز ثبت نشده')]
        parts = [f"{label}: {post_counts[key]} مورد"
                 for key, label in order if post_counts.get(key)]
        return ("تغییر پس از مداخله (وضعیت هر زمینه در سال بعد از مداخله): "
                + "؛ ".join(parts)
                + ". این جمع‌بندی از رفتارهای ثبت‌شدهٔ همان زمینه ساخته شده و "
                  "نبودِ مشاهده، «برطرف شدن» تلقی نشده است.")

    def _overall_course_text(self, direction, per_year, persistent_strengths,
                             persistent_needs, changed_areas):
        """یک جملهٔ جمع‌بندی دربارهٔ مسیر کلی؛ فقط مقایسهٔ دانش‌آموز با خودش."""
        if not per_year:
            return "هنوز مشاهده‌ای در هیچ سالی ثبت نشده است؛ مسیر رشد قابل جمع‌بندی نیست."
        first, last = per_year[0], per_year[-1]
        span = (f"از سال {first['year']} ({self._share_text(first['positive_share'])}٪ مثبت) "
                f"تا سال {last['year']} ({self._share_text(last['positive_share'])}٪ مثبت)")
        improved = [c for c in changed_areas if c['change'] == 'need_to_strength']
        resolved = [c for c in changed_areas if c['change'] == 'need_resolved']
        worsened = [c for c in changed_areas if c['change'] in ('strength_to_need', 'need_emerged')]
        continuing = [p for p in persistent_needs if p['still_present']]
        status = direction['status']
        if status == 'insufficient':
            core = ("تا کنون فقط یک سال با مشاهدهٔ ثبت‌شده وجود دارد؛ مسیر کلی رشد پس از "
                    "ثبت سال‌های بعد قابل جمع‌بندی است")
        elif status == 'improving':
            core = f"مسیر کلی رشد {span} به سمت رفتارهای مثبت‌تر بوده و در سال‌های میانی افت معناداری ثبت نشده است"
        elif status == 'declining':
            core = f"مسیر کلی رشد {span} با افزایش سهم رفتارهای منفی همراه بوده و در سال‌های میانی بهبود معناداری ثبت نشده است"
        elif status == 'stable':
            core = f"ترکیب رفتارها {span} در مجموع تغییر معناداری نداشته است"
        else:
            core = (f"مسیر رشد {span} یکنواخت نبوده است؛ در بعضی سال‌ها بهبود و در بعضی "
                    "دیگر افت ثبت شده و روند کلی «غیرقطعی» است")
        details = []
        if improved:
            details.append("زمینه‌های بهبودیافته (از نیاز به توانمندی): "
                           + "، ".join(c['competency'] for c in improved[:4]))
        if resolved:
            details.append("نیازهای پیشین که در آخرین سال دیگر به‌صورت الگو ثبت نشده‌اند "
                           "(نیازمند تأیید با مشاهدهٔ بیشتر): "
                           + "، ".join(c['competency'] for c in resolved[:4]))
        if worsened:
            details.append("زمینه‌های نیازمند بررسی تازه: "
                           + "، ".join(c['competency'] for c in worsened[:4]))
        if continuing:
            details.append("نیازهای ادامه‌دار: "
                           + "، ".join(p['competency'] for p in continuing[:4]))
        text = core + "."
        if details:
            text += " " + "؛ ".join(details) + "."
        return text

    def growth_narrative_lines(self, narrative, include_years=True):
        """
        رندر مشترک روایت چندساله برای همهٔ خروجی‌ها (بازرسی سیزدهم)

        گزارش معمول برنامه، PDF و Excel همگی از همین فهرست استفاده می‌کنند
        تا روایت در همه‌جا به یک شکل «کامل و یکپارچه» باشد.

        Returns:
            list[tuple[str, str]]: (kind, text) که kind یکی از
            'text' | 'heading' | 'bullet' است.
        """
        lines = []
        if not narrative or not narrative.get('has_data'):
            return lines
        if narrative.get('narrative'):
            lines.append(('text', narrative['narrative']))
        synthesis = narrative.get('synthesis') or {}
        for section in synthesis.get('sections') or []:
            lines.append(('heading', section['title']))
            for line in section.get('lines') or []:
                lines.append(('bullet', line))
        if include_years:
            lines.append(('heading', 'اطلاعات هر سال (مبنای جمع‌بندی)'))
            for year in narrative.get('years') or []:
                if not year.get('has_data'):
                    lines.append(('bullet', f"سال {year['year']} (پایهٔ {year['grade']}): "
                                            "مشاهده‌ای ثبت نشده است."))
                    continue
                strength_text = "؛ ".join(
                    f"{s['competency']} ({s['positive']} رفتار مثبت از {s['count']})"
                    for s in year.get('strengths', [])) or "ثبت نشده"
                need_text = "؛ ".join(
                    f"{n['competency']} ({n['negative']} رفتار منفی از {n['count']})"
                    for n in year.get('needs_attention', [])) or "ثبت نشده"
                outcomes = year.get('outcomes') or {}
                counts = year.get('counts') or {}
                lines.append(('bullet', (
                    f"سال {year['year']} (پایهٔ {year['grade']}): توانمندی‌ها: {strength_text} | "
                    f"نیازمند توجه: {need_text} | ترکیب رفتارها: {counts.get('positive', 0)} مثبت، "
                    f"{counts.get('negative', 0)} منفی، {counts.get('neutral', 0)} خنثی از "
                    f"{counts.get('total', 0)} مشاهده | مداخلات: {year.get('interventions_count', 0)} "
                    f"| نتایج پیگیری: بهبود {outcomes.get('improved', 0)}، بدون تغییر "
                    f"{outcomes.get('no_change', 0)}، تداوم {outcomes.get('continued', 0)}، "
                    f"نیازمند پیگیری بیشتر {outcomes.get('needs_more', 0)}، بدون پیگیری "
                    f"{outcomes.get('no_followup', 0)}")))
        return lines

    def _build_traceability(self, observations, interventions, followups):
        """ساخت داده‌های ردیابی"""
        traceability = {
            'observations': {obs.id: obs for obs in observations},
            'interventions': {inter.id: inter for inter in interventions},
            'followups': {follow.id: follow for follow in followups},
            'observation_ids': [obs.id for obs in observations],
            'intervention_ids': [inter.id for inter in interventions],
            'followup_ids': [follow.id for follow in followups],
        }
        return traceability
    
     
    # ============================================================
    # خروجی PDF (اصلاح شده - بدون Emoji)
    # ============================================================
    
    def export_to_pdf(self, profile_id, file_path):
        """
        خروجی گزارش به صورت فایل PDF - بدون Emoji
        
        Args:
            profile_id: شناسه پرونده دانش‌آموز
            file_path: مسیر ذخیره فایل PDF
            
        Returns:
            tuple: (success, message)
        """
        try:
            report = self.generate_student_report(profile_id)
            if not report:
                return False, "امکان تولید گزارش وجود ندارد."
            
            # ایجاد PDF
            pdf = PersianPDF(file_path)
            
            # ===== عنوان =====
            pdf.add_title("گزارش پرونده سالانه دانش آموز")
            pdf.add_spacer(0.2)
            
            # ===== اطلاعات دانش‌آموز =====
            student = report['student']
            profile = report['profile']
            academic_year = report['academic_year']
            
            info_items = [
                f"نام دانش آموز: {student.full_name}",
                f"پایه: {profile.grade_display}",
                f"کلاس: {profile.class_name or '-'}",
                f"سال تحصیلی: {academic_year.title if academic_year else '-'}",
                f"شناسه پرونده: {profile.id}",
            ]
            
            for item in info_items:
                pdf.add_text(item)
            pdf.add_spacer(0.3)
            
            # ===== خلاصه آماری =====
            pdf.add_subtitle("خلاصه آماری")
            stats_items = [
                f"* تعداد مشاهدات: {report['observations_count']}",
                f"* تعداد مداخلات: {report['interventions_count']}",
                f"* تعداد پیگیری‌ها: {report['followups_count']}",
            ]
            for item in stats_items:
                pdf.add_text(item)
            pdf.add_spacer(0.3)
            
            # ===== توانمندی‌ها (الگوی رفتار مثبت) =====
            pdf.add_subtitle("توانمندی‌ها (الگوی تکرارشوندهٔ رفتار مثبت)")
            if report['strengths']:
                for strength in report['strengths']:
                    line = (f"{strength['competency']} — "
                            f"{strength['positive']} رفتار مثبت از "
                            f"{strength['count']} مشاهدهٔ ثبت‌شده")
                    if strength.get('examples'):
                        line += f" (نمونه: {strength['examples'][0]})"
                    pdf.add_strength(line)
            else:
                pdf.add_text("* الگوی تکرارشوندهٔ رفتار مثبت ثبت نشده است.")
            pdf.add_spacer(0.3)

            # ===== زمینه‌های نیازمند توجه =====
            pdf.add_subtitle("زمینه‌های نیازمند توجه (الگوی تکرارشوندهٔ رفتار منفی)")
            if report['weaknesses']:
                for weakness in report['weaknesses']:
                    line = (f"{weakness['competency']} — "
                            f"{weakness['negative']} رفتار منفی از "
                            f"{weakness['count']} مشاهدهٔ ثبت‌شده")
                    if weakness.get('examples'):
                        line += f" (نمونه: {weakness['examples'][0]})"
                    pdf.add_weakness(line)
            else:
                pdf.add_text("* الگوی تکرارشوندهٔ رفتار منفی ثبت نشده است.")
            pdf.add_text("توجه: این فهرست بر پایهٔ رفتارهای ثبت‌شده است و "
                         "به‌معنای تشخیص یا برچسب نیست.")
            pdf.add_spacer(0.3)
            
            # ===== پیشنهادات =====
            pdf.add_subtitle("پیشنهادات")
            
            pdf.add_bold("به معلم:")
            for rec in report['recommendations']['teacher']:
                pdf.add_text(f"* {rec}")
            pdf.add_spacer(0.2)
            
            pdf.add_bold("به والدین:")
            for rec in report['recommendations']['parents']:
                pdf.add_text(f"* {rec}")
            pdf.add_spacer(0.2)
            
            pdf.add_bold("به مشاور:")
            for rec in report['recommendations']['counselor']:
                pdf.add_text(f"* {rec}")
            pdf.add_spacer(0.3)
            
            # ===== روند تغییر رفتار =====
            if report['trend_data']:
                pdf.add_subtitle("روند تغییر رفتار (بر پایهٔ ترکیب رفتارها)")
                for item in report['trend_data']:
                    pdf.add_text(
                        f"* {item['month']}: {item['positive']} مثبت، "
                        f"{item['negative']} منفی، {item['neutral']} خنثی "
                        f"(سهم مثبت {item['positive_share']}٪) — "
                        f"حجم ثبت: {item['count']} مشاهده"
                    )
                direction = self.calculate_trend_direction(report['trend_data'])
                pdf.add_text(f"* جهت تغییر: {direction['label']} — {direction['message']}")
                # بازرسی سیزدهم: مسیر بین بازه‌ها (نه فقط ابتدا/انتها)
                if direction.get('path_text'):
                    pdf.add_text(f"* مسیر تغییر بین بازه‌ها: {direction['path_text']}")
                for caution in direction.get('caution_notes') or []:
                    pdf.add_text(f"* توجه: {caution}")
                pdf.add_text(f"* {direction['volume_note']}")
            pdf.add_spacer(0.3)

            # ===== مقایسه نیمسال‌ها =====
            if report['semester_stats']:
                stats = report['semester_stats']
                pdf.add_subtitle("مقایسه نیمسال‌ها")
                pdf.add_text(
                    f"* نیمسال اول: {stats['first']['positive']} مثبت، "
                    f"{stats['first']['negative']} منفی از {stats['first']['count']} "
                    f"مشاهده (سهم مثبت {stats['first'].get('positive_share', 0)}٪)")
                pdf.add_text(
                    f"* نیمسال دوم: {stats['second']['positive']} مثبت، "
                    f"{stats['second']['negative']} منفی از {stats['second']['count']} "
                    f"مشاهده (سهم مثبت {stats['second'].get('positive_share', 0)}٪)")
                pdf.add_text(f"* جهت تغییر رفتار: {stats.get('trend', '-')}")
                if stats.get('trend_note'):
                    pdf.add_text(f"* {stats['trend_note']}")
                if stats.get('volume_note'):
                    pdf.add_text(f"* {stats['volume_note']}")
            pdf.add_spacer(0.3)

            # ===== اثربخشی مداخلات (مداخله ← پیگیری ← نتیجه) =====
            effectiveness = report.get('intervention_effectiveness') or {}
            if effectiveness.get('items'):
                pdf.add_subtitle("اثربخشی مداخلات (اقدام ← پیگیری ← نتیجه)")
                for item in effectiveness['items'][:10]:
                    pdf.add_bold(f"{item['type']} — {item['date'] or ''}")
                    if item.get('goal'):
                        pdf.add_text(f"* هدف: {item['goal']}")
                    if item['followups']:
                        for follow in item['followups'][:3]:
                            pdf.add_text(
                                f"* پیگیری {follow['date'] or ''} "
                                f"({follow['method']}): {follow['result_label']}"
                                + (f" — {follow['result_description']}"
                                   if follow['result_description'] else "")
                            )
                    else:
                        pdf.add_text("* پیگیری ثبت نشده است؛ نتیجهٔ مداخله "
                                     "قابل ارزیابی نیست.")
                pdf.add_text(f"* {effectiveness['summary']}")
                pdf.add_spacer(0.3)

            # ===== زمینهٔ خانوادگی (اطلاعات زمینه‌ای) =====
            family = report.get('family_background') or {}
            if family.get('contexts') or family.get('interviews'):
                pdf.add_subtitle("زمینهٔ خانوادگی و گفت‌وگو با والدین")
                pdf.add_text(f"* {family.get('note', '')}")
                for ctx in family['contexts'][:3]:
                    pdf.add_text(
                        f"* وضعیت سرپرست: {ctx['guardian_status']} | "
                        f"حمایت والدین: {ctx['parental_support']} | "
                        f"فضای مطالعه: {ctx['study_space']}"
                    )
                for iv in family['interviews'][:5]:
                    line = (f"* گفت‌وگو {iv['date'] or ''} ({iv['method']}) — "
                            f"{iv['topic'] or 'بدون موضوع ثبت‌شده'} | "
                            f"وضعیت: {iv['status']}")
                    pdf.add_text(line)
                    if iv.get('result'):
                        pdf.add_text(f"  نتیجهٔ گفت‌وگو: {iv['result']}")
                pdf.add_spacer(0.3)

            # ===== تفکیک لایه‌های اطلاعاتی =====
            layers = report.get('information_layers') or {}
            if layers:
                pdf.add_subtitle("لایه‌های اطلاعاتی (مشاهده / غربالگری / تفسیر)")
                pdf.add_text(f"* {layers.get('chain_note', '')}")
                obs_layer = layers.get('observation', {})
                pdf.add_bold(f"{obs_layer.get('title', 'مشاهدهٔ رفتاری')}")
                pdf.add_text(f"* {obs_layer.get('note', '')}")
                scr = layers.get('screening', {})
                pdf.add_bold(scr.get('title', 'غربالگری'))
                if scr.get('items'):
                    for item in scr['items'][:5]:
                        pdf.add_text(
                            f"* {item['tool']} — {item['date'] or ''} | "
                            f"وضعیت: {item['status']} | یافته: "
                            f"{item['total_score'] if item['total_score'] is not None else '-'}"
                        )
                else:
                    pdf.add_text("* نتیجهٔ غربالگری ثبت نشده است.")
                pdf.add_text(f"* {scr.get('note', '')}")
                interp = layers.get('interpretation', {})
                pdf.add_bold(interp.get('title', 'تفسیر حرفه‌ای'))
                if interp.get('items'):
                    for item in interp['items'][:5]:
                        pdf.add_text(f"* [{item['level']}] {item['title']} — "
                                     f"{item['status']}")
                        if item.get('summary'):
                            pdf.add_text(f"  {item['summary']}")
                else:
                    pdf.add_text("* تفسیر حرفه‌ای ثبت نشده است.")
                pdf.add_text(f"* {interp.get('note', '')}")
                pdf.add_spacer(0.3)

            # ===== سابقهٔ رشد چندساله =====
            # (بازرسی سیزدهم) همان روایت منسجمی که در گزارش معمول برنامه و
            # Excel هست، با رندر مشترک؛ نه فهرست جداگانهٔ سال‌ها.
            narrative = report.get('growth_narrative')
            if not narrative or not narrative.get('has_data'):
                try:
                    narrative = self.generate_growth_narrative(student.id)
                except Exception as exc:
                    narrative = None
                    self.logger.warning(f"ساخت روایت رشد ممکن نشد: {exc}")
            if narrative and narrative.get('has_data'):
                pdf.add_subtitle("سابقهٔ رشد و مسیر طی‌شده (چندساله)")
                for kind, line in self.growth_narrative_lines(narrative):
                    if kind == 'heading':
                        pdf.add_bold(line)
                    elif kind == 'bullet':
                        pdf.add_text(f"* {line}")
                    else:
                        pdf.add_text(line)
                pdf.add_spacer(0.3)
            
            # ===== جدول مشاهدات =====
            if report['observations']:
                pdf.add_subtitle("لیست مشاهدات")
                table_data = [["ردیف", "تاریخ", "محیط", "نوع", "شدت"]]
                for idx, obs in enumerate(report['observations'][:20], 1):
                    severity_str = ""
                    if obs.severity:
                        severity_str = "|" * obs.severity
                    table_data.append([
                        str(idx),
                        obs.observation_date or "",
                        obs.location or "",
                        obs.behavior_type or "خنثی",
                        severity_str
                    ])
                pdf.add_table(table_data)
            
            # ===== خلاصه نهایی =====
            pdf.add_subtitle("خلاصه")
            pdf.add_text(report['summary'])
            pdf.add_spacer(0.3)
            
            # ===== متادیتا =====
            pdf.add_separator()
            import jdatetime
            try:
                today = jdatetime.date.today()
                date_str = f"{today.year:04d}/{today.month:02d}/{today.day:02d}"
            except Exception as _exc:
                self.logger.debug(f"خطای مدیریت‌شده در export_to_pdf (مسیر جایگزین): {_exc}")
                date_str = utc_now().strftime("%Y/%m/%d")
            
            pdf.add_text(f"تاریخ تهیه گزارش: {date_str}")
            pdf.add_text("PARTO - سامانه مدیریت پرونده دانش آموزان")
            pdf.add_text("پشتیبانی: support@partow.ir")
            
            # ساخت PDF
            pdf.build(file_path)
            
            return True, f"فایل PDF با موفقیت در {file_path} ذخیره شد."
            
        except Exception as e:
            self.logger.debug(f"خطای مدیریت‌شده در export_to_pdf (مسیر جایگزین): {e}")
            return False, "خطا در ساخت فایل PDF."
    
    # ============================================================
    # خروجی Excel (موجود)
    # ============================================================
    
    def export_to_excel(self, profile_id, file_path):
        """
        خروجی گزارش به صورت فایل Excel با قابلیت ردیابی
        """
        if not OPENPYXL_AVAILABLE:
            return False, "کتابخانه openpyxl نصب نیست. لطفاً با دستور pip install openpyxl نصب کنید."
        
        report = self.generate_student_report(profile_id)
        if not report:
            return False, "امکان تولید گزارش وجود ندارد."
        
        try:
            wb = Workbook()
            student = report['student']
            profile = report['profile']
            academic_year = report['academic_year']
            
            # ===== برگه اول: اطلاعات اصلی (مشاهدات) =====
            ws1 = wb.active
            ws1.title = "گزارش اصلی"
            
            header_font = Font(name='B Nazanin', size=12, bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
            header_alignment = Alignment(horizontal='center', vertical='center')
            
            headers = ['ردیف', 'شناسه', 'نام', 'نام خانوادگی', 'پایه', 'کلاس', 'تاریخ', 'محیط', 'شایستگی', 'نوع', 'شدت', 'شرح']
            for col, header in enumerate(headers, 1):
                cell = ws1.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            
            observations = report.get('observations', [])
            # (بازرسی دوازدهم) یک کوئری دسته‌ای به‌جای یکی برای هر مشاهده
            _titles = self.competency_dal.get_titles_by_ids(
                [o.competency_id for o in observations])
            for row, obs in enumerate(observations, 2):
                ws1.cell(row=row, column=1, value=row-1)
                ws1.cell(row=row, column=2, value=obs.id)
                ws1.cell(row=row, column=3, value=student.first_name)
                ws1.cell(row=row, column=4, value=student.last_name)
                ws1.cell(row=row, column=5, value=profile.grade_display)
                ws1.cell(row=row, column=6, value=profile.class_name or '')
                ws1.cell(row=row, column=7, value=obs.observation_date or '')
                ws1.cell(row=row, column=8, value=obs.location or '')
                
                competency_name = ""
                if obs.competency_id:
                    competency_name = _titles.get(obs.competency_id, "")
                ws1.cell(row=row, column=9, value=competency_name)
                
                ws1.cell(row=row, column=10, value=obs.behavior_type or '')
                ws1.cell(row=row, column=11, value=obs.severity)
                ws1.cell(row=row, column=12, value=obs.description[:100] if obs.description else '')
            
            for col in range(1, 13):
                ws1.column_dimensions[get_column_letter(col)].width = 15
            
            # ===== برگه دوم: خلاصه با قابلیت ردیابی =====
            ws2 = wb.create_sheet("خلاصه")
            
            ws2.cell(row=1, column=1, value="گزارش رشد دانش‌آموز").font = Font(name='B Nazanin', size=16, bold=True)
            ws2.merge_cells('A1:C1')
            
            info_data = [
                ("شناسه پرونده", profile.id),
                ("نام دانش‌آموز", f"{student.first_name} {student.last_name}"),
                ("پایه", profile.grade_display),
                ("کلاس", profile.class_name or ''),
                ("سال تحصیلی", academic_year.title if academic_year else ''),
                ("تعداد مشاهدات", report['observations_count']),
                ("تعداد مداخلات", report['interventions_count']),
                ("تعداد پیگیری‌ها", report['followups_count']),
            ]
            
            for row, (label, value) in enumerate(info_data, 3):
                ws2.cell(row=row, column=1, value=label).font = Font(name='B Nazanin', size=12, bold=True)
                ws2.cell(row=row, column=2, value=value).font = Font(name='B Nazanin', size=12)
                ws2.merge_cells(f'A{row}:B{row}')
            
            # نقاط قوت با شناسه
            row = 13
            ws2.cell(row=row, column=1, value="⭐ نقاط قوت").font = Font(name='B Nazanin', size=12, bold=True, color='27AE60')
            row += 1
            if report['strengths']:
                for strength in report['strengths']:
                    obs_ids = ', '.join([str(i) for i in strength.get('observation_ids', [])])
                    ws2.cell(row=row, column=1, value=(
                        f"• {strength['competency']} — {strength['positive']} رفتار مثبت "
                        f"از {strength['count']} مشاهده"))
                    ws2.cell(row=row, column=2, value=f"شناسه مشاهده‌ها: {obs_ids}").font = Font(name='B Nazanin', size=9, color='7F8C8D')
                    row += 1
            else:
                ws2.cell(row=row, column=1, value="• موردی یافت نشد").font = Font(name='B Nazanin', size=11)
                row += 1
            
            # نیازهای رشدی با شناسه
            row += 1
            ws2.cell(row=row, column=1, value="🔴 زمینه‌های نیازمند حمایت").font = Font(name='B Nazanin', size=12, bold=True, color='E74C3C')
            row += 1
            if report['weaknesses']:
                for weakness in report['weaknesses']:
                    obs_ids = ', '.join([str(i) for i in weakness.get('observation_ids', [])])
                    ws2.cell(row=row, column=1, value=(
                        f"• {weakness['competency']} — {weakness['negative']} رفتار منفی "
                        f"از {weakness['count']} مشاهده"))
                    ws2.cell(row=row, column=2, value=f"شناسه مشاهده‌ها: {obs_ids}").font = Font(name='B Nazanin', size=9, color='7F8C8D')
                    row += 1
            else:
                ws2.cell(row=row, column=1, value="• موردی یافت نشد").font = Font(name='B Nazanin', size=11)
                row += 1
            
            # روند تغییرات با شناسه
            row += 1
            if report['trend_data']:
                ws2.cell(row=row, column=1, value="📈 روند تغییرات").font = Font(name='B Nazanin', size=12, bold=True, color='3498DB')
                row += 1
                for item in report['trend_data']:
                    obs_ids = ', '.join([str(i) for i in item.get('observation_ids', [])])
                    ws2.cell(row=row, column=1, value=(
                        f"• {item['month']}: {item['positive']} مثبت، {item['negative']} منفی، "
                        f"{item['neutral']} خنثی (سهم مثبت {item['positive_share']}٪) — "
                        f"حجم ثبت: {item['count']}"))
                    ws2.cell(row=row, column=2, value=f"شناسه مشاهده‌ها: {obs_ids}").font = Font(name='B Nazanin', size=9, color='7F8C8D')
                    row += 1
                row += 1
            
            # اثربخشی مداخلات (مداخله ← پیگیری ← نتیجه)
            effectiveness = report.get('intervention_effectiveness') or {}
            if effectiveness.get('items'):
                row += 1
                ws2.cell(row=row, column=1,
                         value="🔗 اثربخشی مداخلات (اقدام ← پیگیری ← نتیجه)"
                         ).font = Font(name='B Nazanin', size=12, bold=True, color='8E44AD')
                row += 1
                for item in effectiveness['items']:
                    outcome = item['outcome']
                    ws2.cell(row=row, column=1, value=(
                        f"• {item['type']} ({item['date'] or '-'}): {outcome}"))
                    if item['followups']:
                        ws2.cell(row=row, column=2, value=(
                            f"پیگیری‌ها: {len(item['followups'])} | نتیجهٔ آخر: "
                            f"{item['followups'][-1]['result_label']}"),
                            ).font = Font(name='B Nazanin', size=9, color='7F8C8D')
                    else:
                        ws2.cell(row=row, column=2,
                                 value="پیگیری ثبت نشده است؛ نتیجه قابل ارزیابی نیست."
                                 ).font = Font(name='B Nazanin', size=9, color='95A5A6')
                    row += 1
                ws2.cell(row=row, column=1, value=f"• {effectiveness['summary']}")
                row += 1

            # پیام "داده ناکافی"
            if report['observations_count'] < 3:
                ws2.cell(row=row, column=1, value="⚠️ داده کافی برای تحلیل روند وجود ندارد.").font = Font(name='B Nazanin', size=11, color='F39C12')
                ws2.cell(row=row, column=2, value=f"تعداد مشاهدات ثبت‌شده: {report['observations_count']} (حداقل ۳ مورد نیاز است)").font = Font(name='B Nazanin', size=11, color='F39C12')
                row += 1
            
            ws2.column_dimensions['A'].width = 60
            ws2.column_dimensions['B'].width = 50
            
            # ===== برگه سوم: شایستگی‌ها =====
            ws3 = wb.create_sheet("شایستگی‌ها")
            
            ws3.cell(row=1, column=1, value="شایستگی").font = Font(name='B Nazanin', size=12, bold=True)
            ws3.cell(row=1, column=2, value="تعداد").font = Font(name='B Nazanin', size=12, bold=True)
            ws3.cell(row=1, column=3, value="میانگین شدت").font = Font(name='B Nazanin', size=12, bold=True)
            ws3.cell(row=1, column=4, value="مثبت").font = Font(name='B Nazanin', size=12, bold=True)
            ws3.cell(row=1, column=5, value="منفی").font = Font(name='B Nazanin', size=12, bold=True)
            ws3.cell(row=1, column=6, value="الگو (بر پایهٔ نوع رفتار)").font = Font(name='B Nazanin', size=12, bold=True)
            ws3.cell(row=1, column=7, value="شناسه مشاهده‌ها").font = Font(name='B Nazanin', size=12, bold=True)
            
            row = 2
            for competency, stats in report['competency_stats'].items():
                ws3.cell(row=row, column=1, value=competency).font = Font(name='B Nazanin', size=11)
                ws3.cell(row=row, column=2, value=stats['count']).font = Font(name='B Nazanin', size=11)
                ws3.cell(row=row, column=3, value=stats['avg_severity']).font = Font(name='B Nazanin', size=11)
                ws3.cell(row=row, column=4, value=stats['positive']).font = Font(name='B Nazanin', size=11)
                ws3.cell(row=row, column=5, value=stats['negative']).font = Font(name='B Nazanin', size=11)
                
                obs_ids = ', '.join([str(i) for i in stats.get('observation_ids', [])])
                ws3.cell(row=row, column=7, value=obs_ids).font = Font(name='B Nazanin', size=9, color='7F8C8D')
                
                # بازرسی یازدهم: وضعیت هر زمینه از «نوع رفتارهای ثبت‌شده»
                # می‌آید، نه از میانگین شدت. شدت فقط در ستون جداگانه
                # به‌عنوان اطلاعات تکمیلی نمایش داده می‌شود.
                _positive = stats.get('positive', 0) or 0
                _negative = stats.get('negative', 0) or 0
                _total = stats.get('count', 0) or 0
                kind = classify_pattern(_positive, _negative,
                                        max(_total - _positive - _negative, 0), _total)
                if kind == PATTERN_STRENGTH:
                    status, color = "الگوی مثبت تکرارشونده", '27AE60'
                elif kind == PATTERN_NEEDS_ATTENTION:
                    status, color = "الگوی منفی تکرارشونده", 'E74C3C'
                elif kind == PATTERN_MIXED:
                    status, color = "الگوی ترکیبی", 'F39C12'
                else:
                    status, color = "داده ناکافی", '95A5A6'
                
                cell = ws3.cell(row=row, column=6, value=status)
                cell.font = Font(name='B Nazanin', size=11, color=color)
                row += 1
            
            ws3.column_dimensions['A'].width = 30
            ws3.column_dimensions['B'].width = 15
            ws3.column_dimensions['C'].width = 20
            ws3.column_dimensions['D'].width = 15
            ws3.column_dimensions['E'].width = 15
            ws3.column_dimensions['F'].width = 20
            ws3.column_dimensions['G'].width = 30

            # ===== برگه چهارم: مسیر رشد چندساله (بازرسی سیزدهم) =====
            # همان روایت منسجم گزارش معمول/PDF با رندر مشترک
            growth = report.get('growth_narrative') or {}
            if growth.get('has_data'):
                ws4 = wb.create_sheet("مسیر رشد چندساله")
                ws4.cell(row=1, column=1,
                         value="سابقهٔ رشد و مسیر طی‌شده (مقایسهٔ دانش‌آموز با خودش)"
                         ).font = Font(name='B Nazanin', size=14, bold=True)
                row = 3
                for kind, line in self.growth_narrative_lines(growth):
                    if kind == 'heading':
                        row += 1
                        cell = ws4.cell(row=row, column=1, value=line)
                        cell.font = Font(name='B Nazanin', size=12, bold=True, color='2C3E50')
                    elif kind == 'bullet':
                        ws4.cell(row=row, column=1, value=f"• {line}").font = Font(
                            name='B Nazanin', size=11)
                    else:
                        ws4.cell(row=row, column=1, value=line).font = Font(
                            name='B Nazanin', size=11)
                    ws4.cell(row=row, column=1).alignment = Alignment(wrap_text=True)
                    row += 1
                ws4.column_dimensions['A'].width = 120
            
            wb.save(file_path)
            return True, f"فایل با موفقیت در {file_path} ذخیره شد."
            
        except Exception as e:
            self.logger.debug(f"خطای مدیریت‌شده در export_to_excel (مسیر جایگزین): {e}")
            return False, "خطا در ساخت فایل Excel."