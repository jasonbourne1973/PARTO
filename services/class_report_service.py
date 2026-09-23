"""
سرویس گزارش کلاس - بدون رتبه‌بندی و مقایسه
نمایش وضعیت کلی کلاس بر اساس داده‌های ثبت‌شده
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import jdatetime

from dal.academic_year_dal import AcademicYearDAL
from dal.class_dal import ClassDAL
from dal.competency_dal import CompetencyDAL
from dal.observation_dal import ObservationDAL
from dal.staff_dal import StaffDAL
from dal.student_academic_profile_dal import StudentAcademicProfileDAL
from dal.student_dal import StudentDAL
from services.base_service import BaseService
from utils.behavior_analysis import classify_pattern, pattern_label
from utils.error_handler import ServiceError
from utils.logger import get_logger
from utils.time_utils import utc_now


class ClassReportService(BaseService):
    """
    سرویس گزارش کلاس - بدون رتبه‌بندی
    
    ویژگی‌ها:
    - نمایش وضعیت کلی کلاس
    - آمار مشاهدات، مداخلات و پیگیری‌ها
    - تحلیل شایستگی‌ها در سطح کلاس
    - روند تغییرات در طول زمان
    - بدون مقایسه دانش‌آموزان با یکدیگر
    """
    
    def __init__(self):
        super().__init__()
        self.class_dal = ClassDAL()
        self.observation_dal = ObservationDAL()
        self.student_dal = StudentDAL()
        self.profile_dal = StudentAcademicProfileDAL()
        self.competency_dal = CompetencyDAL()
        self.staff_dal = StaffDAL()
        self.academic_year_dal = AcademicYearDAL()
        self.logger = get_logger(self.__class__.__name__)
    
    def get_class_report(self, class_id, start_date=None, end_date=None):
        """
        دریافت گزارش کامل یک کلاس
        
        Args:
            class_id: شناسه کلاس
            start_date: تاریخ شروع (اختیاری)
            end_date: تاریخ پایان (اختیاری)
        
        Returns:
            dict: گزارش کامل کلاس
        """
        try:
            # دریافت اطلاعات کلاس
            class_obj = self.class_dal.get_by_id(class_id)
            if not class_obj:
                raise ServiceError(f"کلاس با شناسه {class_id} یافت نشد.")
            
            # دریافت معلم کلاس
            teacher_name = None
            if class_obj.teacher_id:
                teacher = self.staff_dal.get_by_id(class_obj.teacher_id)
                if teacher:
                    teacher_name = teacher.full_name
            
            # دریافت سال تحصیلی
            academic_year = None
            if class_obj.academic_year_id:
                academic_year = self.academic_year_dal.get_by_id(class_obj.academic_year_id)
            
            # دریافت آمار کلاس
            class_summary = self.class_dal.get_class_summary(class_id, start_date, end_date)
            
            if not class_summary:
                return {
                    'class': class_obj,
                    'teacher_name': teacher_name,
                    'academic_year': academic_year,
                    'has_data': False,
                    'message': 'داده‌ای برای این کلاس وجود ندارد.'
                }
            
            # دریافت روند
            trend_data = self.observation_dal.get_trend_by_class(
                class_obj.name, 'monthly', start_date, end_date
            )
            
            # دریافت آمار مداخلات و پیگیری‌ها
            intervention_stats = self._get_class_intervention_stats(
                class_obj.name, class_obj.academic_year_id, start_date, end_date
            )
            followup_stats = self._get_class_followup_stats(
                class_obj.name, class_obj.academic_year_id, start_date, end_date
            )
            
            # تحلیل شایستگی‌ها
            competency_analysis = self._analyze_competencies(class_summary.get('competency_stats', {}))
            
            # تولید پیشنهادات
            recommendations = self._generate_recommendations(class_summary, competency_analysis)
            
            return {
                'class': class_obj,
                'teacher_name': teacher_name,
                'academic_year': academic_year,
                'student_count': class_summary.get('student_count', 0),
                'observations_stats': class_summary.get('observations_stats', {}),
                'competency_stats': class_summary.get('competency_stats', {}),
                'student_stats': class_summary.get('student_stats', []),
                'trend_data': trend_data,
                'intervention_stats': intervention_stats,
                'followup_stats': followup_stats,
                'competency_analysis': competency_analysis,
                'recommendations': recommendations,
                'has_data': class_summary.get('has_data', False),
                'start_date': start_date,
                'end_date': end_date
            }
            
        except Exception as e:
            self.logger.error(f"خطا در دریافت گزارش کلاس: {e}")
            raise ServiceError(f"خطا در دریافت گزارش: {e!s}")
    
    def get_class_list_report(self, academic_year_id=None, start_date=None, end_date=None):
        """
        دریافت گزارش لیست کلاس‌ها
        
        Args:
            academic_year_id: شناسه سال تحصیلی (اختیاری)
            start_date: تاریخ شروع (اختیاری)
            end_date: تاریخ پایان (اختیاری)
        
        Returns:
            list: لیست گزارش‌های کلاس‌ها
        """
        try:
            classes = self.class_dal.get_all(academic_year_id)
            reports = []
            
            for class_obj in classes:
                report = self.get_class_report(class_obj.id, start_date, end_date)
                if report:
                    reports.append(report)
            
            return reports
            
        except Exception as e:
            self.logger.error(f"خطا در دریافت گزارش لیست کلاس‌ها: {e}")
            raise ServiceError(f"خطا در دریافت گزارش: {e!s}")
    
    def get_class_summary_for_dashboard(self, class_id):
        """
        دریافت خلاصه کلاس برای داشبورد
        
        Args:
            class_id: شناسه کلاس
        
        Returns:
            dict: خلاصه اطلاعات
        """
        try:
            class_obj = self.class_dal.get_by_id(class_id)
            if not class_obj:
                return None
            
            # دریافت آمار سریع
            stats = self.class_dal.get_class_observations_stats(class_id)
            
            if not stats:
                return {
                    'class_name': class_obj.display_name,
                    'total_observations': 0,
                    'students_count': 0,
                    'has_data': False
                }
            
            return {
                'class_id': class_obj.id,
                'class_name': class_obj.display_name,
                'grade': class_obj.grade,
                'teacher_id': class_obj.teacher_id,
                'total_observations': stats.get('total', 0),
                'positive': stats.get('positive', 0),
                'negative': stats.get('negative', 0),
                'neutral': stats.get('neutral', 0),
                'students_count': len(stats.get('student_ids', [])),
                'has_data': stats.get('total', 0) > 0
            }
            
        except Exception as e:
            self.logger.error(f"خطا در دریافت خلاصه کلاس: {e}")
            return None
    
    def _get_class_intervention_stats(
        self, class_name, academic_year_id, start_date=None, end_date=None
    ):
        """دریافت آمار مداخلات همان کلاس در همان سال تحصیلی."""
        try:
            conn = self.db.get_connection()
            cursor = conn.cursor()
            
            query = """
                SELECT i.*
                FROM interventions i
                JOIN student_academic_profiles sap ON i.student_profile_id = sap.id
                WHERE sap.class_name = ?
                AND sap.academic_year_id = ?
                AND i.is_deleted = 0
            """
            params = [class_name, academic_year_id]
            
            if start_date:
                query += " AND i.date >= ?"
                params.append(start_date)
            if end_date:
                query += " AND i.date <= ?"
                params.append(end_date)
            
            cursor.execute(query, params)
            rows = cursor.fetchall()
            
            total = len(rows)
            if total == 0:
                return {'total': 0, 'planned': 0, 'in_progress': 0, 'completed': 0, 'cancelled': 0}
            
            planned = sum(1 for r in rows if r['status'] == 'planned')
            in_progress = sum(1 for r in rows if r['status'] == 'in_progress')
            completed = sum(1 for r in rows if r['status'] in ['completed', 'done'])
            cancelled = sum(1 for r in rows if r['status'] == 'cancelled')
            
            return {
                'total': total,
                'planned': planned,
                'in_progress': in_progress,
                'completed': completed,
                'cancelled': cancelled
            }
            
        except Exception as e:
            self.logger.error(f"خطا در دریافت آمار مداخلات کلاس: {e}")
            return {'total': 0, 'planned': 0, 'in_progress': 0, 'completed': 0, 'cancelled': 0}
    
    def _get_class_followup_stats(
        self, class_name, academic_year_id, start_date=None, end_date=None
    ):
        """دریافت آمار پیگیری‌های همان کلاس در همان سال تحصیلی."""
        try:
            conn = self.db.get_connection()
            cursor = conn.cursor()
            
            query = """
                SELECT f.*
                FROM followups f
                JOIN interventions i ON f.intervention_id = i.id
                JOIN student_academic_profiles sap ON i.student_profile_id = sap.id
                WHERE sap.class_name = ?
                AND sap.academic_year_id = ?
                AND f.is_deleted = 0
            """
            params = [class_name, academic_year_id]
            
            if start_date:
                query += " AND f.date >= ?"
                params.append(start_date)
            if end_date:
                query += " AND f.date <= ?"
                params.append(end_date)
            
            cursor.execute(query, params)
            rows = cursor.fetchall()
            
            total = len(rows)
            if total == 0:
                return {'total': 0, 'pending': 0, 'done': 0, 'continued': 0, 'closed': 0, 'cancelled': 0}
            
            pending = sum(1 for r in rows if r['status'] == 'pending')
            done = sum(1 for r in rows if r['status'] == 'done')
            continued = sum(1 for r in rows if r['status'] == 'continued')
            closed = sum(1 for r in rows if r['status'] == 'closed')
            cancelled = sum(1 for r in rows if r['status'] == 'cancelled')
            
            return {
                'total': total,
                'pending': pending,
                'done': done,
                'continued': continued,
                'closed': closed,
                'cancelled': cancelled
            }
            
        except Exception as e:
            self.logger.error(f"خطا در دریافت آمار پیگیری‌های کلاس: {e}")
            return {'total': 0, 'pending': 0, 'done': 0, 'continued': 0, 'closed': 0, 'cancelled': 0}
    
    def _analyze_competencies(self, competency_stats):
        """
        تحلیل زمینه‌های کلاس بر پایهٔ **نوع رفتار ثبت‌شده** (بازرسی یازدهم)

        پیش از این، قوت/ضعف هر زمینه با آستانهٔ «میانگین شدت» تعیین
        می‌شد؛ اکنون مبنا الگوی تکرارشوندهٔ رفتارهای مثبت/منفی است و
        شدت فقط اطلاعات تکمیلی است. این تحلیل در سطح **کلاس** است و
        رتبه‌بندی یا مقایسهٔ دانش‌آموزان با یکدیگر نیست.
        """
        if not competency_stats:
            return {
                'strong': [],
                'weak': [],
                'mixed': [],
                'total': 0,
                'has_data': False
            }

        strong = []
        weak = []
        mixed = []

        for name, stats in competency_stats.items():
            count = stats.get('count', 0)
            positive = stats.get('positive', 0)
            negative = stats.get('negative', 0)
            neutral = max(count - positive - negative, 0)
            kind = classify_pattern(positive, negative, neutral, count)
            entry = {
                'name': name,
                'count': count,
                'positive': positive,
                'negative': negative,
                'pattern': kind,
                'pattern_label': pattern_label(kind),
                # شدت صرفاً تکمیلی
                'avg_severity': stats.get('avg_severity', 0),
                'positive_ratio': positive / count if count else 0,
                'negative_ratio': negative / count if count else 0,
            }
            if kind == 'strength':
                strong.append(entry)
            elif kind == 'needs_attention':
                weak.append(entry)
            elif kind == 'mixed':
                mixed.append(entry)

        # مرتب‌سازی بر پایهٔ تعداد رفتار جهت‌دار (نه شدت)
        strong.sort(key=lambda x: (x['positive'], x['count']), reverse=True)
        weak.sort(key=lambda x: (x['negative'], x['count']), reverse=True)
        mixed.sort(key=lambda x: x['count'], reverse=True)

        return {
            'strong': strong[:5],  # الگوهای تکرارشوندهٔ رفتار مثبت در کلاس
            'weak': weak[:5],      # الگوهای تکرارشوندهٔ رفتار منفی در کلاس
            'mixed': mixed[:5],
            'total': len(competency_stats),
            'has_data': len(competency_stats) > 0
        }
    
    def _generate_recommendations(self, class_summary, competency_analysis):
        """
        تولید پیشنهادات بر اساس داده‌های کلاس
        بدون رتبه‌بندی و مقایسه
        """
        recommendations = {
            'teacher': [],
            'counselor': [],
            'general': []
        }
        
        obs_stats = class_summary.get('observations_stats', {})
        total_obs = obs_stats.get('total', 0)
        positive = obs_stats.get('positive', 0)
        negative = obs_stats.get('negative', 0)
        
        # ===== پیشنهادات بر اساس تعداد مشاهدات =====
        if total_obs == 0:
            recommendations['teacher'].append(
                "⚠️ هیچ مشاهده‌ای برای این کلاس ثبت نشده است. "
                "ثبت مشاهدات منظم به تحلیل دقیق‌تر وضعیت کلاس کمک می‌کند."
            )
        elif total_obs < 10:
            recommendations['teacher'].append(
                f"📊 تعداد مشاهدات ثبت‌شده ({total_obs} مورد) کمتر از حد مطلوب است. "
                "ثبت مشاهدات بیشتر به تحلیل دقیق‌تر وضعیت کلاس کمک می‌کند."
            )
        
        # ===== پیشنهادات بر اساس نسبت مثبت/منفی =====
        if total_obs > 0:
            positive_ratio = positive / total_obs
            
            if positive_ratio >= 0.6:
                recommendations['teacher'].append(
                    f"✅ {positive_ratio*100:.0f}% مشاهدات کلاس مثبت است. "
                    "فضای آموزشی کلاس مطلوب است. به تقویت رفتارهای مثبت ادامه دهید."
                )
            elif positive_ratio >= 0.4:
                recommendations['teacher'].append(
                    f"🟡 {positive_ratio*100:.0f}% از رفتارهای ثبت‌شدهٔ کلاس مثبت "
                    "است. با تقویت رفتارهای مثبت و تشویق دانش‌آموزان می‌توان "
                    "این نسبت را بهتر کرد."
                )
            else:
                recommendations['teacher'].append(
                    f"🔴 {positive_ratio*100:.0f}% از رفتارهای ثبت‌شدهٔ کلاس مثبت "
                    "است. بررسی زمینه‌ها و هماهنگی با مشاور مدرسه پیشنهاد می‌شود؛ "
                    "این آمار «الگوی مشاهده‌شده» است، نه تشخیص."
                )
        
        # ===== پیشنهادات بر اساس شایستگی‌ها =====
        weak_competencies = competency_analysis.get('weak', [])
        if weak_competencies:
            weak_names = [
                f"{w['name']} ({w.get('negative', 0)} رفتار منفی از "
                f"{w.get('count', 0)} مشاهده)" for w in weak_competencies[:3]
            ]
            recommendations['teacher'].append(
                f"📋 زمینه‌های با الگوی تکرارشوندهٔ رفتار منفی در کلاس:\n"
                f"{chr(10).join(['   • ' + name for name in weak_names])}\n"
                f"بررسی این الگوها و طراحی فعالیت‌های هدفمند پیشنهاد می‌شود."
            )
        
        strong_competencies = competency_analysis.get('strong', [])
        if strong_competencies:
            strong_names = [
                f"{s['name']} ({s.get('positive', 0)} رفتار مثبت از "
                f"{s.get('count', 0)} مشاهده)" for s in strong_competencies[:3]
            ]
            recommendations['teacher'].append(
                f"⭐ زمینه‌های با الگوی تکرارشوندهٔ رفتار مثبت در کلاس:\n"
                f"{chr(10).join(['   • ' + name for name in strong_names])}\n"
                f"تقویت این الگوها پیشنهاد می‌شود."
            )
        
        # ===== پیشنهادات بر اساس آمار دانش‌آموزان =====
        student_stats = class_summary.get('student_stats', [])
        if student_stats:
            students_without_obs = [s for s in student_stats if s.get('observations_count', 0) == 0]
            if students_without_obs:
                names = [s['student_name'] for s in students_without_obs[:5]]
                recommendations['teacher'].append(
                    f"👤 {len(students_without_obs)} دانش‌آموز هیچ مشاهده‌ای ندارند:\n"
                    f"{chr(10).join(['   • ' + name for name in names])}\n"
                    f"توجه ویژه به این دانش‌آموزان برای ثبت مشاهدات توصیه می‌شود."
                )
        
        # ===== پیشنهادات برای مشاور =====
        if total_obs > 20 and negative > positive * 0.5:
            recommendations['counselor'].append(
                f"📊 {negative} مشاهده منفی در این کلاس ثبت شده است. "
                f"بررسی الگوهای رفتاری و ارائه راهکارهای تخصصی توصیه می‌شود."
            )
        
        if len(weak_competencies) >= 3:
            recommendations['counselor'].append(
                f"📋 {len(weak_competencies)} شایستگی نیازمند توجه در این کلاس شناسایی شده است. "
                f"طراحی برنامه‌های حمایتی تخصصی توصیه می‌شود."
            )
        
        # ===== پیشنهادات عمومی =====
        if not recommendations['teacher'] and not recommendations['counselor']:
            recommendations['general'].append(
                "✅ وضعیت عمومی کلاس مطلوب است. "
                "به ثبت مستمر مشاهدات و پیگیری‌های منظم ادامه دهید."
            )
        
        return recommendations
    
    def export_class_report_pdf(self, class_id, file_path, start_date=None, end_date=None):
        """
        خروجی گزارش کلاس به صورت PDF
        
        Args:
            class_id: شناسه کلاس
            file_path: مسیر ذخیره فایل
            start_date: تاریخ شروع (اختیاری)
            end_date: تاریخ پایان (اختیاری)
        
        Returns:
            tuple: (success, message)
        """
        try:
            report_data = self.get_class_report(class_id, start_date, end_date)
            if not report_data or not report_data.get('has_data', False):
                return False, "داده‌ای برای تولید گزارش وجود ندارد."
            
            from utils.persian_pdf import PersianPDF
            
            pdf = PersianPDF(file_path)
            
            # ===== عنوان =====
            class_obj = report_data['class']
            pdf.add_title(f"گزارش کلاس {class_obj.display_name}")
            pdf.add_spacer(0.2)
            
            # ===== اطلاعات کلاس =====
            info_items = [
                f"نام کلاس: {class_obj.display_name}",
                f"پایه: {class_obj.grade_display if hasattr(class_obj, 'grade_display') else class_obj.grade}",
                f"معلم: {report_data.get('teacher_name', 'نامشخص')}",
                f"تعداد دانش‌آموزان: {report_data.get('student_count', 0)}",
            ]
            
            if report_data.get('academic_year'):
                info_items.append(f"سال تحصیلی: {report_data['academic_year'].title}")
            
            if start_date and end_date:
                info_items.append(f"بازه زمانی: {start_date} تا {end_date}")
            
            for item in info_items:
                pdf.add_text(item)
            pdf.add_spacer(0.3)
            
            # ===== خلاصه آماری =====
            obs_stats = report_data.get('observations_stats', {})
            pdf.add_subtitle("خلاصه آماری")
            pdf.add_text(f"کل مشاهدات: {obs_stats.get('total', 0)}")
            pdf.add_text(f"مشاهدات مثبت: {obs_stats.get('positive', 0)}")
            pdf.add_text(f"مشاهدات منفی: {obs_stats.get('negative', 0)}")
            pdf.add_text(f"مشاهدات خنثی: {obs_stats.get('neutral', 0)}")
            
            if obs_stats.get('total', 0) > 0:
                avg_severity = obs_stats.get('avg_severity', 0)
                pdf.add_text(f"میانگین شدت مشاهدات: {avg_severity}")
            pdf.add_spacer(0.3)
            
            # ===== آمار مداخلات و پیگیری‌ها =====
            inter_stats = report_data.get('intervention_stats', {})
            follow_stats = report_data.get('followup_stats', {})
            
            if inter_stats.get('total', 0) > 0:
                pdf.add_subtitle("مداخلات")
                pdf.add_text(f"کل مداخلات: {inter_stats.get('total', 0)}")
                pdf.add_text(f"در حال اجرا: {inter_stats.get('in_progress', 0)}")
                pdf.add_text(f"تکمیل شده: {inter_stats.get('completed', 0)}")
                pdf.add_spacer(0.2)
            
            if follow_stats.get('total', 0) > 0:
                pdf.add_subtitle("پیگیری‌ها")
                pdf.add_text(f"کل پیگیری‌ها: {follow_stats.get('total', 0)}")
                pdf.add_text(f"در انتظار: {follow_stats.get('pending', 0)}")
                pdf.add_text(f"انجام شده: {follow_stats.get('done', 0)}")
                pdf.add_spacer(0.3)
            
            # ===== شایستگی‌ها =====
            comp_analysis = report_data.get('competency_analysis', {})
            if comp_analysis.get('has_data', False):
                pdf.add_subtitle("تحلیل شایستگی‌ها")
                
                strong = comp_analysis.get('strong', [])
                weak = comp_analysis.get('weak', [])
                
                if strong:
                    pdf.add_bold("⭐ شایستگی‌های برتر:")
                    for item in strong[:5]:
                        pdf.add_text(f"• {item['name']} (میانگین شدت: {item['avg_severity']})")
                    pdf.add_spacer(0.1)
                
                if weak:
                    pdf.add_bold("🔴 شایستگی‌های نیازمند توجه:")
                    for item in weak[:5]:
                        pdf.add_text(f"• {item['name']} (میانگین شدت: {item['avg_severity']})")
                    pdf.add_spacer(0.1)
            
            # ===== روند =====
            trend_data = report_data.get('trend_data', [])
            if trend_data:
                pdf.add_subtitle("روند تغییرات")
                for item in trend_data[-6:]:  # ۶ ماه اخیر
                    pdf.add_text(
                        f"{item['label']}: {item['total']} مشاهده "
                        f"(مثبت: {item['positive']}, منفی: {item['negative']})"
                    )
                pdf.add_spacer(0.3)
            
            # ===== پیشنهادات =====
            recommendations = report_data.get('recommendations', {})
            pdf.add_subtitle("پیشنهادات")
            
            if recommendations.get('teacher'):
                pdf.add_bold("👨‍🏫 به معلم:")
                for rec in recommendations['teacher'][:5]:
                    pdf.add_text(f"• {rec}")
                pdf.add_spacer(0.2)
            
            if recommendations.get('counselor'):
                pdf.add_bold("🫂 به مشاور:")
                for rec in recommendations['counselor'][:3]:
                    pdf.add_text(f"• {rec}")
                pdf.add_spacer(0.2)
            
            if recommendations.get('general'):
                pdf.add_bold("💡 عمومی:")
                for rec in recommendations['general'][:3]:
                    pdf.add_text(f"• {rec}")
                pdf.add_spacer(0.3)
            
            # ===== متادیتا =====
            pdf.add_separator()
            try:
                today = jdatetime.date.today()
                date_str = f"{today.year:04d}/{today.month:02d}/{today.day:02d}"
            except Exception as _exc:
                self.logger.debug(f"خطای مدیریت‌شده در export_class_report_pdf (مسیر جایگزین): {_exc}")
                date_str = utc_now().strftime("%Y/%m/%d")
            
            pdf.add_text(f"تاریخ تهیه گزارش: {date_str}")
            pdf.add_text("PARTO - سامانه مدیریت پرونده دانش‌آموزان")
            pdf.add_text("گزارش کلاس - بدون رتبه‌بندی")
            
            pdf.build(file_path)
            return True, f"فایل PDF با موفقیت در {file_path} ذخیره شد."
            
        except Exception as e:
            self.logger.error(f"خطا در خروجی PDF گزارش کلاس: {e}")
            return False, f"خطا در ساخت فایل PDF: {e!s}"
    
    def export_class_report_excel(self, class_id, file_path, start_date=None, end_date=None):
        """
        خروجی گزارش کلاس به صورت Excel
        
        Args:
            class_id: شناسه کلاس
            file_path: مسیر ذخیره فایل
            start_date: تاریخ شروع (اختیاری)
            end_date: تاریخ پایان (اختیاری)
        
        Returns:
            tuple: (success, message)
        """
        try:
            report_data = self.get_class_report(class_id, start_date, end_date)
            if not report_data:
                return False, "داده‌ای برای تولید گزارش وجود ندارد."
            
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Alignment, Font, PatternFill
                from openpyxl.utils import get_column_letter
            except ImportError as _exc:
                self.logger.debug(f"خطای مدیریت‌شده در export_class_report_excel (مسیر جایگزین): {_exc}")
                return False, "کتابخانه openpyxl نصب نیست. pip install openpyxl"
            
            wb = Workbook()
            
            # ===== برگه ۱: خلاصه =====
            ws1 = wb.active
            ws1.title = "خلاصه"
            
            class_obj = report_data['class']
            
            # عنوان
            ws1.cell(row=1, column=1, value=f"گزارش کلاس {class_obj.display_name}")
            ws1.cell(row=1, column=1).font = Font(name='B Nazanin', size=16, bold=True)
            ws1.merge_cells('A1:C1')
            
            # اطلاعات
            info_data = [
                ("نام کلاس", class_obj.display_name),
                ("پایه", class_obj.grade_display if hasattr(class_obj, 'grade_display') else class_obj.grade),
                ("معلم", report_data.get('teacher_name', 'نامشخص')),
                ("تعداد دانش‌آموزان", report_data.get('student_count', 0)),
            ]
            
            if report_data.get('academic_year'):
                info_data.append(("سال تحصیلی", report_data['academic_year'].title))
            
            row = 3
            for label, value in info_data:
                ws1.cell(row=row, column=1, value=label).font = Font(name='B Nazanin', size=11, bold=True)
                ws1.cell(row=row, column=2, value=value).font = Font(name='B Nazanin', size=11)
                row += 1
            
            # آمار
            obs_stats = report_data.get('observations_stats', {})
            row += 2
            ws1.cell(row=row, column=1, value="آمار مشاهدات").font = Font(name='B Nazanin', size=12, bold=True)
            row += 1
            
            stats_data = [
                ("کل مشاهدات", obs_stats.get('total', 0)),
                ("مشاهدات مثبت", obs_stats.get('positive', 0)),
                ("مشاهدات منفی", obs_stats.get('negative', 0)),
                ("مشاهدات خنثی", obs_stats.get('neutral', 0)),
                ("میانگین شدت", obs_stats.get('avg_severity', 0)),
            ]
            
            for label, value in stats_data:
                ws1.cell(row=row, column=1, value=label)
                ws1.cell(row=row, column=2, value=value)
                row += 1
            
            ws1.column_dimensions['A'].width = 30
            ws1.column_dimensions['B'].width = 20
            
            # ===== برگه ۲: دانش‌آموزان =====
            ws2 = wb.create_sheet("دانش‌آموزان")
            
            headers = ["ردیف", "نام دانش‌آموز", "تعداد مشاهدات", "مثبت", "منفی", "خنثی", "میانگین شدت", "وضعیت"]
            for col, header in enumerate(headers, 1):
                cell = ws2.cell(row=1, column=col, value=header)
                cell.font = Font(name='B Nazanin', size=11, bold=True)
                cell.fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            student_stats = report_data.get('student_stats', [])
            for row, student in enumerate(student_stats, 2):
                ws2.cell(row=row, column=1, value=row-1)
                ws2.cell(row=row, column=2, value=student.get('student_name', 'نامشخص'))
                ws2.cell(row=row, column=3, value=student.get('observations_count', 0))
                ws2.cell(row=row, column=4, value=student.get('positive', 0))
                ws2.cell(row=row, column=5, value=student.get('negative', 0))
                ws2.cell(row=row, column=6, value=student.get('neutral', 0))
                ws2.cell(row=row, column=7, value=student.get('avg_severity', 0))
                ws2.cell(row=row, column=8, value=student.get('status', 'بدون مشاهده'))
            
            for col in range(1, 9):
                ws2.column_dimensions[get_column_letter(col)].width = 18
            
            # ===== برگه ۳: شایستگی‌ها =====
            ws3 = wb.create_sheet("شایستگی‌ها")
            
            headers = ["شایستگی", "تعداد", "میانگین شدت", "مثبت", "منفی",
                   "الگو (بر پایهٔ نوع رفتار)"]
            for col, header in enumerate(headers, 1):
                cell = ws3.cell(row=1, column=col, value=header)
                cell.font = Font(name='B Nazanin', size=11, bold=True)
                cell.fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            comp_stats = report_data.get('competency_stats', {})
            row = 2
            for name, stats in comp_stats.items():
                ws3.cell(row=row, column=1, value=name)
                ws3.cell(row=row, column=2, value=stats.get('count', 0))
                ws3.cell(row=row, column=3, value=stats.get('avg_severity', 0))
                ws3.cell(row=row, column=4, value=stats.get('positive', 0))
                ws3.cell(row=row, column=5, value=stats.get('negative', 0))
                
                # وضعیت بر پایهٔ نوع رفتار ثبت‌شده (نه میانگین شدت)
                status = pattern_label(classify_pattern(
                    stats.get('positive', 0), stats.get('negative', 0),
                    max(stats.get('count', 0) - stats.get('positive', 0)
                        - stats.get('negative', 0), 0),
                    stats.get('count', 0)))
                ws3.cell(row=row, column=6, value=status)
                row += 1
            
            for col in range(1, 7):
                ws3.column_dimensions[get_column_letter(col)].width = 20
            
            wb.save(file_path)
            return True, f"فایل Excel با موفقیت در {file_path} ذخیره شد."
            
        except Exception as e:
            self.logger.error(f"خطا در خروجی Excel گزارش کلاس: {e}")
            return False, f"خطا در ساخت فایل Excel: {e!s}"