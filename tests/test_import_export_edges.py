"""
آزمون‌های لبه‌های ایمپورت/خروجی Excel (دور هفدهم — مرحلهٔ ۴، بندهای ۱۴، ۱۵، ۱۹)

این فایل Qt لازم ندارد؛ اجرای واقعی صفحه‌ها و دیالوگ‌ها در
`verify_fixes17.py §D` انجام می‌شود.

پوشش:
  • تکراری داخل یک فایل (کد ملی یکسان در دو ردیف)
  • کد ملی/تلفن نامعتبر (عددی نبودن، طول نادرست)
  • ارقام فارسی/عربی (یکدست‌سازی به ASCII تا تکراری‌یابی و جست‌وجو کار کند)
  • تاریخ تولد نامعتبر، سلول تاریخ/عددِ اکسل، و نرمال‌سازی «1395/1/1»
  • متن یونیکد/فارسی با نیم‌فاصله و آدرس بلند
  • فایل بزرگ (۳۰۰ ردیف)
  • شکست ساخت پروندهٔ سالانه برای یک ردیف (بدون دانش‌آموز یتیم)
  • صفر رکورد واردشده = شکست صریح (نه پیام موفقیت)
  • فایل CSV با پسوند .xlsx → شکست صریح (هرگز CSV جا زده نمی‌شود)
  • خروجی: فایل واقعی xlsx، همهٔ ردیف‌های داده‌شده (نه صفحهٔ جاری)، و
    برچسب صریح وقتی فهرست «حذف‌شده‌ها» صادر می‌شود

جداسازی: هر تست روی دیتابیس موقت تازه اجرا می‌شود و در پایان وضعیت قبلی
برگردانده می‌شود.
"""

import contextlib
import io
import os
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import database.connection as _dbc
from config import settings as _settings
from dal.academic_year_dal import AcademicYearDAL
from dal.student_academic_profile_dal import StudentAcademicProfileDAL
from dal.student_dal import StudentDAL
from models.academic_year import AcademicYear
from models.student import Student
from utils.excel_importer import ExcelImporter

HEADERS = ["نام", "نام خانوادگی", "کد ملی", "تاریخ تولد", "نام پدر",
           "نام ولی", "تلفن ولی", "آدرس", "پایه", "کلاس"]


class _TempDbTestCase(unittest.TestCase):
    """دیتابیس موقت برای هر تست + بازگردانی وضعیت قبلی"""

    def setUp(self):
        self._saved = (
            _settings.DB_PATH,
            _dbc.DB_PATH,
            _dbc.DatabaseConnection._instance,
            _dbc.DatabaseConnection._connection,
            _dbc.DatabaseConnection._initialized,
        )
        self.tmp = tempfile.mkdtemp(prefix="partow_import_")
        _settings.DB_PATH = os.path.join(self.tmp, "partow.db")
        _dbc.DB_PATH = _settings.DB_PATH
        _dbc.DatabaseConnection._instance = None
        _dbc.DatabaseConnection._connection = None
        _dbc.DatabaseConnection._initialized = False
        with contextlib.redirect_stdout(io.StringIO()):
            self.conn = _dbc.DatabaseConnection().get_connection(user_id=1)

        self.importer = ExcelImporter()
        self.year = self._make_year("1403-1404", "1403/07/01", "1404/03/31")

    def tearDown(self):
        (_settings.DB_PATH, _dbc.DB_PATH, _dbc.DatabaseConnection._instance,
         _dbc.DatabaseConnection._connection,
         _dbc.DatabaseConnection._initialized) = self._saved

    # ------------------------------------------------------------------
    def _make_year(self, title, start, end):
        year = AcademicYear()
        year.title = title
        year.start_date = start
        year.end_date = end
        year.is_active = 1
        created = AcademicYearDAL().create(year)
        AcademicYearDAL().set_active(created.id)
        return created

    def _write_xlsx(self, rows, name="import.xlsx", headers=None, sheet=None):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers or HEADERS)
        for row in rows:
            ws.append(row)
        path = os.path.join(self.tmp, name)
        wb.save(path)
        return path

    def _import(self, path):
        with contextlib.redirect_stdout(io.StringIO()):
            return self.importer.import_students_from_excel(path, self.year.id)

    def _students(self, where="", params=()):
        query = "SELECT * FROM students"
        if where:
            query += " WHERE " + where
        return self.conn.execute(query + " ORDER BY id", params).fetchall()

    def _profiles_count(self, student_id=None):
        if student_id is None:
            return self.conn.execute(
                "SELECT COUNT(*) FROM student_academic_profiles").fetchone()[0]
        return self.conn.execute(
            "SELECT COUNT(*) FROM student_academic_profiles WHERE student_id = ?",
            (student_id,)).fetchone()[0]


class TestImportEdges(_TempDbTestCase):
    """بند ۱۴ — لبه‌های ایمپورت"""

    def test_duplicate_national_code_inside_one_file(self):
        path = self._write_xlsx([
            ["دوقلو", "اولی", "7777777777", "1395/01/01", "پدر", "ولی",
             "09121234567", "آدرس", 1, "الف"],
            ["دوقلو", "دومی", "7777777777", "1395/01/01", "پدر", "ولی",
             "09121234567", "آدرس", 1, "الف"],
        ])
        success, _message, imported, errors = self._import(path)
        self.assertTrue(success)
        self.assertEqual(imported, 1)
        self.assertEqual(len(errors), 1)
        self.assertIn("تکراری", errors[0])
        self.assertEqual(len(self._students("national_code = ?", ("7777777777",))), 1)

    def test_invalid_national_code_and_phone_are_reported(self):
        path = self._write_xlsx([
            ["کدنقص", "کوتاه", "12345", "1395/01/01", "پدر", "ولی",
             "09121234567", "آدرس", 1, "الف"],
            ["تلفن‌غلط", "حرفی", "1111111111", "1395/01/01", "پدر", "ولی",
             "abc", "آدرس", 1, "الف"],
            ["درست", "سالم", "2222222222", "1395/01/01", "پدر", "ولی",
             "09121234567", "آدرس", 1, "الف"],
        ])
        success, _message, imported, errors = self._import(path)
        self.assertTrue(success)
        self.assertEqual(imported, 1)
        self.assertEqual(len(errors), 2)
        self.assertIn("کد ملی باید دقیقاً ۱۰ رقم باشد", " ".join(errors))
        self.assertIn("شماره تماس سرپرست باید ۱۱ رقم باشد", " ".join(errors))
        self.assertEqual(len(self._students()), 1)

    def test_persian_digits_are_normalized_and_detect_duplicates(self):
        """«۱۲۳۴۵۶۷۸۹۰» و «1234567890» باید یک مقدار باشند"""
        path = self._write_xlsx([
            ["فارسی", "ارقام", "۱۲۳۴۵۶۷۸۹۰", "۱۳۹۵/۰۱/۰۱", "پدر", "ولی",
             "۰۹۱۲۱۲۳۴۵۶۷", "آدرس", 1, "الف"],
            ["ویِ‌اسکی", "ارقام", "1234567890", "1395/01/01", "پدر", "ولی",
             "09121234567", "آدرس", 1, "الف"],
        ])
        success, _message, imported, errors = self._import(path)
        self.assertTrue(success)
        self.assertEqual(imported, 1)
        self.assertEqual(len(errors), 1)
        self.assertIn("تکراری", errors[0])

        row = self._students()[0]
        self.assertEqual(row["national_code"], "1234567890")
        self.assertEqual(row["guardian_phone"], "09121234567")
        self.assertEqual(row["birth_date"], "1395/01/01")

    def test_invalid_birth_date_and_excel_date_cell_are_rejected(self):
        import datetime as _datetime

        path = self._write_xlsx([
            ["تاریخ‌غلط", "ماه‌سیزده", "3333333333", "1395/13/45", "پدر", "ولی",
             "", "آدرس", 1, "الف"],
            ["تاریخ", "عددی", "4444444444", "", "پدر", "ولی", "", "آدرس", 1, "الف"],
            ["تاریخ", "درست", "4444444445", "1395/01/01", "پدر", "ولی",
             "", "آدرس", 1, "الف"],
        ])
        # ردیف دوم: سلول واقعی «تاریخ» اکسل (datetime) — نباید بی‌صدا ذخیره شود
        wb = openpyxl.load_workbook(path)
        wb.active.cell(row=3, column=4, value=_datetime.datetime(2016, 3, 21))
        wb.save(path)

        _success, _message, imported, errors = self._import(path)
        self.assertEqual(imported, 1)
        self.assertEqual(len(errors), 2)
        self.assertIn("معتبر نیست", errors[0])
        self.assertIn("باید متن باشد", " ".join(errors))
        rows = self._students()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["last_name"], "درست")
        self.assertEqual(rows[0]["birth_date"], "1395/01/01")

    def test_valid_date_is_normalized_and_unicode_text_is_preserved(self):
        address = "خیابان ولیعصر، کوچهٔ ۱۲، پلاک ۳۴، واحد ۵ — نزدیک مسجد"
        path = self._write_xlsx([
            ["علی‌اکبر", "محمّدزاده", "", "1395/1/1", "حسین‌علی",
             "زهرا‌سادات", "", address, 3, "سوم-الف"],
        ])
        success, _message, imported, errors = self._import(path)
        self.assertTrue(success)
        self.assertEqual(imported, 1, errors)
        row = self._students()[0]
        self.assertEqual(row["first_name"], "علی‌اکبر")
        self.assertEqual(row["last_name"], "محمّدزاده")
        self.assertEqual(row["father_name"], "حسین‌علی")
        self.assertEqual(row["guardian_name"], "زهرا‌سادات")
        self.assertEqual(row["address"], address)
        self.assertEqual(row["birth_date"], "1395/01/01")

    def test_large_file_imports_every_row(self):
        rows = [[f"نام{i}", f"خانوادگی{i}", f"{2000000000 + i}", "1395/01/01",
                 "پدر", "ولی", "09121234567", "آدرس", 1, "الف"]
                for i in range(300)]
        success, _message, imported, errors = self._import(self._write_xlsx(rows))
        self.assertTrue(success)
        self.assertEqual(imported, 300, errors[:3])
        self.assertEqual(len(self._students()), 300)
        self.assertEqual(self._profiles_count(), 300)

    def test_profile_failure_leaves_no_orphan_and_other_rows_import(self):
        path = self._write_xlsx([
            ["بدشانس", "پرونده", "5555555555", "1395/01/01", "پدر", "ولی",
             "", "آدرس", 1, "الف"],
            ["خوش‌شانس", "پرونده", "6666666666", "1395/01/01", "پدر", "ولی",
             "", "آدرس", 1, "الف"],
        ])
        real_create = StudentAcademicProfileDAL.create
        calls = {"count": 0}

        def fake_create(self_dal, profile):
            calls["count"] += 1
            if calls["count"] == 1:
                raise RuntimeError("پرونده ساخته نشد")
            return real_create(self_dal, profile)

        StudentAcademicProfileDAL.create = fake_create
        try:
            _success, _message, imported, errors = self._import(path)
        finally:
            StudentAcademicProfileDAL.create = real_create

        self.assertEqual(imported, 1)
        self.assertEqual(len(errors), 1)
        self.assertIn("پرونده ساخته نشد", errors[0])
        self.assertEqual(len(self._students()), 1)              # دانش‌آموز یتیم نماند
        self.assertEqual(len(self._students("first_name = ?", ("خوش‌شانس",))), 1)
        self.assertEqual(self._profiles_count(), 1)

    def test_zero_imported_is_not_reported_as_success(self):
        rows = [["", "بدون‌نام", "8888888888", "1395/01/01", "پدر", "ولی",
                 "", "آدرس", 1, "الف"]]
        success, message, imported, errors = self._import(self._write_xlsx(rows))
        self.assertFalse(success)
        self.assertEqual(imported, 0)
        self.assertTrue(errors)
        self.assertIn("هیچ دانش‌آموزی ایمپورت نشد", message)
        self.assertNotIn("✅", message)
        self.assertEqual(len(self._students()), 0)

    def test_csv_renamed_to_xlsx_is_not_accepted(self):
        path = os.path.join(self.tmp, "fake.xlsx")
        with open(path, "w", encoding="utf-8") as fh:
            fh.write("نام,نام خانوادگی,کد ملی\nالف,ب,1234567890\n")
        success, message, imported, _errors = self._import(path)
        self.assertFalse(success)
        self.assertEqual(imported, 0)
        self.assertTrue(message)
        self.assertEqual(len(self._students()), 0)


class TestExportEdges(_TempDbTestCase):
    """بند ۱۵ و ۱۹ — خروجی Excel"""

    def _add_students(self, count):
        students = []
        for index in range(count):
            student = Student()
            student.first_name = f"خروجی{index}"
            student.last_name = "آزمون"
            student.national_code = f"{3000000000 + index}"
            student.birth_date = "1395/01/01"
            student.guardian_phone = "09121234567"
            student.guardian_name = "ولیِ آزمون"
            student.address = "نشانی آزمایشی با نیم‌فاصله‌ و ارقام ۱۲۳"
            student.is_active = 1
            students.append(StudentDAL().create(student))
        return students

    @staticmethod
    def _read_rows(file_path):
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        data_rows = []
        for row in sheet.iter_rows(values_only=True):
            if row and isinstance(row[0], int):
                data_rows.append(row)
        return sheet, data_rows

    def test_export_covers_all_given_students_in_a_real_xlsx(self):
        students = self._add_students(45)
        path = os.path.join(self.tmp, "export.xlsx")
        success, message = self.importer.export_students_to_excel(
            students, path, self.year)
        self.assertTrue(success, message)

        # فایل واقعاً xlsx است (zip با ساختار اکسل)، نه CSV با پسوند
        self.assertTrue(zipfile.is_zipfile(path))
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
        self.assertIn("xl/workbook.xml", names)
        self.assertIn("[Content_Types].xml", names)

        sheet, data_rows = self._read_rows(path)
        self.assertEqual(len(data_rows), 45)
        self.assertEqual(data_rows[0][1], "خروجی0")
        self.assertEqual(data_rows[-1][1], "خروجی44")
        self.assertEqual(data_rows[3][8], "نشانی آزمایشی با نیم‌فاصله‌ و ارقام ۱۲۳")
        self.assertEqual(sheet.cell(row=3, column=2).value, "45")
        self.assertNotIn("حذف‌شده", str(sheet.cell(row=1, column=1).value))

    def test_deleted_view_export_is_labeled(self):
        students = self._add_students(2)
        path = os.path.join(self.tmp, "deleted.xlsx")
        success, _message = self.importer.export_students_to_excel(
            students, path, self.year, deleted_view=True)
        self.assertTrue(success)
        sheet, data_rows = self._read_rows(path)
        self.assertEqual(len(data_rows), 2)
        title = str(sheet.cell(row=1, column=1).value)
        self.assertIn("حذف‌شده", title)
        status_values = [str(cell.value) for row in sheet.iter_rows()
                         for cell in row if cell.value]
        self.assertTrue(any("حذف‌شده (این ردیف‌ها" in value for value in status_values))


if __name__ == "__main__":
    unittest.main(verbosity=2)
