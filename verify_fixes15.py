"""
بررسی‌های دور پانزدهم — داوری دوم مدیر پروژه (مرحله‌های ۱ و ۲)

مرحلهٔ ۱ — سه ایراد محتوایی:
  A) روند رشد: بازه‌های میانی در همهٔ مسیرها (سرویس چندساله + گزارش سالانه) ... ۹ بررسی
  B) روایت چندساله: مسیر هر زمینه و «تغییر پس از مداخله» .................. ۱۰ بررسی
  C) گزارش والدین: یک منبع واحد، بدون اطلاعات داخلی ......................... ۹ بررسی
  D) نگهبان‌های بدون پس‌روندگی ................................................ ۴ بررسی
مرحلهٔ ۲ — باقی‌ماندهٔ فنی در کد فعلی:
  E) تریگرهای Audit جدول notifications: جایگزینی روی دیتابیس‌های موجود بدون حذف داده ... ۶ بررسی
  F) توضیحات/داک‌استرینگ‌های کهنه که کد خطرناک قدیمی را نقل می‌کردند ........ ۳ بررسی
  G) ثبت خطای بارگذاری Migration در لاگ برنامه (نه فقط print) ................ ۲ بررسی
  H) دلیل واقعی check_same_thread=False: فقط بستن بین‌نخی در close_all ........ ۱ بررسی

جمع: ۴۴ بررسی

هر بررسی روی یک دیتابیس موقت اجرا می‌شود و به داده‌های کاربر دست نمی‌زند.
"""

import contextlib
import io
import os
import re
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

PASS = 0
FAIL = 0
FAILURES = []


def check(section, name, condition, detail=""):
    global PASS, FAIL
    if condition:
        PASS += 1
        print(f"  ✅ [{section}] {name}")
    else:
        FAIL += 1
        FAILURES.append(f"[{section}] {name} {detail}")
        print(f"  ❌ [{section}] {name}  {detail}")


def read(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


# ============================================================
# آماده‌سازی دیتابیس موقت
# ============================================================
TMP = tempfile.mkdtemp(prefix="round15_verify_")
TEST_DB = os.path.join(TMP, "partow.db")

import config.settings as settings  # noqa: E402
import database.connection as dbc  # noqa: E402

settings.DB_PATH = TEST_DB
dbc.DB_PATH = TEST_DB
dbc.DatabaseConnection._instance = None
dbc.DatabaseConnection._connection = None
dbc.DatabaseConnection._initialized = False
dbc.DatabaseConnection._current_user_id = None

with contextlib.redirect_stdout(io.StringIO()):
    db = dbc.DatabaseConnection()
    conn = db.get_connection(user_id=1)

from dal.competency_dal import CompetencyDAL  # noqa: E402
from dal.followup_dal import FollowUpDAL  # noqa: E402
from dal.intervention_dal import InterventionDAL  # noqa: E402
from dal.observation_dal import ObservationDAL  # noqa: E402
from dal.student_academic_profile_dal import StudentAcademicProfileDAL  # noqa: E402
from dal.student_dal import StudentDAL  # noqa: E402
from models.followup import FollowUp  # noqa: E402
from models.intervention import Intervention  # noqa: E402
from models.observation import Observation  # noqa: E402
from models.student import Student  # noqa: E402
from models.student_academic_profile import StudentAcademicProfile  # noqa: E402
from services.parent_report_service import (  # noqa: E402
    PARENT_EXCLUDED_LAYERS,
    ParentReportService,
)
from services.report_generator import ReportGenerator  # noqa: E402
from services.trend_analysis_service import TrendAnalysisService  # noqa: E402
from utils.behavior_analysis import (  # noqa: E402
    DOMINANCE_RATIO,
    MEANINGFUL_SHARE_CHANGE,
    MIN_PATTERN_COUNT,
    classify_pattern,
    growth_direction,
)

report_gen = ReportGenerator()
trend_service = TrendAnalysisService()
parent_service = ParentReportService(report_generator=report_gen)


def _period(label, positive, negative, neutral=0):
    return {'label': label, 'positive': positive, 'negative': negative,
            'neutral': neutral, 'total': positive + negative + neutral}


# ------------------------------------------------------------
# دادهٔ آزمایشی: سناریوی دقیق مدیر پروژه
#   سال ۱: مشکل در زمینهٔ X (۳ منفی) + مداخله با پیگیری «بهبود»
#   سال ۲: X کمتر شده (۱ منفی، ۲ مثبت) + مداخلهٔ دوم (بدون پیگیری)
#   سال ۳: X برطرف شده (فقط مثبت)
#   زمینهٔ Y: نیاز ادامه‌دار با مداخله در سال ۲ (سال ۳ هنوز نیاز)
#   زمینهٔ Z: نیاز در سال ۱ با مداخله؛ در سال ۲ هیچ مشاهده‌ای برای Z نیست
# ------------------------------------------------------------
with contextlib.redirect_stdout(io.StringIO()):
    st = Student()
    st.first_name = "سناریو"
    st.last_name = "پانزدهم"
    st.national_code = "1515151515"
    st.is_active = 1
    student_id = StudentDAL().create(st).id

    for title in ('1403-1404', '1404-1405'):
        conn.execute(
            "INSERT INTO academic_years (title, is_active, start_date) VALUES (?, 0, ?)",
            (title, title[:4] + '/07/01'))
    conn.commit()
    year_ids = [r[0] for r in conn.execute(
        "SELECT id FROM academic_years WHERE is_deleted = 0 ORDER BY start_date, id"
    ).fetchall()][:3]

    comps = CompetencyDAL().get_all()
    c_x, c_s, c_y, c_z = [c.id for c in comps[:4]]
    titles = {c.id: c.title for c in comps[:4]}

    profile_ids = []
    for yid, grade in zip(year_ids, (1, 2, 3)):
        p = StudentAcademicProfile()
        p.student_id = student_id
        p.academic_year_id = yid
        p.grade = grade
        p.class_name = "الف"
        profile_ids.append(StudentAcademicProfileDAL().create(p).id)

    obs_dal, int_dal, fu_dal = ObservationDAL(), InterventionDAL(), FollowUpDAL()

    def _add_obs(pid, comp_id, btype, date, text="شرح آزمایشی"):
        o = Observation()
        o.student_profile_id = pid
        o.staff_id = 1
        o.competency_id = comp_id
        o.observation_date = date
        o.description = text
        o.behavior = "رفتار ثبت‌شدهٔ آزمایشی"
        o.behavior_type = btype
        o.severity = 2
        return obs_dal.create(o).id

    def _add_int(pid, obs_id, itype, date, result=None, status="done"):
        i = Intervention()
        i.student_profile_id = pid
        i.staff_id = 1
        i.observation_id = obs_id
        i.type = itype
        i.date = date
        i.description = "شرح داخلی مداخله SECRET_INTERVENTION_TEXT"
        i.goal = "هدف آزمایشی"
        i.status = "done"
        iid = int_dal.create(i).id
        if result:
            f = FollowUp()
            f.intervention_id = iid
            f.staff_id = 1
            f.date = date
            f.method = "گفتگو"
            f.status = status
            f.result_type = result
            f.result_description = "نتیجهٔ آزمایشی"
            f.description = "پیگیری آزمایشی"
            fu_dal.create(f)
        return iid

    # سال ۱
    o1 = _add_obs(profile_ids[0], c_x, "منفی", "1403/08/01", "SECRET_OBS_TEXT_1")
    _add_obs(profile_ids[0], c_x, "منفی", "1403/09/01")
    _add_obs(profile_ids[0], c_x, "منفی", "1403/10/01")
    _add_obs(profile_ids[0], c_s, "مثبت", "1403/08/02")
    _add_obs(profile_ids[0], c_s, "مثبت", "1403/09/02")
    o1z = _add_obs(profile_ids[0], c_z, "منفی", "1403/08/03")
    _add_obs(profile_ids[0], c_z, "منفی", "1403/09/03")
    _add_int(profile_ids[0], o1, "individual_talk", "1403/09/05", "improved")
    _add_int(profile_ids[0], o1z, "warning", "1403/09/06", "no_change")
    # سال ۲
    o2 = _add_obs(profile_ids[1], c_x, "منفی", "1404/08/01")
    _add_obs(profile_ids[1], c_x, "مثبت", "1404/09/01")
    _add_obs(profile_ids[1], c_x, "مثبت", "1404/10/01")
    _add_obs(profile_ids[1], c_s, "مثبت", "1404/08/02")
    _add_obs(profile_ids[1], c_s, "مثبت", "1404/09/02")
    o2y = _add_obs(profile_ids[1], c_y, "منفی", "1404/08/04")
    _add_obs(profile_ids[1], c_y, "منفی", "1404/09/04")
    _add_int(profile_ids[1], o2, "parent_call", "1404/09/05")
    _add_int(profile_ids[1], o2y, "individual_talk", "1404/09/07", "needs_more", status="pending")
    # سال ۳
    _add_obs(profile_ids[2], c_x, "مثبت", "1405/08/01")
    _add_obs(profile_ids[2], c_x, "مثبت", "1405/09/01")
    _add_obs(profile_ids[2], c_s, "مثبت", "1405/08/02")
    _add_obs(profile_ids[2], c_s, "مثبت", "1405/09/02")
    o3y = _add_obs(profile_ids[2], c_y, "منفی", "1405/08/04")
    _add_obs(profile_ids[2], c_y, "منفی", "1405/09/04")
    _add_int(profile_ids[2], o3y, "responsibility", "1405/09/05", "improved")

    # لایه‌های داخلی/حساس روی پروندهٔ سال ۱ (نباید وارد گزارش والدین شوند)
    conn.execute(
        "INSERT INTO family_contexts (student_profile_id, economic_status, family_stress, notes) "
        "VALUES (?, 'SECRET_ECONOMIC', 'SECRET_STRESS', 'SECRET_FAMILY_NOTE')",
        (profile_ids[0],))
    conn.execute(
        "INSERT INTO parent_interviews (student_profile_id, staff_id, interview_date, parent_name, "
        "topic, summary) VALUES (?, 1, '1403/09/10', 'ولی', 'SECRET_INTERVIEW_TOPIC', "
        "'SECRET_INTERVIEW_SUMMARY')", (profile_ids[0],))
    conn.execute(
        "INSERT INTO professional_interpretations (student_profile_id, staff_id, level, title, "
        "detailed_text) VALUES (?, 1, 'counselor', 'SECRET_INTERPRETATION_TITLE', "
        "'SECRET_INTERPRETATION_TEXT')", (profile_ids[0],))
    conn.commit()

# ============================================================
print("=" * 76)
print("بخش A: روند رشد — بازه‌های میانی در همهٔ مسیرها")
print("=" * 76)

pm_example = growth_direction([_period('مهر', 20, 70), _period('آبان', 60, 30),
                               _period('آذر', 25, 65)])
check("A", "مثال مدیر پروژه (۲۰/۷۰ ← ۶۰/۳۰ ← ۲۵/۶۵) → «تغییر ترکیبی / روند غیرقطعی» با نقطهٔ برگشت",
      pm_example['status'] == 'mixed' and 'غیرقطعی' in pm_example['label']
      and pm_example['turning_points'] == ['آبان'],
      f"{pm_example['status']} {pm_example['turning_points']}")

mono = growth_direction([_period('۱', 20, 70), _period('۲', 40, 50), _period('۳', 60, 30)])
spike = growth_direction([_period('۱', 50, 50), _period('۲', 90, 10), _period('۳', 50, 50)])
check("A", "تغییر یک‌جهت → همان جهت؛ ابتدا=انتها با جهش میانی → غیرقطعی (نه «ثابت»)",
      mono['status'] == 'improving' and spike['status'] == 'mixed'
      and spike['overall_status'] == 'stable',
      f"{mono['status']}/{spike['status']}")

fake_years = [
    {'year': 'y1', 'positive': 20, 'negative': 70, 'neutral': 0,
     'observations_count': 90, 'positive_share': 22.2, 'has_data': True},
    {'year': 'y2', 'positive': 60, 'negative': 30, 'neutral': 0,
     'observations_count': 90, 'positive_share': 66.7, 'has_data': True},
    {'year': 'y3', 'positive': 25, 'negative': 65, 'neutral': 0,
     'observations_count': 90, 'positive_share': 27.8, 'has_data': True},
]
multi = trend_service._analyze_multi_year_overall([90, 90, 90], fake_years)
check("A", "روند چندسالهٔ سرویس دیگر فقط سال اول/آخر را مقایسه نمی‌کند: مثال مدیر → غیرقطعی + نقطهٔ برگشت",
      multi['status'] == 'mixed' and 'غیرقطعی' in multi['label']
      and multi['turning_points'] == ['y2'] and multi.get('path_text'),
      f"{multi['status']} {multi.get('turning_points')}")

multi_mono = trend_service._analyze_multi_year_overall([9, 9, 9], [
    {'year': 'y1', 'positive': 2, 'negative': 7, 'neutral': 0, 'observations_count': 9,
     'positive_share': 22.2, 'has_data': True},
    {'year': 'y2', 'positive': 5, 'negative': 4, 'neutral': 0, 'observations_count': 9,
     'positive_share': 55.6, 'has_data': True},
    {'year': 'y3', 'positive': 8, 'negative': 1, 'neutral': 0, 'observations_count': 9,
     'positive_share': 88.9, 'has_data': True},
])
multi_one = trend_service._analyze_multi_year_overall([9], [fake_years[0]])
check("A", "چندساله: بهبود پیوسته → improving؛ یک سال → insufficient (کلیدهای سازگار حفظ شده)",
      multi_mono['status'] == 'improving' and multi_one['status'] == 'insufficient'
      and all(k in multi for k in ('status', 'message', 'icon', 'color', 'years_with_data',
                                   'total_years', 'first_year', 'last_year',
                                   'positive_change', 'volume_note', 'direction')),
      f"{multi_mono['status']}/{multi_one['status']}")

ts_src = read('services/trend_analysis_service.py')
multi_body = ts_src.split("def _analyze_multi_year_overall")[1].split("def _get_multi_year_top_competencies")[0]
check("A", "کد چندساله از growth_direction مشترک استفاده می‌کند و آستانه‌های جداگانهٔ اول/آخر حذف شده‌اند",
      'growth_direction(' in multi_body and 'slightly_improving' not in ts_src
      and 'positive_change > 0.15' not in multi_body,
      "")

real_multi = trend_service.analyze_multi_year_trend(student_id)
overall = real_multi.get('overall_trend') or {}
expected_dir = growth_direction([
    {'label': y['year'], 'positive': y['positive'], 'negative': y['negative'],
     'neutral': y['neutral'], 'total': y['observations_count']}
    for y in real_multi.get('year_data', [])
])
check("A", "روی دادهٔ واقعی سه‌ساله، نتیجهٔ سرویس همان مسیر همهٔ سال‌ها است (با path_text)",
      real_multi.get('success') and overall.get('status') == expected_dir['status']
      and overall.get('path_text') == expected_dir['path_text'] and overall['path_text'],
      f"{overall.get('status')} vs {expected_dir['status']}")

# --- ایراد واقعیِ کشف‌شده: جهت تغییر گزارش سالانه همیشه «دادهٔ کافی نیست» بود
year2_report = report_gen.generate_student_report(profile_ids[1])
trend_data = year2_report['trend_data']
annual_dir = report_gen.calculate_trend_direction(trend_data)
check("A", "جهت تغییر گزارش سالانه (برنامه و PDF) با چند ماه داده دیگر «دادهٔ کافی نیست» نمی‌گوید",
      trend_data and len(trend_data) >= 2 and annual_dir['status'] == 'improving'
      and annual_dir['periods_with_data'] == len(trend_data),
      f"{annual_dir['status']} periods={annual_dir.get('periods_with_data')}")

check("A", "دادهٔ ماهانه هر دو کلید count و total را دارد و calculate_trend_direction با هر دو کار می‌کند",
      all('total' in p and p['total'] == p['count'] for p in trend_data)
      and report_gen.calculate_trend_direction(
          [{k: v for k, v in p.items() if k != 'total'} for p in trend_data]
      )['status'] == 'improving',
      "")

check("A", "اصل «تعداد مشاهدات شاخص رشد نیست» در پیام‌های روند حفظ شده است",
      'نه تعداد مشاهدات' in annual_dir['message']
      and 'حجم ثبت' in annual_dir['volume_note']
      and 'تعداد مشاهدات' in (multi.get('volume_note') or ''),
      "")

# ============================================================
print()
print("=" * 76)
print("بخش B: روایت چندساله — مسیر هر زمینه و تغییر پس از مداخله")
print("=" * 76)

narrative = report_gen.generate_growth_narrative(student_id)
synth = narrative.get('synthesis') or {}
paths = {a['competency']: a for a in synth.get('area_paths', [])}
x_path = paths.get(titles[c_x], {})
y_path = paths.get(titles[c_y], {})
z_path = paths.get(titles[c_z], {})

check("B", "برای هر زمینهٔ نیازمند توجه/دارای مداخله، مسیر سال‌به‌سال ساخته می‌شود (X، Y، Z)",
      set(paths) == {titles[c_x], titles[c_y], titles[c_z]}
      and [p['status'] for p in x_path.get('path', [])] == ['need', 'strength', 'strength'],
      str({k: [p['status'] for p in v['path']] for k, v in paths.items()}))

x_post = {p['year']: p['status'] for p in x_path.get('post_intervention', [])}
check("B", "سناریوی مدیر: پس از مداخلهٔ سال ۱ → «کمتر شده»؛ پس از مداخلهٔ سال ۲ → «برطرف شده» (با احتیاط)",
      x_post == {'1403-1404': 'reduced', '1404-1405': 'resolved'}
      and 'کمتر شده' in x_path.get('text', '') and 'برطرف شده' in x_path.get('text', '')
      and 'مشاهدهٔ بیشتر' in x_path.get('text', ''),
      str(x_post))

check("B", "تغییر پس از مداخله با شمارش رفتارها مستند می‌شود (۳ منفی از ۳ ← ۱ منفی از ۳)",
      '3 منفی از 3 ← 1 منفی از 3' in x_path.get('text', '')
      and 'گفتگوی فردی (پیگیری: بهبود مشاهده شد)' in x_path.get('text', '')
      and 'تماس با والدین (بدون پیگیری)' in x_path.get('text', ''),
      x_path.get('text', '')[:160])

y_post = {p['year']: p['status'] for p in y_path.get('post_intervention', [])}
check("B", "نیاز ادامه‌دار پس از مداخله → «ادامه یافته»؛ مداخلهٔ آخرین سال → «سال بعدی هنوز ثبت نشده»",
      y_post == {'1404-1405': 'continued', '1405-1406': 'pending'}
      and 'ادامه یافته' in y_path.get('text', ''),
      str(y_post))

z_post = {p['year']: p['status'] for p in z_path.get('post_intervention', [])}
check("B", "نبودِ مشاهده در سال بعد «برطرف شدن» تلقی نمی‌شود (no_data، نه resolved)",
      z_post.get('1403-1404') == 'no_data'
      and z_path['path'][1]['status'] == 'not_observed'
      and 'قابل جمع‌بندی نیست' in z_path.get('text', ''),
      str(z_post))

changes_section = next((s for s in synth.get('sections', []) if s['key'] == 'changes'), {})
conclusion = next((s for s in synth.get('sections', []) if s['key'] == 'conclusion'), {})
check("B", "مسیر زمینه‌ها در بخش ۳ روایت و جمع‌بندی «تغییر پس از مداخله» در بخش ۵ می‌آید",
      any(ln == x_path.get('text') for ln in changes_section.get('lines', []))
      and 'تغییر پس از مداخله' in changes_section.get('title', '')
      and any('تغییر پس از مداخله' in ln and 'کمتر شده: 1 مورد' in ln
              and 'ادامه‌یافته: 1 مورد' in ln for ln in conclusion.get('lines', []))
      and synth.get('post_intervention_counts', {}).get('resolved') == 1,
      str(synth.get('post_intervention_counts')))

lines = report_gen.growth_narrative_lines(narrative)
headings = [t for k, t in lines if k == 'heading']
check("B", "ساختار روایت حفظ شده (همان ۶ بخش + جدول سال‌ها؛ بدون بخش جدید) و مسیر زمینه در رندر مشترک است",
      [s['key'] for s in synth.get('sections', [])]
      == ['trajectory', 'persistent', 'changes', 'interventions', 'conclusion', 'caveats']
      and len(headings) == 7
      and any(t == x_path.get('text') for k, t in lines if k == 'bullet')
      and 'تغییر پس از مداخله' in narrative.get('narrative', ''),
      str(len(headings)))

xlsx_path = os.path.join(TMP, "r15.xlsx")
pdf_path = os.path.join(TMP, "r15.pdf")
ok_xlsx, _ = report_gen.export_to_excel(profile_ids[2], xlsx_path)
ok_pdf, pdf_msg = report_gen.export_to_pdf(profile_ids[2], pdf_path)
sheet_ok = False
if ok_xlsx:
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path)
    if 'مسیر رشد چندساله' in wb.sheetnames:
        cells = [str(c.value) for row in wb['مسیر رشد چندساله'].iter_rows() for c in row if c.value]
        sheet_ok = any('تغییر پس از مداخله' in v for v in cells)
check("B", "PDF و Excel همان روایت (با تغییر پس از مداخله) را بدون خطا می‌سازند",
      ok_xlsx and sheet_ok and ok_pdf and os.path.getsize(pdf_path) > 1000,
      f"xlsx={ok_xlsx} sheet={sheet_ok} pdf={ok_pdf} {pdf_msg[:40]}")

all_text = narrative.get('narrative', '') + " " + " ".join(t for _, t in lines)
check("B", "مقایسه فقط دانش‌آموز با خودش است و شرح داخلی مداخله/مشاهده وارد روایت نمی‌شود",
      'میانگین کلاس' not in all_text and 'رتبه' not in all_text
      and 'SECRET_' not in all_text
      and 'دانش‌آموزان دیگر' in narrative.get('narrative', ''),
      "")

with contextlib.redirect_stdout(io.StringIO()):
    st1 = Student()
    st1.first_name = "تک"
    st1.last_name = "سال"
    st1.national_code = "1515151516"
    st1.is_active = 1
    single_id = StudentDAL().create(st1).id
    p1 = StudentAcademicProfile()
    p1.student_id = single_id
    p1.academic_year_id = year_ids[-1]
    p1.grade = 1
    p1.class_name = "ب"
    single_pid = StudentAcademicProfileDAL().create(p1).id
    so = _add_obs(single_pid, c_x, "منفی", "1405/08/01")
    _add_obs(single_pid, c_x, "منفی", "1405/09/01")
    _add_int(single_pid, so, "warning", "1405/09/05", "improved")
single = report_gen.generate_growth_narrative(single_id)
single_paths = single['synthesis'].get('area_paths', [])
check("B", "دانش‌آموز تک‌سال: مسیر با «سال بعدی هنوز ثبت نشده» و بدون ادعای بهبود/برطرف‌شدن",
      len(single_paths) == 1
      and [p['status'] for p in single_paths[0]['post_intervention']] == ['pending']
      and 'برطرف' not in single_paths[0]['text'] and 'کمتر شده' not in single_paths[0]['text'],
      str([p['status'] for p in single_paths[0]['post_intervention']]) if single_paths else "no path")

# ============================================================
print()
print("=" * 76)
print("بخش C: گزارش والدین — یک منبع واحد، بدون اطلاعات داخلی")
print("=" * 76)

prs_src = read('services/parent_report_service.py')
rg_src = read('services/report_generator.py')
page_src = read('views/pages/reports_page.py')
parent_method = rg_src.split("def generate_parent_report(")[1].split("\n    def ")[0]
check("C", "منبع واحد: صفحهٔ گزارش‌ها و ReportGenerator.generate_parent_report هر دو به ParentReportService می‌رسند",
      'parent_report_service.generate_parent_report_data(' in page_src
      and 'ParentReportService(' in parent_method
      and "full_report['strengths']" not in parent_method
      and 'generate_student_report(' in prs_src,
      "")

check("C", "گزارش والدین منطق جداگانهٔ الگو/روند ندارد (summarize_by_competency و گروه‌بندی ماهانهٔ مستقل حذف شد)",
      'summarize_by_competency' not in prs_src
      and 'observation_date[:7]' not in prs_src
      and 'calculate_trend_direction(' in prs_src,
      "")

parent = parent_service.generate_parent_report_data(profile_ids[0])
full = report_gen.generate_student_report(profile_ids[0])
check("C", "توانمندی‌ها و زمینه‌های نیازمند توجه والدین دقیقاً همان الگوهای گزارش داخلی‌اند (یک منبع)",
      parent is not None
      and [(e['competency'], e['positive'], e['count']) for e in parent['strengths']]
      == [(e['competency'], e['positive'], e['count']) for e in full['strengths'][:5]]
      and [(e['competency'], e['negative'], e['count']) for e in parent['weaknesses']]
      == [(e['competency'], e['negative'], e['count']) for e in full['weaknesses'][:5]]
      and len(parent['weaknesses']) == 2,
      str([e['competency'] for e in (parent or {}).get('weaknesses', [])]))

legacy = report_gen.generate_parent_report(profile_ids[0])
check("C", "مسیر قدیمی (generate_parent_report) همان دادهٔ واحد را برمی‌گرداند (بدون نسخهٔ دوم)",
      legacy is not None and sorted(legacy.keys()) == sorted(parent.keys())
      and legacy['summary'] == parent['summary']
      and legacy['recommendations'] == parent['recommendations'],
      "")


def _flatten(value, acc):
    if isinstance(value, dict):
        for k, v in value.items():
            acc.append(str(k))
            _flatten(v, acc)
    elif isinstance(value, (list, tuple)):
        for v in value:
            _flatten(v, acc)
    else:
        acc.append(str(value))


blob_parts = []
_flatten({k: v for k, v in parent.items() if k not in ('student', 'profile', 'academic_year')},
         blob_parts)
blob = " ".join(blob_parts)
check("C", "اطلاعات داخلی وارد گزارش والدین نمی‌شود (متن مشاهده، شناسهٔ مشاهدات، زمینهٔ خانوادگی، مصاحبه، تفسیر، شرح مداخله)",
      'SECRET_' not in blob and 'observation_ids' not in blob and 'examples' not in blob
      and 'شناسهٔ مشاهدات' not in blob and 'economic' not in blob
      and parent['has_family_context'] is True and parent['has_interviews'] is True
      and set(PARENT_EXCLUDED_LAYERS) <= set(parent['excluded_layers'])
      and 'تشخیص' in parent['privacy_note'],
      "")

internal_parent_recs = full['recommendations']['parents']
stripped = [parent_service._strip_trace(t) for t in internal_parent_recs]
check("C", "پیشنهادهای والدین از همان پیشنهادهای گزارش داخلی‌اند، فقط بدون شناسهٔ ردیابی؛ پیگیری در انتظار هم یادآوری می‌شود",
      all(s in parent['recommendations']['parents'] for s in stripped)
      and any('شناسهٔ مشاهدات' in t for t in internal_parent_recs)
      and all('شناسهٔ مشاهدات' not in t for t in parent['recommendations']['parents']),
      str(parent['recommendations']['parents'])[:120])

parent2 = parent_service.generate_parent_report_data(profile_ids[1])
check("C", "روند گزارش والدین همان جهت تغییر گزارش داخلی است (نه محاسبهٔ جداگانه) و پیگیری در انتظار شمرده می‌شود",
      parent2['trend']['has_data'] is True
      and parent2['trend']['direction']['status'] == annual_dir['status'] == 'improving'
      and parent2['trend_direction']['status'] == 'improving'
      and parent2['pending_count'] == 1
      and any('پیگیری' in r and 'همکاری' in r for r in parent2['recommendations']['parents']),
      f"{parent2['trend'].get('trend_text')} pending={parent2['pending_count']}")

ui_keys = ('student_name', 'grade', 'class', 'summary', 'strengths', 'weaknesses',
           'recommendations', 'trend', 'observations_count', 'interventions_count',
           'pending_count', 'privacy_note', 'strength_lines', 'weakness_lines',
           'intervention_effectiveness', 'trend_direction', 'balance_note')
check("C", "کلیدهای موردنیاز صفحهٔ گزارش‌ها و PDF در دادهٔ واحد موجودند و بدون شرح/هدف مداخله",
      all(k in parent for k in ui_keys)
      and all({'competency', 'positive', 'count'} <= set(e) for e in parent['strengths'])
      and 'items' not in parent['intervention_effectiveness']
      and parent['intervention_effectiveness']['outcome_counts'].get('improved') == 1,
      str([k for k in ui_keys if k not in parent]))

parent_pdf = os.path.join(TMP, "parent15.pdf")
ok_ppdf, ppdf_msg = parent_service.export_parent_report_pdf(profile_ids[0], parent_pdf)
check("C", "خروجی PDF والدین از همان دادهٔ واحد ساخته می‌شود",
      ok_ppdf and os.path.getsize(parent_pdf) > 1000,
      f"{ok_ppdf} {ppdf_msg[:50]}")

# ============================================================
print()
print("=" * 76)
print("بخش D: نگهبان‌های بدون پس‌روندگی")
print("=" * 76)

check("D", "منطق قوت/ضعف (utils/behavior_analysis) دست‌نخورده است: ۲ تکرار، غلبهٔ ۶۰٪، تغییر ۱۰ واحد",
      MIN_PATTERN_COUNT == 2 and abs(DOMINANCE_RATIO - 0.6) < 1e-9
      and abs(MEANINGFUL_SHARE_CHANGE - 10.0) < 1e-9
      and classify_pattern(2, 0, 0, 2) == 'strength'
      and classify_pattern(0, 2, 0, 2) == 'needs_attention'
      and classify_pattern(0, 0, 2, 2) == 'insufficient'
      and classify_pattern(1, 0, 0, 1) not in ('strength', 'needs_attention'))

check("D", "قراردادهای دورهای ۱۱–۱۳ روایت حفظ شده‌اند (کلیدهای synthesis و اولویت‌ها)",
      all(k in synth for k in ('text', 'sections', 'trajectory', 'persistent_strengths',
                               'persistent_needs', 'changed_areas', 'effective_interventions',
                               'area_interventions', 'priorities', 'caveats'))
      and set(p['kind'] for p in synth['priorities']) <= {'persistent_need', 'reversed',
                                                          'new_need', 'build_on'},
      "")

check("D", "بدون قابلیت جدید: نسخهٔ دیتابیس ۹ و نسخهٔ برنامه ۲۷.۲.۴ (بدون migration و جدول تازه)",
      settings.DB_VERSION == 9 and settings.APP_VERSION == '27.2.4'
      and not [r[0] for r in conn.execute(
          "SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'parent_report%'")],
      f"{settings.DB_VERSION}/{settings.APP_VERSION}")

check("D", "دکمهٔ گزارش والدین و مسیر اصلی صفحهٔ گزارش‌ها حفظ شده است",
      'self.parent_btn.clicked.connect(self.show_parent_report)' in page_src
      and 'def show_parent_report' in page_src
      and 'def export_parent_report_pdf' in prs_src,
      "")

# ============================================================
print()
print("=" * 76)
print("بخش E: تریگرهای Audit جدول notifications (IF NOT EXISTS → DROP/CREATE مدیریت‌شده)")
print("=" * 76)

import json  # noqa: E402
import threading  # noqa: E402

from dal.notification_dal import NotificationDAL  # noqa: E402
from models.notification import Notification  # noqa: E402
from utils.security import normalize_entity_type  # noqa: E402

notif_dal = NotificationDAL()


def _notification_triggers(connection):
    return {r[0]: (r[1] or '') for r in connection.execute(
        "SELECT name, sql FROM sqlite_master WHERE type = 'trigger' "
        "AND tbl_name = 'notifications'")}


def _make_notification(title, recipient=1):
    n = Notification()
    n.user_id = recipient
    n.type = 'reminder'
    n.title = title
    n.message = 'پیام آزمایشی اعلان'
    n.entity_type = 'followup'
    n.entity_id = 1
    return notif_dal.create(n).id


fresh_triggers = _notification_triggers(conn)
v6_src = read('database/migrations/migration_v6.py')
check("E", "notifications در فهرست مدیریت‌شده است: ۵ تریگر با JSON کامل ردیف؛ migration_v6 دیگر تریگر IF NOT EXISTS نمی‌سازد",
      'notifications' in dbc.DatabaseConnection._AUDIT_TABLES
      and set(fresh_triggers) == {f'trg_notifications_{k}_audit' for k in
                                  ('insert', 'update', 'soft_delete', 'restore', 'hard_delete')}
      and "'title', NEW.\"title\"" in fresh_triggers['trg_notifications_insert_audit']
      and "'type', NEW.type)" not in fresh_triggers['trg_notifications_insert_audit']
      and 'CREATE TRIGGER' not in v6_src,
      str(sorted(fresh_triggers)))

# شبیه‌سازی دیتابیس موجود با تریگرهای قدیمی migration_v6 (فقط id، IF NOT EXISTS)
for name in fresh_triggers:
    conn.execute(f'DROP TRIGGER IF EXISTS "{name}"')
conn.execute("""
    CREATE TRIGGER IF NOT EXISTS trg_notifications_insert_audit
    AFTER INSERT ON notifications
    BEGIN
        INSERT INTO audit_logs (user_id, action, entity_type, entity_id, new_value)
        VALUES (NEW.user_id, 'create', 'notification', NEW.id,
                json_object('id', NEW.id, 'type', NEW.type));
    END
""")
conn.execute("""
    CREATE TRIGGER IF NOT EXISTS trg_notifications_update_audit
    AFTER UPDATE ON notifications
    WHEN NEW.is_deleted = 0 AND OLD.is_deleted = 0
    BEGIN
        INSERT INTO audit_logs (user_id, action, entity_type, entity_id, old_value, new_value)
        VALUES (NEW.user_id, 'edit', 'notification', NEW.id,
                json_object('id', OLD.id), json_object('id', NEW.id));
    END
""")
conn.commit()
legacy_id = _make_notification('اعلان قدیمی')
audit_before = conn.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
legacy_row = conn.execute(
    "SELECT new_value FROM audit_logs WHERE entity_type = 'notification' AND entity_id = ?",
    (legacy_id,)).fetchone()

# «راه‌اندازی دوباره»: بستن همهٔ اتصال‌ها و بازکردن (نسل جدید → تضمین اسکیما/تریگرها)
with contextlib.redirect_stdout(io.StringIO()):
    db.close_all()
    conn = db.get_connection(user_id=1)
upgraded = _notification_triggers(conn)
audit_after = conn.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
legacy_still = conn.execute(
    "SELECT COUNT(*) FROM audit_logs WHERE entity_type = 'notification' AND entity_id = ?",
    (legacy_id,)).fetchone()[0]
check("E", "روی دیتابیس موجود، تریگرهای قدیمی با همان نام‌ها جایگزین می‌شوند (۵ تریگر، JSON کامل) و هیچ ردیفی از audit_logs حذف نمی‌شود",
      legacy_row is not None and json.loads(legacy_row[0]) == {'id': legacy_id, 'type': 'reminder'}
      and len(upgraded) == 5
      and "'title', NEW.\"title\"" in upgraded['trg_notifications_insert_audit']
      and "'type', NEW.type)" not in upgraded['trg_notifications_insert_audit']
      and audit_after == audit_before and legacy_still == 1,
      f"triggers={len(upgraded)} audit {audit_before}->{audit_after}")

managed_id = _make_notification('اعلان مدیریت‌شده')
create_rows = conn.execute(
    "SELECT user_id, new_value FROM audit_logs WHERE entity_type = 'notifications' "
    "AND action = 'create' AND entity_id = ?", (managed_id,)).fetchall()
create_json = json.loads(create_rows[0][1]) if create_rows else {}
check("E", "ایجاد اعلان: دقیقاً یک ردیف Audit با انجام‌دهندهٔ واقعی (نه گیرنده)، نوع موجودیت یکدست و تصویر کامل ردیف",
      len(create_rows) == 1 and create_rows[0][0] == 1
      and create_json.get('title') == 'اعلان مدیریت‌شده'
      and create_json.get('user_id') == 1 and 'message' in create_json,
      str(create_rows[:1]))

notif_dal.mark_as_read(managed_id)
edit_rows = conn.execute(
    "SELECT old_value, new_value FROM audit_logs WHERE entity_type = 'notifications' "
    "AND action = 'edit' AND entity_id = ?", (managed_id,)).fetchall()
check("E", "خوانده‌شدن اعلان: یک ردیف «ویرایش» با old/new کامل (is_read از ۰ به ۱)",
      len(edit_rows) == 1
      and json.loads(edit_rows[0][0]).get('is_read') == 0
      and json.loads(edit_rows[0][1]).get('is_read') == 1,
      str(len(edit_rows)))

scheduler_result = {}


def _scheduler_like_worker():
    with db.worker_context(None):
        scheduler_result['id'] = _make_notification('اعلان زمان‌بند')


worker = threading.Thread(target=_scheduler_like_worker)
worker.start()
worker.join()
sched_row = conn.execute(
    "SELECT user_id FROM audit_logs WHERE entity_type = 'notifications' "
    "AND action = 'create' AND entity_id = ?", (scheduler_result.get('id'),)).fetchone()
check("E", "اعلانِ ساختهٔ نخ زمان‌بند (worker_context بدون کاربر) با user_id تهی ثبت می‌شود، نه به نام گیرنده",
      sched_row is not None and sched_row[0] is None,
      str(sched_row))

conn.execute("UPDATE notifications SET created_at = '2000-01-01 00:00:00' WHERE id = ?",
             (managed_id,))
conn.commit()
notif_dal.delete_old(days=30)
soft_rows = conn.execute(
    "SELECT old_value FROM audit_logs WHERE entity_type = 'notifications' "
    "AND action = 'delete_soft' AND entity_id = ?", (managed_id,)).fetchall()
dash_src = read('views/pages/dashboard_page.py')
check("E", "پاکسازی اعلان‌های قدیمی (حذف منطقی) ردیف delete_soft با تصویر قبلی دارد؛ نام‌های قدیمی/جدید یکدست و برچسب داشبورد موجود",
      len(soft_rows) == 1 and json.loads(soft_rows[0][0]).get('title') == 'اعلان مدیریت‌شده'
      and normalize_entity_type('notification') == 'notifications'
      and normalize_entity_type('notifications') == 'notifications'
      and "'notifications': 'اعلان'" in dash_src,
      str(len(soft_rows)))

# ============================================================
print()
print("=" * 76)
print("بخش F: توضیحات کهنه‌ای که کد خطرناک قدیمی را نقل می‌کردند")
print("=" * 76)

conn_src = read('database/connection.py')
migrate_body = conn_src.split("def _migrate_database")[1].split("def _create_all_tables")[0]
check("F", "connection.py: نه متد مردهٔ _migrate_to_v1 و نه نقلِ «migrations = {1: …}»؛ _migrate_database فقط MigrationManager، خطا قبل از مُهر نسخه",
      '_migrate_to_v1' not in conn_src and 'migrations = {1' not in conn_src
      and 'هیچ مسیر جایگزین' in migrate_body
      and 'raise' in migrate_body
      and migrate_body.index("MigrationManager.migrate(") < migrate_body.index("self._set_db_version(to_version)"),
      "")

connect_block = conn_src.split("def _open_thread_connection")[1].split("conn.row_factory")[0]
check("F", "توضیح کنار check_same_thread=False دلیل واقعی (بستن بین‌نخی در close_all زیر قفل) را می‌گوید، نه «برای مسیرهای قدیمی»",
      'close_all()' in connect_block and 'DB_THREAD_LOCK' in connect_block
      and 'مجوز «استفادهٔ هم‌زمان» نیست' in connect_block
      and 'اگر مسیری قدیمی اتصال را جابه‌جا کرد' not in connect_block,
      "")

backup_src = read('utils/backup.py')
restore_body = backup_src.split("def restore_backup")[1]
check("F", "backup.py: نقلِ کد قدیمی (کپی زنده / استخراج بدون اعتبارسنجی) از داک‌استرینگ‌ها حذف شد؛ copy2 فقط در جایگزینی فایل هنگام بازیابی، پس از اعتبارسنجی و زیر قفل",
      backup_src.count('shutil.copy2(db_backup, self.db_path)') == 1
      and 'shutil.copy2(db_backup, self.db_path)' in restore_body
      and 'extractall' not in backup_src
      and '_safe_extract(' in restore_body
      and restore_body.index('_safe_extract(')
      < restore_body.index('_verify_sqlite_file(')
      < restore_body.index('_quiesce_database()')
      < restore_body.index('shutil.copy2(db_backup, self.db_path)'),
      "")

# ============================================================
print()
print("=" * 76)
print("بخش G: خطای بارگذاری Migration در لاگ برنامه")
print("=" * 76)

import logging  # noqa: E402
import shutil  # noqa: E402

import database.migrations.manager as migration_manager  # noqa: E402

manager_src = read('database/migrations/manager.py')
discover_src = manager_src.split("def _discover_migrations")[1].split("class MigrationManager")[0]
check("G", "_discover_migrations دیگر print ندارد و از logger برنامه استفاده می‌کند",
      'print(' not in discover_src and 'logger.error(' in discover_src
      and 'logger.warning(' in discover_src and 'from utils.logger import get_logger' in manager_src,
      "")


class _CaptureHandler(logging.Handler):
    def __init__(self):
        super().__init__(level=logging.DEBUG)
        self.records = []

    def emit(self, record):
        self.records.append(record)


capture = _CaptureHandler()
migration_manager.logger.addHandler(capture)
# (دور ۱۶) کشف migration از روی فایل‌سیستم است؛ یک پوشهٔ موقت با یک فایل
# سالم و یک فایل خراب ساخته می‌شود. قرارداد جدید: فایلِ موجودِ خراب →
# MigrationLoadError با traceback در لاگ (نه نادیده‌گرفتن خاموش).
_broken_dir = tempfile.mkdtemp(prefix="r15_broken_mig_")
with open(os.path.join(_broken_dir, "migration_v1.py"), "w", encoding="utf-8") as fh:
    fh.write("def upgrade(connection):\n    pass\n")
with open(os.path.join(_broken_dir, "migration_v51.py"), "w", encoding="utf-8") as fh:
    fh.write("import module_that_does_not_exist_for_parto_test\n\ndef upgrade(connection):\n    pass\n")
stdout_buf = io.StringIO()
load_error = None
try:
    with contextlib.redirect_stdout(stdout_buf):
        migration_manager._discover_migrations(directory=_broken_dir)
except migration_manager.MigrationLoadError as e:
    load_error = e
finally:
    migration_manager.logger.removeHandler(capture)
    shutil.rmtree(_broken_dir, ignore_errors=True)
logged = [r for r in capture.records if 'migration_v51' in r.getMessage()]
check("G", "ماژول migration خراب: در لاگ برنامه (ERROR با traceback) ثبت می‌شود، چیزی چاپ نمی‌شود و Migration با MigrationLoadError متوقف می‌شود",
      len(logged) == 1 and logged[0].levelno == logging.ERROR and logged[0].exc_info is not None
      and 'migration_v51' not in stdout_buf.getvalue()
      and load_error is not None and load_error.filename == 'migration_v51.py',
      f"logged={len(logged)} error={load_error}")

# ============================================================
print()
print("=" * 76)
print("بخش H: check_same_thread=False فقط برای بستن بین‌نخی")
print("=" * 76)

worker_state = {}
worker_ready = threading.Event()
worker_release = threading.Event()


def _holding_worker():
    worker_conn = db.get_connection()
    worker_state['conn'] = worker_conn
    worker_state['same_object_as_main'] = worker_conn is conn
    worker_ready.set()
    worker_release.wait(timeout=30)
    try:
        worker_conn.execute("SELECT 1")
        worker_state['usable_after_close_all'] = True
    except Exception as e:
        worker_state['usable_after_close_all'] = False
        worker_state['error'] = type(e).__name__


holder = threading.Thread(target=_holding_worker)
holder.start()
worker_ready.wait(timeout=30)
close_error = None
try:
    with contextlib.redirect_stdout(io.StringIO()):
        db.close_all()
except Exception as e:
    close_error = e
worker_release.set()
holder.join(timeout=30)
with contextlib.redirect_stdout(io.StringIO()):
    conn = db.get_connection(user_id=1)
check("H", "اتصال نخ کارگر جدا از نخ اصلی است؛ close_all از نخ اصلی آن را بدون خطا می‌بندد و نخ کارگر دیگر نمی‌تواند از آن استفاده کند",
      worker_state.get('same_object_as_main') is False and close_error is None
      and worker_state.get('usable_after_close_all') is False
      and worker_state.get('error') == 'ProgrammingError',
      str(worker_state))

# ============================================================
print()
print("=" * 76)
print(f"نتیجهٔ دور پانزدهم (مرحله‌های ۱ و ۲):  {PASS} موفق / {FAIL} ناموفق  از {PASS + FAIL}")
if FAILURES:
    print("موارد ناموفق:")
    for item in FAILURES:
        print("   -", item)
print("=" * 76)

with contextlib.suppress(Exception):
    db.close_all()
sys.exit(1 if FAIL else 0)
