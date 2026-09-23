"""
دیالوگ خروجی داده برای هوش مصنوعی - نسخه با پشتیبانی از راهنما
"""

import csv
import json
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from PySide6.QtWidgets import (
    QApplication,
    QButtonGroup,
    QDialog,
    QFileDialog,
    QFormLayout,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QMessageBox,
    QPushButton,
    QRadioButton,
    QScrollArea,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from dal.competency_dal import CompetencyDAL
from dal.followup_dal import FollowUpDAL
from dal.intervention_dal import InterventionDAL
from dal.observation_dal import ObservationDAL
from dal.student_academic_profile_dal import StudentAcademicProfileDAL
from dal.student_dal import StudentDAL
from utils.persian_pdf import PersianPDF
from utils.time_utils import utc_now, utc_now_iso


class ExportAIDialog(QDialog):
    """دیالوگ خروجی داده برای هوش مصنوعی با پشتیبانی از راهنما"""

    def __init__(self, student_id=None, profile_id=None, parent=None):
        super().__init__(parent)

        self.student_id = student_id
        self.profile_id = profile_id

        self.student_dal = StudentDAL()
        self.profile_dal = StudentAcademicProfileDAL()
        self.observation_dal = ObservationDAL()
        self.intervention_dal = InterventionDAL()
        self.followup_dal = FollowUpDAL()
        self.competency_dal = CompetencyDAL()

        self.setWindowTitle("📤 خروجی داده برای هوش مصنوعی")
        self.setModal(True)
        self.resize(700, 650)

        self.setup_ui()
        self.load_data()

    def setup_ui(self):
        main_layout = QVBoxLayout()
        self.setLayout(main_layout)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QScrollArea.Shape.NoFrame)
        scroll.setStyleSheet("""
            QScrollArea {
                border: none;
                background-color: transparent;
            }
            QScrollBar:vertical {
                background-color: #f0f0f0;
                width: 10px;
                border-radius: 5px;
            }
            QScrollBar::handle:vertical {
                background-color: #D9C36A;
                border-radius: 5px;
                min-height: 30px;
            }
            QScrollBar::handle:vertical:hover {
                background-color: #D9C36A;
            }
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {
                height: 0px;
            }
        """)

        container = QWidget()
        container.setStyleSheet("background-color: transparent;")
        container_layout = QVBoxLayout()
        container_layout.setSpacing(12)
        container_layout.setContentsMargins(5, 5, 5, 5)
        container.setLayout(container_layout)

        # ===== توضیحات =====
        desc_label = QLabel("""
📋 **خروجی داده برای تحلیل با هوش مصنوعی**

این ابزار داده‌های دانش‌آموز را در فرمت‌های مختلف خروجی می‌دهد تا بتوانید
با هوش مصنوعی مورد نظر خود (ChatGPT، Claude، Gemini، DeepSeek و ...) تحلیل کنید.

✅ فرمت مورد نظر خود را انتخاب کنید و روی دکمه خروجی کلیک کنید.
""")
        desc_label.setWordWrap(True)
        desc_label.setStyleSheet("""
            QLabel {
                background-color: #08223A;
                padding: 12px;
                border-radius: 8px;
                border: 1px solid #D9C36A;
                font-size: 13px;
                line-height: 1.6;
            }
        """)
        container_layout.addWidget(desc_label)

        # ===== اطلاعات دانش‌آموز =====
        info_group = QGroupBox("👤 اطلاعات دانش‌آموز")
        info_group.setStyleSheet("""
            QGroupBox {
    color: #111111;
    background-color: #66BB6A;
                font-weight: bold;
                border: 2px solid #8BC34A;
                border-radius: 8px;
                margin-top: 10px;
                padding-top: 10px;
            }
            QGroupBox::title {
    background-color: #8BC34A;
                subcontrol-origin: margin;
                left: 15px;
                padding: 0 8px 0 8px;
                color: #111111;
            }
        """)
        info_layout = QFormLayout()
        info_group.setLayout(info_layout)

        self.student_name_label = QLabel("نامشخص")
        self.student_name_label.setStyleSheet("font-weight: bold; color: #F4C542;")
        info_layout.addRow("👤 دانش‌آموز:", self.student_name_label)

        self.student_grade_label = QLabel("نامشخص")
        info_layout.addRow("📚 پایه:", self.student_grade_label)

        self.student_class_label = QLabel("نامشخص")
        info_layout.addRow("🏫 کلاس:", self.student_class_label)

        self.data_count_label = QLabel("در حال بارگذاری...")
        info_layout.addRow("📊 تعداد رکوردها:", self.data_count_label)

        container_layout.addWidget(info_group)

        # ===== انتخاب فرمت =====
        format_group = QGroupBox("📁 انتخاب فرمت خروجی")
        format_group.setStyleSheet("""
            QGroupBox {
    color: #111111;
    background-color: #66BB6A;
                font-weight: bold;
                border: 2px solid #8BC34A;
                border-radius: 8px;
                margin-top: 10px;
                padding-top: 10px;
            }
            QGroupBox::title {
    background-color: #8BC34A;
                subcontrol-origin: margin;
                left: 15px;
                padding: 0 8px 0 8px;
                color: #111111;
            }
        """)
        format_layout = QVBoxLayout()
        format_group.setLayout(format_layout)

        format_help = QLabel("💡 هر هوش مصنوعی فرمت خاصی را بهتر پشتیبانی می‌کند. فرمت مناسب را انتخاب کنید:")
        format_help.setStyleSheet("font-size: 12px; color: #D9C36A;")
        format_help.setWordWrap(True)
        format_layout.addWidget(format_help)

        self.format_group = QButtonGroup(self)

        formats = [
            ("pdf", "📄 PDF (همراه با پرامپت)", "مناسب برای ChatGPT، Claude، Gemini، DeepSeek"),
            ("excel", "📊 Excel (XLSX)", "مناسب برای تحلیل داده در Excel و Power BI"),
            ("csv", "📃 CSV", "سبک و قابل استفاده در همه ابزارها"),
            ("json", "📝 JSON", "ساختاریافته و دقیق - مناسب برای برنامه‌نویسان"),
            ("txt", "📋 TXT (ساده)", "همه‌جا قابل استفاده"),
            ("zip", "📦 ZIP (همه فرمت‌ها)", "یکجا دریافت کنید")
        ]

        self.format_radio_buttons = {}
        for i, (value, label, tooltip) in enumerate(formats):
            radio = QRadioButton(label)
            radio.setToolTip(tooltip)
            radio.setStyleSheet("""
                QRadioButton {
                    padding: 5px 10px;
                    font-size: 13px;
                }
                QRadioButton::indicator {
                    width: 18px;
                    height: 18px;
                }
            """)
            self.format_group.addButton(radio, i)
            self.format_radio_buttons[value] = radio
            if value == "pdf":
                radio.setChecked(True)
            format_layout.addWidget(radio)

            tip_label = QLabel(f"   ↳ {tooltip}")
            tip_label.setStyleSheet("font-size: 11px; color: #D9C36A; padding-left: 30px;")
            format_layout.addWidget(tip_label)

        container_layout.addWidget(format_group)

        # ===== پرامپت =====
        prompt_group = QGroupBox("📋 پرامپت پیشنهادی برای هوش مصنوعی")
        prompt_group.setStyleSheet("""
            QGroupBox {
    color: #111111;
    background-color: #66BB6A;
                font-weight: bold;
                border: 2px solid #8BC34A;
                border-radius: 8px;
                margin-top: 10px;
                padding-top: 10px;
            }
            QGroupBox::title {
    background-color: #8BC34A;
                subcontrol-origin: margin;
                left: 15px;
                padding: 0 8px 0 8px;
                color: #111111;
            }
        """)
        prompt_layout = QVBoxLayout()
        prompt_group.setLayout(prompt_layout)

        self.prompt_text = QTextEdit()
        self.prompt_text.setReadOnly(True)
        self.prompt_text.setMinimumHeight(150)
        self.prompt_text.setMaximumHeight(200)
        self.prompt_text.setStyleSheet("""
            QTextEdit {
    color: #F4C542;
                background-color: #08223A;
                border: 1px solid #8BC34A;
                border-radius: 5px;
                padding: 8px;
                font-size: 12px;
                line-height: 1.6;
            }
        """)
        self.prompt_text.setPlainText(self._get_standard_prompt())
        prompt_layout.addWidget(self.prompt_text)

        container_layout.addWidget(prompt_group)

        scroll.setWidget(container)
        main_layout.addWidget(scroll)

        # ===== دکمه‌ها =====
        button_layout = QHBoxLayout()
        button_layout.setSpacing(10)

        self.export_btn = QPushButton("📤 خروجی و دانلود")
        self.export_btn.setStyleSheet("""
            QPushButton {
                background-color: #66BB6A;
                color: #111111;
                padding: 12px 30px;
                border: none;
                border-radius: 6px;
                font-weight: bold;
                font-size: 14px;
                min-height: 40px;
            }
            QPushButton:hover { background-color: #8BC34A; }
        """)
        self.export_btn.clicked.connect(self.export_data)
        button_layout.addWidget(self.export_btn)

        self.copy_prompt_btn = QPushButton("📋 کپی پرامپت")
        self.copy_prompt_btn.setStyleSheet("""
            QPushButton {
                background-color: #0B2E4F;
                color: #F4C542;
                padding: 12px 20px;
                border: none;
                border-radius: 6px;
                font-weight: bold;
                font-size: 13px;
                min-height: 40px;
            }
            QPushButton:hover { background-color: #08223A; }
        """)
        self.copy_prompt_btn.clicked.connect(self.copy_prompt)
        button_layout.addWidget(self.copy_prompt_btn)

        button_layout.addStretch()

        self.cancel_btn = QPushButton("❌ انصراف")
        self.cancel_btn.setStyleSheet("""
            QPushButton {
                background-color: #C62828;
                color: #F4C542;
                padding: 12px 20px;
                border: none;
                border-radius: 6px;
                font-weight: bold;
                font-size: 13px;
                min-height: 40px;
            }
            QPushButton:hover { background-color: #D94B4B; }
        """)
        self.cancel_btn.clicked.connect(self.reject)
        button_layout.addWidget(self.cancel_btn)

        main_layout.addLayout(button_layout)
    
    # ===== بقیه متدها (بدون تغییر) =====
    
    def load_data(self):
        """بارگذاری اطلاعات دانش‌آموز"""
        try:
            if self.profile_id:
                profile = self.profile_dal.get_by_id(self.profile_id)
                if profile:
                    student = self.student_dal.get_by_id(profile.student_id)
                    if student:
                        self.student_name_label.setText(student.full_name)
                        self.student_grade_label.setText(profile.grade_display)
                        self.student_class_label.setText(profile.class_name or "-")
                        
                        observations = self.observation_dal.get_by_student_profile(profile.id)
                        interventions = self.intervention_dal.get_by_student_profile(profile.id)
                        followups = self.followup_dal.get_by_student_profile(profile.id)
                        
                        count_text = f"مشاهدات: {len(observations)} | مداخلات: {len(interventions)} | پیگیری‌ها: {len(followups)}"
                        self.data_count_label.setText(count_text)
                        
            elif self.student_id:
                profile = self.profile_dal.get_active_by_student(self.student_id)
                if profile:
                    self.profile_id = profile.id
                    self.load_data()
                else:
                    self.student_name_label.setText("دانش‌آموز بدون پرونده")
                    self.data_count_label.setText("هیچ داده‌ای یافت نشد")
                    
        except Exception as e:
            self.data_count_label.setText(f"خطا: {e!s}")
    
    def _get_standard_prompt(self):
        """دریافت پرامپت استاندارد برای هوش مصنوعی"""
        return """
📋 **دستورالعمل تحلیل داده‌های دانش‌آموز**

لطفاً داده‌های ارسالی را با دقت بررسی کنید و تحلیل زیر را ارائه دهید:

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

**۱. شناسایی الگوهای رفتاری:**
• چه الگوهای تکراری در مشاهدات وجود دارد؟
• آیا بین مشاهدات و مداخلات ارتباط مشخصی وجود دارد؟
• چه زمانی از روز یا چه موقعیت‌هایی بیشترین مشاهدات را داشته است؟

**۲. تحلیل شایستگی‌ها:**
• کدام شایستگی‌ها در وضعیت خوب و کدام نیازمند توجه هستند؟
• روند تغییرات شایستگی‌ها در طول زمان چگونه است؟
• آیا بین شایستگی‌ها ارتباطی وجود دارد؟

**۳. پیشنهاد مداخلات:**
• چه مداخلاتی برای هر دانش‌آموز مناسب‌تر است؟
• آیا الگوی خاصی در پاسخ به مداخلات وجود دارد؟
• چه مداخلاتی بیشترین تأثیر را داشته‌اند؟

**۴. هشدارها و نقاط بحرانی:**
• آیا دانش‌آموزی نیاز به توجه فوری دارد؟
• آیا الگوی خطرناکی مشاهده می‌شود؟
• چه عواملی می‌توانند باعث تشدید رفتارهای منفی شوند؟

**۵. پیشنهادات کلی:**
• چه تغییراتی در رویکرد آموزشی می‌تواند مفید باشد؟
• چه نقاط قوت و ضعفی در داده‌ها وجود دارد؟
• چه توصیه‌هایی برای معلم، والدین و مشاور دارید؟

**۶. تحلیل روند بلندمدت:**
• روند کلی رشد دانش‌آموز چگونه است؟
• آیا بهبود یا افت قابل توجهی مشاهده می‌شود؟

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

📌 **نکات مهم:**
• تحلیل را به صورت ساختاریافته و با زبان فارسی روان ارائه دهید.
• از آمار و ارقام برای پشتیبانی از تحلیل‌های خود استفاده کنید.
• پیشنهادات عملی و قابل اجرا ارائه دهید.
• در صورت مشاهده هرگونه الگوی غیرعادی، به آن اشاره کنید.
"""
    
    def copy_prompt(self):
        """کپی پرامپت در کلیپ‌بورد"""
        clipboard = QApplication.clipboard()
        clipboard.setText(self.prompt_text.toPlainText())
        QMessageBox.information(self, "موفقیت", "✅ پرامپت با موفقیت در کلیپ‌بورد کپی شد.")
    
    def export_data(self):
        """خروجی داده با فرمت انتخاب شده"""
        if not self.profile_id:
            QMessageBox.warning(self, "توجه", "هیچ داده‌ای برای خروجی وجود ندارد.")
            return
        
        checked_id = self.format_group.checkedId()
        if checked_id < 0:
            QMessageBox.warning(self, "توجه", "لطفاً یک فرمت را انتخاب کنید.")
            return
        
        format_keys = ["pdf", "excel", "csv", "json", "txt", "zip"]
        selected_format = format_keys[checked_id]
        
        ext_map = {
            "pdf": "pdf", "excel": "xlsx", "csv": "csv",
            "json": "json", "txt": "txt", "zip": "zip"
        }
        ext = ext_map.get(selected_format, "pdf")
        
        student_name = self.student_name_label.text().replace(" ", "_")
        default_name = f"داده_هوش_مصنوعی_{student_name}_{utc_now().strftime('%Y%m%d_%H%M%S')}.{ext}"
        
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "ذخیره فایل",
            default_name,
            f"{selected_format.upper()} Files (*.{ext})"
        )
        
        if not file_path:
            return
        
        try:
            data = self._collect_data()
            
            if selected_format == "pdf":
                self._export_pdf(data, file_path)
            elif selected_format == "excel":
                self._export_excel(data, file_path)
            elif selected_format == "csv":
                self._export_csv(data, file_path)
            elif selected_format == "json":
                self._export_json(data, file_path)
            elif selected_format == "txt":
                self._export_txt(data, file_path)
            elif selected_format == "zip":
                self._export_zip(data, file_path)
            
            QMessageBox.information(
                self,
                "موفقیت",
                f"✅ فایل با موفقیت در {file_path} ذخیره شد.\n\n"
                f"📋 پرامپت پیشنهادی را کپی کنید و همراه فایل به هوش مصنوعی ارسال کنید."
            )
            
            self.accept()
            
        except Exception as e:
            QMessageBox.critical(self, "خطا", f"مشکل در خروجی:\n{e!s}")
    
    def _collect_data(self):
        """جمع‌آوری تمام داده‌های مورد نیاز"""
        profile = self.profile_dal.get_by_id(self.profile_id)
        if not profile:
            return {}
        
        student = self.student_dal.get_by_id(profile.student_id)
        observations = self.observation_dal.get_by_student_profile(profile.id)
        interventions = self.intervention_dal.get_by_student_profile(profile.id)
        
        all_followups = []
        for inter in interventions:
            followups = self.followup_dal.get_by_intervention(inter.id)
            all_followups.extend(followups)
        
        competencies = self.competency_dal.get_all()
        comp_dict = {c.id: c.title for c in competencies}
        
        return {
            'student': {
                'id': student.id if student else None,
                'full_name': student.full_name if student else "",
                'national_code': student.national_code if student else "",
                'birth_date': student.birth_date if student else "",
                'father_name': student.father_name if student else "",
                'guardian_name': student.guardian_name if student else "",
                'guardian_phone': student.guardian_phone if student else "",
                'address': student.address if student else "",
            },
            'profile': {
                'id': profile.id,
                'grade': profile.grade,
                'grade_display': profile.grade_display,
                'class_name': profile.class_name or "",
                'status': profile.status,
            },
            'observations': [
                {
                    'id': o.id,
                    'date': o.observation_date,
                    'location': o.location,
                    'competency': comp_dict.get(o.competency_id, "نامشخص"),
                    'behavior_type': o.behavior_type,
                    'severity': o.severity,
                    'antecedent': o.antecedent,
                    'behavior': o.behavior,
                    'consequence': o.consequence,
                    'description': o.description,
                }
                for o in observations
            ],
            'interventions': [
                {
                    'id': i.id,
                    'date': i.date,
                    'type': i.type,
                    'type_display': i.type_display,
                    'description': i.description,
                    'goal': i.goal,
                    'status': i.status,
                    'result': i.result,
                }
                for i in interventions
            ],
            'followups': [
                {
                    'id': f.id,
                    'date': f.date,
                    'status': f.status,
                    'status_display': f.status_display,
                    'result_type': f.result_type,
                    'result_type_display': f.result_type_display,
                    'result_description': f.result_description,
                    'next_action_date': f.next_action_date,
                    'description': f.description,
                }
                for f in all_followups
            ],
            'metadata': {
                'export_date': utc_now_iso(),
                'total_observations': len(observations),
                'total_interventions': len(interventions),
                'total_followups': len(all_followups),
            }
        }
    
    def _export_pdf(self, data, file_path):
        """خروجی PDF همراه با پرامپت"""
        pdf = PersianPDF(file_path)
        
        pdf.add_title("📊 گزارش داده‌های دانش‌آموز برای هوش مصنوعی")
        pdf.add_spacer(0.3)
        
        student = data.get('student', {})
        profile = data.get('profile', {})
        
        pdf.add_subtitle("👤 اطلاعات دانش‌آموز")
        pdf.add_text(f"نام: {student.get('full_name', 'نامشخص')}")
        pdf.add_text(f"پایه: {profile.get('grade_display', 'نامشخص')}")
        pdf.add_text(f"کلاس: {profile.get('class_name', 'نامشخص')}")
        pdf.add_text(f"کد ملی: {student.get('national_code', 'ندارد')}")
        pdf.add_spacer(0.3)
        
        pdf.add_subtitle("📊 خلاصه آماری")
        meta = data.get('metadata', {})
        pdf.add_text(f"تعداد مشاهدات: {meta.get('total_observations', 0)}")
        pdf.add_text(f"تعداد مداخلات: {meta.get('total_interventions', 0)}")
        pdf.add_text(f"تعداد پیگیری‌ها: {meta.get('total_followups', 0)}")
        pdf.add_spacer(0.3)
        
        pdf.add_subtitle("📝 مشاهدات")
        for i, obs in enumerate(data.get('observations', []), 1):
            pdf.add_text(f"{i}. تاریخ: {obs.get('date', '')} | {obs.get('location', '')}")
            pdf.add_text(f"   شایستگی: {obs.get('competency', 'نامشخص')}")
            pdf.add_text(f"   نوع: {obs.get('behavior_type', 'نامشخص')} | شدت: {obs.get('severity', 0)}")
            pdf.add_text(f"   زمینه: {obs.get('antecedent', '-')}")
            pdf.add_text(f"   رفتار: {obs.get('behavior', '-')}")
            pdf.add_text(f"   پیامد: {obs.get('consequence', '-')}")
            pdf.add_spacer(0.1)
        
        pdf.add_page_break()
        pdf.add_title("📋 پرامپت پیشنهادی برای هوش مصنوعی")
        pdf.add_spacer(0.3)
        
        for line in self.prompt_text.toPlainText().split('\n'):
            if line.strip():
                pdf.add_text(line)
        
        pdf.build(file_path)
    
    def _export_excel(self, data, file_path):
        """خروجی Excel"""
        try:
            from openpyxl import Workbook
            
            wb = Workbook()
            
            ws = wb.active
            ws.title = "اطلاعات"
            
            student = data.get('student', {})
            profile = data.get('profile', {})
            
            info_data = [
                ["نام دانش‌آموز", student.get('full_name', 'نامشخص')],
                ["پایه", profile.get('grade_display', 'نامشخص')],
                ["کلاس", profile.get('class_name', 'نامشخص')],
                ["کد ملی", student.get('national_code', 'ندارد')],
                ["تاریخ تولد", student.get('birth_date', '')],
                ["نام پدر", student.get('father_name', '')],
                ["نام ولی", student.get('guardian_name', '')],
                ["شماره تماس ولی", student.get('guardian_phone', '')],
            ]
            
            for row, (key, value) in enumerate(info_data, 1):
                ws.cell(row=row, column=1, value=key)
                ws.cell(row=row, column=2, value=value)
            
            ws2 = wb.create_sheet("مشاهدات")
            headers = ["ردیف", "تاریخ", "محیط", "شایستگی", "نوع", "شدت", "زمینه", "رفتار", "پیامد"]
            for col, h in enumerate(headers, 1):
                ws2.cell(row=1, column=col, value=h)
            
            for row, obs in enumerate(data.get('observations', []), 2):
                ws2.cell(row=row, column=1, value=row-1)
                ws2.cell(row=row, column=2, value=obs.get('date', ''))
                ws2.cell(row=row, column=3, value=obs.get('location', ''))
                ws2.cell(row=row, column=4, value=obs.get('competency', 'نامشخص'))
                ws2.cell(row=row, column=5, value=obs.get('behavior_type', 'نامشخص'))
                ws2.cell(row=row, column=6, value=obs.get('severity', 0))
                ws2.cell(row=row, column=7, value=obs.get('antecedent', ''))
                ws2.cell(row=row, column=8, value=obs.get('behavior', ''))
                ws2.cell(row=row, column=9, value=obs.get('consequence', ''))
            
            ws3 = wb.create_sheet("مداخلات")
            headers = ["ردیف", "تاریخ", "نوع", "توضیحات", "هدف", "وضعیت", "نتیجه"]
            for col, h in enumerate(headers, 1):
                ws3.cell(row=1, column=col, value=h)
            
            for row, inter in enumerate(data.get('interventions', []), 2):
                ws3.cell(row=row, column=1, value=row-1)
                ws3.cell(row=row, column=2, value=inter.get('date', ''))
                ws3.cell(row=row, column=3, value=inter.get('type_display', 'نامشخص'))
                ws3.cell(row=row, column=4, value=inter.get('description', ''))
                ws3.cell(row=row, column=5, value=inter.get('goal', ''))
                ws3.cell(row=row, column=6, value=inter.get('status', ''))
                ws3.cell(row=row, column=7, value=inter.get('result', ''))
            
            wb.save(file_path)
            
        except ImportError:
            self._export_csv(data, file_path.replace('.xlsx', '.csv'))
    
    def _export_csv(self, data, file_path):
        """خروجی CSV"""
        
        with open(file_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            
            writer.writerow(['نوع', 'تاریخ', 'محیط/نوع', 'شایستگی/هدف', 'نوع رفتار/وضعیت', 'شدت/نتیجه', 'توضیحات'])
            
            for obs in data.get('observations', []):
                writer.writerow([
                    'مشاهده',
                    obs.get('date', ''),
                    obs.get('location', ''),
                    obs.get('competency', ''),
                    obs.get('behavior_type', ''),
                    obs.get('severity', ''),
                    obs.get('behavior', '')
                ])
            
            for inter in data.get('interventions', []):
                writer.writerow([
                    'مداخله',
                    inter.get('date', ''),
                    inter.get('type_display', ''),
                    inter.get('goal', ''),
                    inter.get('status', ''),
                    inter.get('result', ''),
                    inter.get('description', '')
                ])
    
    def _export_json(self, data, file_path):
        """خروجی JSON"""
        with open(file_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    
    def _export_txt(self, data, file_path):
        """خروجی TXT ساده"""
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write("=" * 60 + "\n")
            f.write("📊 داده‌های دانش‌آموز برای هوش مصنوعی\n")
            f.write("=" * 60 + "\n\n")
            
            student = data.get('student', {})
            f.write(f"نام: {student.get('full_name', 'نامشخص')}\n")
            f.write(f"پایه: {data.get('profile', {}).get('grade_display', 'نامشخص')}\n")
            f.write(f"کلاس: {data.get('profile', {}).get('class_name', 'نامشخص')}\n\n")
            
            f.write("-" * 40 + "\n")
            f.write("📝 مشاهدات:\n")
            f.write("-" * 40 + "\n")
            for obs in data.get('observations', []):
                f.write(f"تاریخ: {obs.get('date', '')} | {obs.get('location', '')}\n")
                f.write(f"شایستگی: {obs.get('competency', 'نامشخص')} | نوع: {obs.get('behavior_type', 'نامشخص')} | شدت: {obs.get('severity', 0)}\n")
                f.write(f"رفتار: {obs.get('behavior', '-')}\n\n")
            
            f.write("\n" + "-" * 40 + "\n")
            f.write("🛠️ مداخلات:\n")
            f.write("-" * 40 + "\n")
            for inter in data.get('interventions', []):
                f.write(f"تاریخ: {inter.get('date', '')} | نوع: {inter.get('type_display', 'نامشخص')}\n")
                f.write(f"توضیحات: {inter.get('description', '')}\n")
                f.write(f"نتیجه: {inter.get('result', 'هنوز مشخص نشده')}\n\n")
    
    def _export_zip(self, data, file_path):
        """خروجی ZIP شامل همه فرمت‌ها"""
        import shutil
        import tempfile
        
        temp_dir = tempfile.mkdtemp()
        base_name = os.path.splitext(os.path.basename(file_path))[0]
        
        try:
            formats = ['pdf', 'excel', 'csv', 'json', 'txt']
            saved_files = []
            
            for fmt in formats:
                temp_file = os.path.join(temp_dir, f"{base_name}.{fmt}")
                if fmt == 'pdf':
                    self._export_pdf(data, temp_file)
                elif fmt == 'excel':
                    self._export_excel(data, temp_file)
                elif fmt == 'csv':
                    self._export_csv(data, temp_file)
                elif fmt == 'json':
                    self._export_json(data, temp_file)
                elif fmt == 'txt':
                    self._export_txt(data, temp_file)
                saved_files.append(temp_file)
            
            with zipfile.ZipFile(file_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
                for f in saved_files:
                    if os.path.exists(f):
                        zipf.write(f, os.path.basename(f))
            
            shutil.rmtree(temp_dir)
            
        except Exception as e:
            shutil.rmtree(temp_dir)
            raise e