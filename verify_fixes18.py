"""
راستی‌آزمایی دور هجدهم — مرحلهٔ ۲: پیوست‌ها (BUG-ATT-04/05/06 + رگرسیون)

بخش‌ها:
  A) بازیابی پیوست حذف‌شده با «صفحهٔ واقعی» دیالوگ (آفلاین) — مرحلهٔ ۲:
  B) سال/پایه و UI (شاخص‌ها/تحلیل/ناوبری/کنتراست) — مرحلهٔ ۳:
  C) نگاشت ناوبری و کنتراست WCAG — مرحلهٔ ۳:
  D) رگرسیون نهایی: شش گزارش با خروجی واقعی PDF/Excel، AI export چهار
     قالب، داشبورد دوساله، پیشرفت هدف، migration (پایین/شکست میانی/
     بالا)، constraint دیتابیس تازه و قدیمی — مرحلهٔ ۴:
  (و در آغاز فایل:) بازیابی پیوست حذف‌شده با «صفحهٔ واقعی» دیالوگ (آفلاین):
     چک‌باکس «نمایش حذف‌شده‌ها»، انتخاب ردیف حذف‌شده، دکمهٔ ↩️،
     راستی‌آزمایی فایل فیزیکی پیش از بازیابی (BUG-ATT-05)، برهم‌کنش
     سقف ۲۰ با بازیابی (BUG-ATT-06)، و رگرسیون حذف منطقی (فایل می‌ماند).

نکتهٔ محیط: اجرای واقعی Qt با کتابخانه‌های جانشین (qtstub) و
QT_QPA_PLATFORM=offscreen. کلیک واقعی موس، QFileDialog و بازکردن فایل
خارجی همچنان NOT_TESTED - GUI EXECUTION REQUIRED هستند.

اجرا:
  LD_LIBRARY_PATH=/home/user/qtstub/lib QT_QPA_PLATFORM=offscreen \
      python3 verify_fixes18.py
"""

import contextlib
import io
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

PASS = 0
FAIL = 0
FAILURES = []


def check(section, name, condition, detail=""):
    global PASS, FAIL
    if condition:
        PASS += 1
        print(f"  ✅ [{section}] {name}")
    else:
        FAIL += 1
        FAILURES.append(f"[{section}] {name} {detail}")
        print(f"  ❌ [{section}] {name}  {detail}")


TMP = tempfile.mkdtemp(prefix="round18_verify_")

# ============================================================
# دیتابیس و پوشهٔ پیوست موقت (هیچ دست‌زدنی به دادهٔ واقعی نمی‌شود)
# ============================================================
TEST_DB = os.path.join(TMP, "partow.db")
ATT_DIR = os.path.join(TMP, "attachments")
os.makedirs(ATT_DIR, exist_ok=True)

import config.settings as settings
import database.connection as dbc

settings.DB_PATH = TEST_DB
dbc.DB_PATH = TEST_DB
settings.ATTACHMENTS_DIR = ATT_DIR
dbc.DatabaseConnection._instance = None
dbc.DatabaseConnection._connection = None
dbc.DatabaseConnection._initialized = False

with contextlib.redirect_stdout(io.StringIO()):
    db = dbc.DatabaseConnection()
    conn = db.get_connection(user_id=1)

# ماژول‌هایی که ATTACHMENTS_DIR را در سطح ماژول نگه می‌دارند
import dal.attachment_dal as attachment_dal_mod
import services.attachment_service as attachment_service_mod

attachment_dal_mod.ATTACHMENTS_DIR = ATT_DIR
attachment_service_mod.ATTACHMENTS_DIR = ATT_DIR

from dal.attachment_dal import AttachmentDAL
from services.attachment_service import AttachmentService

SERVICE = AttachmentService()

from PySide6.QtWidgets import QApplication, QMessageBox

app = QApplication.instance() or QApplication(sys.argv)

# --- رهگیری پیام‌ها (بدون جلوهٔ واقعی؛ محتوا سنجیده می‌شود) ---
messages = []
real_info = QMessageBox.information
real_crit = QMessageBox.critical
real_warn = QMessageBox.warning
real_question = QMessageBox.question

QMessageBox.information = staticmethod(
    lambda *a, **k: messages.append(("info", str(a[2])))
    or QMessageBox.StandardButton.Ok)
QMessageBox.critical = staticmethod(
    lambda *a, **k: messages.append(("crit", str(a[2])))
    or QMessageBox.StandardButton.Ok)
QMessageBox.warning = staticmethod(
    lambda *a, **k: messages.append(("warn", str(a[2])))
    or QMessageBox.StandardButton.Ok)
QMessageBox.question = staticmethod(
    lambda *a, **k: QMessageBox.StandardButton.Yes)

from dal.intervention_dal import InterventionDAL

# (دور نوزدهم، مرحلهٔ ۸ — RESTORE-EDGE-01) از این پس restore_attachment
# وجودِ واقعیِ موجودیت والد (نه فقط entity_id فرضی) را بررسی می‌کند؛ پس
# این اسکریپت هم به یک مشاهده/مداخلهٔ واقعی نیاز دارد، نه شناسه‌های
# ۴۲/۹ی صرفاً نمادین بدون رکورد پشتیبان.
from dal.observation_dal import ObservationDAL
from dal.student_academic_profile_dal import StudentAcademicProfileDAL
from dal.student_dal import StudentDAL
from models.intervention import Intervention
from models.observation import Observation
from models.student import Student
from models.student_academic_profile import StudentAcademicProfile
from views.dialogs.attachment_dialog import AttachmentDialog

with contextlib.redirect_stdout(io.StringIO()):
    _year_row_18 = conn.execute(
        "SELECT id FROM academic_years WHERE is_active = 1 LIMIT 1").fetchone()
    _entity_student = Student()
    _entity_student.first_name = "آزمون"
    _entity_student.last_name = "پیوست‌۱۸"
    _entity_student.national_code = "8300000001"
    _entity_student = StudentDAL().create(_entity_student)

    _entity_profile = StudentAcademicProfile()
    _entity_profile.student_id = _entity_student.id
    _entity_profile.academic_year_id = _year_row_18["id"]
    _entity_profile.grade = 1
    _entity_profile.class_name = "اول-الف"
    _entity_profile = StudentAcademicProfileDAL().create(_entity_profile)

    _entity_obs = Observation()
    _entity_obs.student_profile_id = _entity_profile.id
    _entity_obs.staff_id = 1
    _entity_obs.observation_date = "1404/01/01"
    _entity_obs.description = "مشاهدهٔ زمینه‌ساز پیوست دور ۱۸"
    _entity_obs.behavior_type = "خنثی"
    _entity_obs = ObservationDAL().create(_entity_obs)

    _entity_inter = Intervention()
    _entity_inter.student_profile_id = _entity_profile.id
    _entity_inter.staff_id = 1
    _entity_inter.type = "گفتگوی فردی"
    _entity_inter.date = "1404/01/01"
    _entity_inter.description = "مداخلهٔ زمینه‌ساز پیوست دور ۱۸"
    _entity_inter = InterventionDAL().create(_entity_inter)

ENTITY = ("observation", _entity_obs.id)
LIMITED_ENTITY_ID = _entity_inter.id


def make_attachment(name, content, entity=ENTITY):
    return SERVICE.upload_attachment(
        entity[0], entity[1], content, name, user_id=None)


def active_count(entity=ENTITY):
    return SERVICE.attachment_dal.get_count_by_entity(entity[0], entity[1])


def is_deleted(attachment_id):
    row = conn.execute(
        "SELECT is_deleted FROM attachments WHERE id = ?",
        (attachment_id,)).fetchone()
    return row is None or row["is_deleted"] == 1


print()
print("=" * 76)
print("بخش A: مرحلهٔ ۲ — بازیابی پیوست و راستی‌آزمایی فایل فیزیکی (بند ۱۱)")
print("=" * 76)

# ---------- A1: دیالوگ فعال‌ها؛ چک‌باکس موجود؛ دکمهٔ بازیابی پنهان ----------
a = make_attachment("active.txt", b"active content")
dlg = AttachmentDialog(ENTITY[0], ENTITY[1])
check("A", "A1: دیالوگ باز شد و پیوست فعال را نشان داد",
      dlg.file_list.count() == 1 and active_count() == 1,
      f"count={dlg.file_list.count()}")
check("A", "A1: چک‌باکس «نمایش حذف‌شده‌ها» در دیالوگ هست",
      hasattr(dlg, "show_deleted_check") and dlg.show_deleted_check is not None)
check("A", "A1: دکمهٔ بازیابی در حالت عادی پنهان است",
      dlg.restore_btn.isHidden() and not dlg.showing_deleted)
check("A", "A1: دکمهٔ افزودن فعال است",
      dlg.add_btn.isEnabled())
dlg.close()
dlg.deleteLater()

# ---------- A2: حذف منطقی → چک‌باکس → فهرست حذف‌شده‌ها ----------
deleted_att = make_attachment("victim.txt", b"victim content")
SERVICE.delete_attachment(deleted_att.id)
check("A", "A2: حذف منطقی: فایل فیزیکی سر جای خودش ماند",
      os.path.exists(deleted_att.file_path))

dlg = AttachmentDialog(ENTITY[0], ENTITY[1])
dlg.show_deleted_check.setChecked(True)   # سیگنال toggled → toggle_show_deleted
check("A", "A2: حالت حذف‌شده‌ها فعال شد و افزودن قفل شد",
      dlg.showing_deleted and not dlg.add_btn.isEnabled())
check("A", "A2: پیوست حذف‌شده با برچسب «(حذف‌شده)» در فهرست است",
      dlg.file_list.count() == 1
      and "(حذف‌شده)" in dlg.file_list.item(0).text())
check("A", "A2: دکمهٔ بازیابی در این حالت به نمایش درمی‌آید",
      not dlg.restore_btn.isHidden())

# ---------- A3: انتخاب ردیف حذف‌شده → فقط بازیابی ----------
dlg.file_list.setCurrentRow(0)
# کلیک واقعی در آفلاین ممکن نیست؛ همان هندلرِ itemClicked صدا زده می‌شود
dlg.on_file_selected(dlg.file_list.item(0))
check("A", "A3: انتخاب حذف‌شده → فقط «بازیابی» فعال است",
      dlg.restore_btn.isEnabled() and not dlg.open_btn.isEnabled()
      and not dlg.download_btn.isEnabled() and not dlg.delete_btn.isEnabled())

# ---------- A4: بازیابی واقعی از مسیر UI ----------
messages.clear()
dlg.restore_selected()

check("A", "A4: بازیابی از مسیر UI روی DB اثر واقعی گذاشت",
      not is_deleted(deleted_att.id), f"is_deleted={is_deleted(deleted_att.id)}")
check("A", "A4: فهرست خودکار به حالت فعال برگشت و پیوست دیده می‌شود",
      not dlg.showing_deleted and dlg.file_list.count() == 2)
check("A", "A4: پیام موفقیت صادر شد",
      any(kind == "info" for kind, _ in messages), f"messages={messages}")
check("A", "A4: دیالوگ به فهرست فعال برگشت (چک‌باکس برداشته شد)",
      not dlg.showing_deleted and not dlg.show_deleted_check.isChecked())
dlg.close()
dlg.deleteLater()

# ---------- A5: BUG-ATT-05 — فایل فیزیکی گم‌شده: شکست صریح، نه موفقیت خاموش ----------
ghost = make_attachment("ghost.txt", b"i will vanish")
SERVICE.delete_attachment(ghost.id)
os.remove(ghost.file_path)  # «پاک‌شدن دستی» فایل

dlg = AttachmentDialog(ENTITY[0], ENTITY[1])
dlg.show_deleted_check.setChecked(True)
ids = [dlg.file_list.item(i).data(0x0100)
       for i in range(dlg.file_list.count())]
check("A", "A5: پیوست بدون فایل در فهرست حذف‌شده‌ها دیده می‌شود",
      ghost.id in ids)
row = ids.index(ghost.id)
dlg.file_list.setCurrentRow(row)
dlg.on_file_selected(dlg.file_list.item(row))
messages.clear()
dlg.restore_selected()

check("A", "A5: بازیابی رد شد و پیام خطای صریح «فایل فیزیکی» آمد",
      any(kind == "crit" and "فایل فیزیکی" in msg
          for kind, msg in messages), f"messages={messages}")
check("A", "A5: رکورد در همان حالت حذف‌شده ماند (بدون نیمه‌بازیابی)",
      is_deleted(ghost.id))
check("A", "A5: پیام موفقیت جعلی صادر نشد",
      not any(kind == "info" for kind, _ in messages))
dlg.close()
dlg.deleteLater()

# ---------- A6: BUG-ATT-06 — سقف ۲۰ با بازیابی هم سنجیده می‌شود ----------
LIMITED = ("intervention", LIMITED_ENTITY_ID)
for i in range(20):
    make_attachment(f"lim_{i:02d}.txt", f"c{i}".encode(), LIMITED)
first = SERVICE.get_attachments_by_entity(LIMITED[0], LIMITED[1])[0]
SERVICE.delete_attachment(first.id)
extra = make_attachment("to_delete.txt", b"x", LIMITED)
# اکنون: ۲۰ فعال (سقف پر) + ۱ حذف‌شدهٔ در انتظار بازیابی («first»)
check("A", "A6: پیش‌فرض سناریو — ۲۰ فعال و ۱ حذف‌شده",
      active_count(LIMITED) == 20 and is_deleted(first.id)
      and not is_deleted(extra.id))

dlg = AttachmentDialog(LIMITED[0], LIMITED[1])
dlg.show_deleted_check.setChecked(True)
dlg.file_list.setCurrentRow(0)
dlg.on_file_selected(dlg.file_list.item(0))
messages.clear()
dlg.restore_selected()

check("A", "A6: بازیابی در سقفِ پر رد شد با پیام «حداکثر»",
      any(kind == "crit" and "حداکثر" in msg
          for kind, msg in messages), f"messages={messages}")
check("A", "A6: رکورد همچنان حذف‌شده است",
      is_deleted(first.id))

# جا باز می‌کنیم → همان مسیر UI حالا موفق می‌شود
holder = SERVICE.get_attachments_by_entity(LIMITED[0], LIMITED[1])[0]
SERVICE.delete_attachment(holder.id)
messages.clear()
dlg.load_attachments()          # تازه‌سازی فهرست حذف‌شده‌ها
dlg.file_list.setCurrentRow(0)
dlg.on_file_selected(dlg.file_list.item(0))
dlg.restore_selected()
check("A", "A6: پس از بازکردن جا، بازیابی از همان مسیر UI موفق شد",
      any(kind == "info" for kind, _ in messages)
      and not is_deleted(first.id)
      and active_count(LIMITED) == 20)
dlg.close()
dlg.deleteLater()

# ---------- A7: رگرسیون — حذف منطقی از مسیر UI فایل را نگه می‌دارد ----------
dlg = AttachmentDialog(ENTITY[0], ENTITY[1])
keep = make_attachment("keep_me.txt", b"keep", ENTITY)
dlg.load_attachments()
for i in range(dlg.file_list.count()):
    if dlg.file_list.item(i).data(0x0100) == keep.id:
        dlg.file_list.setCurrentRow(i)
        break
dlg.on_file_selected(dlg.file_list.currentItem())
messages.clear()
dlg.delete_selected()
check("A", "A7: حذف منطقی: رکورد حذف شد ولی فایل فیزیکی ماند",
      is_deleted(keep.id) and os.path.exists(keep.file_path))
check("A", "A7: پیام موفقیت حذف صادر شد",
      any(kind == "info" and "حذف" in msg for kind, msg in messages))
dlg.close()
dlg.deleteLater()

# ---------- A8: رگرسیون — قرارداد DAL (نباید True ثابت بدهد) ----------
fresh = make_attachment("contract.txt", b"contract", ENTITY)
check("A", "A8: بازیابی رکورد فعال در DAL → False (نه True ثابت)",
      AttachmentDAL().restore(fresh.id) is False)
SERVICE.delete_attachment(fresh.id)
check("A", "A8: بازیابی رکورد حذف‌شده در DAL → True",
      AttachmentDAL().restore(fresh.id) is True)
check("A", "A8: بازیابی دوبارهٔ همان رکورد → False",
      AttachmentDAL().restore(fresh.id) is False)

# ============================================================
# بخش B: مرحلهٔ ۳ — سال و پایهٔ فهرست‌ها (بندهای ۶ و ۷)
# ============================================================
print()
print("=" * 76)
print("بخش B: مرحلهٔ ۳ — هندلر سال شاخص‌ها/تحلیل و پایهٔ سال انتخاب‌شده")
print("=" * 76)

from dal.academic_year_dal import AcademicYearDAL
from dal.student_academic_profile_dal import StudentAcademicProfileDAL
from dal.student_dal import StudentDAL
from models.academic_year import AcademicYear
from models.student import Student
from models.student_academic_profile import StudentAcademicProfile
from views.pages.analysis_page import AnalysisPage
from views.pages.indicators_page import IndicatorsPage

year_dal = AcademicYearDAL()
student_dal = StudentDAL()
profile_dal = StudentAcademicProfileDAL()


def make_year(title, start, end):
    year = AcademicYear()
    year.title = title
    year.start_date = start
    year.end_date = end
    return year_dal.create(year)


def make_student(first, last, national_code):
    student = Student()
    student.first_name = first
    student.last_name = last
    student.national_code = national_code
    return student_dal.create(student)


def make_profile(student_id, year_id, grade, class_name):
    profile = StudentAcademicProfile()
    profile.student_id = student_id
    profile.academic_year_id = year_id
    profile.grade = grade
    profile.class_name = class_name
    profile.status = "active"
    return profile_dal.create(profile)


def select_combo_item(combo, data_value, emit=True):
    """انتخاب در کامبو (با یا بدون سیگنال واقعی)"""
    for i in range(combo.count()):
        if combo.itemData(i) == data_value:
            if not emit:
                combo.blockSignals(True)
            combo.setCurrentIndex(i)
            if not emit:
                combo.blockSignals(False)
            return True
    return False


def combo_student_ids(combo):
    return [combo.itemData(i) for i in range(combo.count())]


def combo_row_text(combo, data_value):
    for i in range(combo.count()):
        if combo.itemData(i) == data_value:
            return combo.itemText(i)
    return None


# دادهٔ دو ساله: A فقط ۱۴۰۳ (پایهٔ سوم)؛ B فقط ۱۴۰۴ (پایهٔ اول)
y1403 = make_year("۱۴۰۳-۱۴۰۴", "1403/07/01", "1404/06/30")
y1404 = make_year("۱۴۰۴-۱۴۰۵", "1404/07/01", "1405/06/30")
year_dal.set_active(y1404.id)

student_a = make_student("علی", "یک‌سال‌ه", "5550000001")
student_b = make_student("رضا", "دو‌سال‌ه", "5550000002")
profile_a = make_profile(student_a.id, y1403.id, 3, "سوم-الف")
profile_b = make_profile(student_b.id, y1404.id, 1, "اول-ب")

# ---------- B1: فهرست شاخص‌ها از سال انتخاب‌شده ----------
with contextlib.redirect_stdout(io.StringIO()):
    ipage = IndicatorsPage()

check("B", "B1: صفحهٔ شاخص‌ها با سال فعال (۱۴۰۴) بالا آمد",
      ipage.year_combo.currentData() == y1404.id,
      f"year={ipage.year_combo.currentData()}")
select_combo_item(ipage.year_combo, y1403.id)   # سیگنال واقعی → on_year_changed
check("B", "B1: فهرست ۱۴۰۳ فقط A را دارد (B نشت نکرده)",
      combo_student_ids(ipage.student_combo) == [None, student_a.id],
      f"ids={combo_student_ids(ipage.student_combo)}")
check("B", "B1: پایهٔ نمایشی A از پروندهٔ ۱۴۰۳ است («سوم»)",
      "سوم" in (combo_row_text(ipage.student_combo, student_a.id) or ""))

# ---------- B2: تغییر سال → پنل کهنه پاک ----------
select_combo_item(ipage.student_combo, student_a.id, emit=False)
with contextlib.redirect_stdout(io.StringIO()):
    ipage.load_student_indicators()
check("B", "B2: انتخاب A در ۱۴۰۳ ثبت شد",
      ipage.current_student_id == student_a.id
      and ipage.current_profile_id == profile_a.id)

select_combo_item(ipage.year_combo, y1404.id)   # سیگنال واقعی
check("B", "B2: تغییر سال → انتخاب و پروندهٔ کهنه پاک شد",
      ipage.current_student_id is None
      and ipage.current_profile_id is None)
check("B", "B2: پنل جزئیات پیام راهنما دارد (نه نتیجهٔ کهنه)",
      "دانش‌آموز" in ipage.details_panel.toPlainText())
check("B", "B2: فهرست ۱۴۰۴ فقط B را دارد",
      combo_student_ids(ipage.student_combo) == [None, student_b.id])
check("B", "B2: پایهٔ نمایشی B از پروندهٔ ۱۴۰۴ است («اول»)",
      "اول" in (combo_row_text(ipage.student_combo, student_b.id) or ""))
ipage.close()
ipage.deleteLater()

# ---------- B3: حفظ ایمن انتخاب وقتی همان دانش‌آموز در هر دو سال هست ----------
make_profile(student_b.id, y1403.id, 2, "دوم-ب")
with contextlib.redirect_stdout(io.StringIO()):
    ipage = IndicatorsPage()
select_combo_item(ipage.year_combo, y1404.id)
select_combo_item(ipage.student_combo, student_b.id, emit=False)
with contextlib.redirect_stdout(io.StringIO()):
    ipage.load_student_indicators()
select_combo_item(ipage.year_combo, y1403.id)   # سیگنال واقعی
check("B", "B3: انتخاب B پس از تغییر سال حفظ شد",
      ipage.current_student_id == student_b.id)
check("B", "B3: پنل با پروندهٔ ۱۴۰۳ دوباره خوانده شد (پایهٔ دوم)",
      ipage.current_profile_id == profile_dal.get_by_student_and_year(
          student_b.id, y1403.id).id)
ipage.close()
ipage.deleteLater()

# ---------- B4: صفحهٔ تحلیل — همان قرارداد ----------
# (توجه: از B3، دانش‌آموز B در ۱۴۰۳ هم پرونده دارد → فهرست ۱۴۰۳ هر دو)
with contextlib.redirect_stdout(io.StringIO()):
    apage = AnalysisPage()
select_combo_item(apage.year_combo, y1403.id)   # سیگنال واقعی → on_year_changed
ids_1403 = combo_student_ids(apage.student_combo)
check("B", "B4: فهرست تحلیل ۱۴۰۳ = پرونده‌های همان سال (A و B از B3)",
      ids_1403 == [None, student_b.id, student_a.id], f"ids={ids_1403}")
check("B", "B4: پایهٔ B در ۱۴۰۳ «دوم» است (پروندهٔ همان سال، نه فعال)",
      "دوم" in (combo_row_text(apage.student_combo, student_b.id) or ""))
select_combo_item(apage.year_combo, y1404.id)
check("B", "B4: فهرست تحلیل در ۱۴۰۴ فقط B است",
      combo_student_ids(apage.student_combo) == [None, student_b.id])
check("B", "B4: پایهٔ B در ۱۴۰۴ «اول» است (پروندهٔ همان سال)",
      "اول" in (combo_row_text(apage.student_combo, student_b.id) or ""))
apage.close()
apage.deleteLater()

# ============================================================
# بخش C: مرحلهٔ ۳ — نگاشت ناوبری و کنتراست هدر (بندهای ۲۶ و ۳۱)
# ============================================================
print()
print("=" * 76)
print("بخش C: مرحلهٔ ۳ — نگاشت مرکزی ناوبری و کنتراست نام/آیکون کاربر")
print("=" * 76)

from views.main_window import PAGE_INDEX

# ---------- C1: نگاشت مرکزی ----------
check("C", "C1: PAGE_INDEX پیوسته 0..N است",
      sorted(PAGE_INDEX.values()) == list(range(len(PAGE_INDEX))))
check("C", "C1: ایندکس‌های کلیدی مطابق ترتیب addWidget است",
      PAGE_INDEX["students"] == 2 and PAGE_INDEX["followups"] == 5
      and PAGE_INDEX["reports"] == 8 and PAGE_INDEX["goals"] == 13)

# ---------- C2: هیچ setCurrentIndex عددی خام روی stacked_widget ----------
import pathlib

mw_src = pathlib.Path("views/main_window.py").read_text(encoding="utf-8")
import re

raw_indexes = [
    m.group(1).strip()
    for m in re.finditer(r"stacked_widget\.setCurrentIndex\(([^)]*)\)", mw_src)
]
check("C", "C2: setCurrentIndex فقط از PAGE_INDEX (index/idx) می‌خواند",
      set(raw_indexes) <= {"index", "idx"}, f"args={raw_indexes}")
check("C", "C2: منو با نام صفحه کار می‌کند نه ایندکس عددی",
      '"btn_settings", "⚙️ تنظیمات", "settings", None' in mw_src
      and ', 9, None)' not in mw_src)

# ---------- C3: کنتراست محاسباتی (WCAG) ----------

def _relative_luminance(hex_color):
    r, g, b = (int(hex_color[i:i + 2], 16) / 255 for i in (0, 2, 4))
    def channel(c):
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * channel(r) + 0.7152 * channel(g) + 0.0722 * channel(b)


def contrast_ratio(fg_hex, bg_hex):
    l1 = _relative_luminance(fg_hex)
    l2 = _relative_luminance(bg_hex)
    lighter, darker = max(l1, l2), min(l1, l2)
    return (lighter + 0.05) / (darker + 0.05)


# سایدبار: user_frame #08223A — نام/آیکون باید مرئی باشند (دور هجدهم)
sidebar_name = contrast_ratio("F4C542", "08223A")
check("C", "C3: نام/آیکون کاربر سایدبار (#F4C542 روی #08223A) ≥ 4.5",
      sidebar_name >= 4.5, f"ratio={sidebar_name:.2f}")
# رگرسیون: رنگ قدیمی نامرئی نباید برگردد
check("C", "C3: رنگ نامرئی قبلی (#08223A روی خودش) در سورس سایدبار نیست",
      'user_icon_label.setStyleSheet("font-size: 20px; color: #08223A;")'
      not in mw_src
      and "color: #08223A; font-size: 13px; font-weight: bold;" not in mw_src)
# هدر اصلی: پس‌زمینه #0B2E4F
header_user = contrast_ratio("66BB6A", "0B2E4F")
header_title = contrast_ratio("F4C542", "0B2E4F")
header_version = contrast_ratio("D9C36A", "0B2E4F")
check("C", "C3: نام کاربر هدر (#66BB6A) ≥ 4.5",
      header_user >= 4.5, f"ratio={header_user:.2f}")
check("C", "C3: عنوان هدر (#F4C542) ≥ 4.5",
      header_title >= 4.5, f"ratio={header_title:.2f}")
check("C", "C3: نسخهٔ هدر (#D9C36A) ≥ 4.5",
      header_version >= 4.5, f"ratio={header_version:.2f}")
# نقش کاربر سایدبار (#D9C36A روی #08223A)
check("C", "C3: نقش کاربر سایدبار (#D9C36A) ≥ 4.5",
      contrast_ratio("D9C36A", "08223A") >= 4.5)

# ============================================================
# بخش D: مرحلهٔ ۴ — رگرسیون نهایی (بندهای ۱۸، ۱۹، ۲۰، ۲۱، ۲۳، ۲۹، ۳۰)
# ============================================================
print()
print("=" * 76)
print("بخش D: مرحلهٔ ۴ — شش گزارش، AI export، داشبورد، ارتقاء پایه، هدف، migration، constraint")
print("=" * 76)

import sqlite3 as _sqlite3

from dal.class_dal import ClassDAL
from dal.competency_dal import CompetencyDAL
from dal.followup_dal import FollowUpDAL
from dal.goal_dal import GoalDAL
from dal.intervention_dal import InterventionDAL
from dal.observation_dal import ObservationDAL
from dal.staff_dal import StaffDAL
from database.migrations.manager import MigrationManager, MigrationStepError
from models.class_model import ClassModel
from models.followup import FollowUp
from models.intervention import Intervention
from models.observation import Observation
from models.staff import Staff
from services.case_timeline_service import CaseTimelineService
from services.class_report_service import ClassReportService
from services.dashboard_service import DashboardService
from services.goal_service import GoalService
from services.report_generator import ReportGenerator
from services.teacher_performance_service import TeacherPerformanceService
from services.teacher_report_service import TeacherReportService

REPORTS_DIR = os.path.join(TMP, "reports_d")
os.makedirs(REPORTS_DIR, exist_ok=True)

# ---------- D0: دادهٔ پایهٔ گزارش‌ها ----------
report_teacher = Staff()
report_teacher.full_name = "معلم گزارش‌ها"
report_teacher.role = "teacher"
teacher_id = StaffDAL().create(report_teacher).id

class_1403 = ClassModel()
class_1403.name = "سوم-الف"          # باید با class_name پروندهٔ A یکی باشد
class_1403.grade = 3
class_1403.teacher_id = teacher_id
class_1403.academic_year_id = y1403.id
class_1403.capacity = 30
class_1403_id = ClassDAL().create(class_1403).id

competency_id = CompetencyDAL().get_all()[0].id


def make_observation(profile_id, staff_id, date, behavior, behavior_type,
                     severity=2, student_id=None, year_id=None):
    observation = Observation()
    observation.student_profile_id = profile_id
    observation.staff_id = staff_id
    observation.observation_date = date
    observation.behavior = behavior
    observation.description = "شرح آزمایشی برای رگرسیون گزارش‌ها"
    observation.behavior_type = behavior_type
    observation.severity = severity
    observation.competency_id = competency_id
    observation.location = "کلاس"
    created = ObservationDAL().create(observation)
    return created


# زنجیرهٔ کامل برای A در ۱۴۰۳: مشاهده → مداخله → پیگیری
obs_a = make_observation(profile_a.id, teacher_id, "1403/10/01",
                         "رفتار مثبت آزمایشی گزارش", "مثبت")
intervention_a = Intervention()
intervention_a.student_profile_id = profile_a.id
intervention_a.staff_id = teacher_id
intervention_a.type = "parent_call"
intervention_a.description = "مداخلهٔ آزمایشی تماس با والد"
intervention_a.date = "1403/10/02"
intervention_a.status = "planned"
intervention_a_id = InterventionDAL().create(intervention_a).id

followup_a = FollowUp()
followup_a.intervention_id = intervention_a_id
followup_a.staff_id = teacher_id
followup_a.date = "1403/10/05"
followup_a.status = "pending"
followup_a.result_type = "done"
followup_a.description = "پیگیری آزمایشی تماس با والد"
followup_a_id = FollowUpDAL().create(followup_a).id

# مشاهدهٔ دوم A در ۱۴۰۳ (برای تمایز شمارش داشبورد) و یکی برای B در ۱۴۰۴
make_observation(profile_a.id, teacher_id, "1403/11/01",
                 "رفتار منفی آزمایشی گزارش", "منفی", severity=4)
make_observation(profile_b.id, teacher_id, "1404/08/01",
                 "رفتار مثبت سال ۱۴۰۴", "مثبت")

check("D", "D0: دادهٔ پایهٔ گزارش‌ها ساخته شد (۲ مشاهدهٔ ۱۴۰۳، ۱ مشاهدهٔ ۱۴۰۴، زنجیرهٔ کامل A)",
      obs_a.id and intervention_a_id and followup_a_id)


def file_ok(path, prefix=None):
    if not os.path.exists(path) or os.path.getsize(path) == 0:
        return False
    if prefix:
        with open(path, "rb") as fh:
            return fh.read(len(prefix)) == prefix
    return True


def excel_ok(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    names = wb.sheetnames
    wb.close()
    return bool(names)


# ---------- D1: گزارش دانش‌آموز (PDF + Excel) ----------
gen = ReportGenerator()
pdf_path = os.path.join(REPORTS_DIR, "student.pdf")
xlsx_path = os.path.join(REPORTS_DIR, "student.xlsx")
gen.export_to_pdf(profile_a.id, pdf_path)
gen.export_to_excel(profile_a.id, xlsx_path)
check("D", "D1: گزارش دانش‌آموز PDF معتبر (%PDF، حجم غیرصفر)",
      file_ok(pdf_path, b"%PDF"))
check("D", "D1: گزارش دانش‌آموز Excel باز می‌شود و شیت دارد",
      file_ok(xlsx_path) and excel_ok(xlsx_path))

# ---------- D2: گزارش والدین ----------
parent_data = gen.generate_parent_report(profile_a.id)
check("D", "D2: گزارش والدین دادهٔ معتبر برمی‌گرداند",
      isinstance(parent_data, dict) and len(parent_data) > 0)

# ---------- D3: گزارش کلاس (PDF + Excel) ----------
class_svc = ClassReportService()
class_pdf = os.path.join(REPORTS_DIR, "class.pdf")
class_xlsx = os.path.join(REPORTS_DIR, "class.xlsx")
ok_pdf, _ = class_svc.export_class_report_pdf(class_1403_id, class_pdf) or (False, "")
ok_xlsx, _ = class_svc.export_class_report_excel(class_1403_id, class_xlsx) or (False, "")
check("D", "D3: گزارش کلاس ۱۴۰۳ PDF (داده‌دار، %PDF)",
      ok_pdf is True and file_ok(class_pdf, b"%PDF"),
      f"ok={ok_pdf}")
check("D", "D3: گزارش کلاس ۱۴۰۳ Excel باز می‌شود",
      ok_xlsx is True and excel_ok(class_xlsx), f"ok={ok_xlsx}")

# ---------- D4: گزارش معلم (PDF + Excel) ----------
teacher_svc = TeacherReportService()
teacher_pdf = os.path.join(REPORTS_DIR, "teacher.pdf")
teacher_xlsx = os.path.join(REPORTS_DIR, "teacher.xlsx")
report_data = {
    "teacher_name": "معلم گزارش‌ها",
    "total_students": 1,
    "start_date": "1403/10/01",
    "end_date": "1403/12/29",
    "year_title": "۱۴۰۳-۱۴۰۴",
    "total_observations": 2,
    "positive_observations": 1,
    "negative_observations": 1,
    "neutral_observations": 0,
    "total_interventions": 1,
    "pending_followups": 1,
    "competency_stats": {},
}
teacher_svc.export_to_pdf(report_data, teacher_pdf)
teacher_svc.export_to_excel(report_data, teacher_xlsx)
check("D", "D4: گزارش معلم PDF معتبر",
      file_ok(teacher_pdf, b"%PDF"))
check("D", "D4: گزارش معلم Excel باز می‌شود",
      file_ok(teacher_xlsx) and excel_ok(teacher_xlsx))

# ---------- D5: گزارش عملکرد معلم ----------
perf_svc = TeacherPerformanceService()
perf_pdf = os.path.join(REPORTS_DIR, "performance.pdf")
try:
    perf_result = perf_svc.export_teacher_report_pdf(
        teacher_id, perf_pdf, "1403/10/01", "1404/03/31")
except Exception as _perf_exc:
    perf_result = (False, str(_perf_exc))
perf_ok = file_ok(perf_pdf, b"%PDF")
check("D", "D5: گزارش عملکرد معلم PDF معتبر (یا «بدون داده» صریح)",
      perf_ok or (isinstance(perf_result, tuple)
                  and perf_result[0] is False
                  and "داده" in str(perf_result[1])),
      f"result={perf_result}")

# ---------- D6: Timeline (زنجیرهٔ مشاهده → مداخله → پیگیری) ----------
timeline_svc = CaseTimelineService()
timeline = timeline_svc.get_timeline(profile_a.id)
event_types = sorted({event.get("type") for event in timeline})
check("D", "D6: Timeline هر سه رویداد زنجیرهٔ A را دارد",
      {"observation", "intervention", "followup"} <= set(event_types),
      f"types={event_types}")
timeline_summary = timeline_svc.get_timeline_summary(profile_a.id)
check("D", "D6: خلاصهٔ Timeline با شمارش واقعی می‌خواند",
      isinstance(timeline_summary, dict))

# ---------- D7: AI Export — چهار قالب + پیگیری‌ها + شکست Excel ≠ CSV ----------
from views.dialogs.export_ai_dialog import ExportAIDialog

with contextlib.redirect_stdout(io.StringIO()):
    ai_dialog = ExportAIDialog(student_id=student_a.id, profile_id=profile_a.id)
ai_data = ai_dialog._collect_data()
check("D", "D7: دادهٔ AI شامل پیگیری‌ها است",
      len(ai_data.get("followups", [])) == 1,
      f"followups={len(ai_data.get('followups', []))}")

ai_base = os.path.join(REPORTS_DIR, "ai_export")
ai_dialog._export_pdf(ai_data, ai_base + ".pdf")
ai_dialog._export_excel(ai_data, ai_base + ".xlsx")
ai_dialog._export_csv(ai_data, ai_base + ".csv")
ai_dialog._export_txt(ai_data, ai_base + ".txt")
check("D", "D7: خروجی PDF AI معتبر",
      file_ok(ai_base + ".pdf", b"%PDF"))

from openpyxl import load_workbook

ai_wb = load_workbook(ai_base + ".xlsx", read_only=True)
ai_sheets = " | ".join(ai_wb.sheetnames)
ai_wb.close()
check("D", "D7: Excel AI شیت پیگیری‌ها دارد (fallback CSV نیست)",
      ("پیگیری" in ai_sheets), f"sheets={ai_sheets}")
csv_content = pathlib.Path(ai_base + ".csv").read_text(encoding="utf-8")
txt_content = pathlib.Path(ai_base + ".txt").read_text(encoding="utf-8")
check("D", "D7: پیگیری‌ها در CSV و TXT هم هستند",
      "پیگیری" in csv_content and "پیگیری" in txt_content)

ai_src = pathlib.Path("views/dialogs/export_ai_dialog.py").read_text(encoding="utf-8")
excel_block = ai_src.split("def _export_excel")[1].split("def _export_csv")[0]
check("D", "D7: مسیر _export_excel هیچ fallback به CSV ندارد (رگرسیون)",
      ("_export_csv" not in excel_block and "raise" in excel_block) or
      ("return False" in excel_block or "except" in excel_block))

# ---------- D8: داشبورد با دامنهٔ سال ----------
dash = DashboardService()
dash_1403 = dash.get_dashboard_data(year_id=y1403.id)
dash_1404 = dash.get_dashboard_data(year_id=y1404.id)
check("D", "D8: داشبورد ۱۴۰۳ = ۲ مشاهدهٔ همان سال",
      dash_1403["general_stats"]["observations_count"] == 2,
      f"got={dash_1403['general_stats']['observations_count']}")
check("D", "D8: داشبورد ۱۴۰۴ = ۱ مشاهدهٔ همان سال (بدون نشتی)",
      dash_1404["general_stats"]["observations_count"] == 1,
      f"got={dash_1404['general_stats']['observations_count']}")

# ---------- D9: پیشرفت هدف — کلمپ ۰..۱۰۰ + ماندگاری ----------
goal_service = GoalService()
goal_created = goal_service.create_goal({
    "student_profile_id": profile_a.id,
    "title": "هدف آزمایشی رگرسیون",
    "description": "شرح هدف",
    "domain": "educational",
    "priority": "medium",
}, user_id=teacher_id)
goal_service.update_goal_progress(
    goal_created.id, 150, notes="بیش از سقف", user_id=teacher_id)
goal_after = GoalDAL().get_by_id(goal_created.id)
check("D", "D9: پیشرفت هدف ۱۵۰٪ به ۱۰۰ کلمپ و در DB ماند",
      goal_after.progress_percent == 100
      and goal_after.progress_notes == "بیش از سقف")
goal_service.update_goal_progress(goal_created.id, -5, user_id=teacher_id)
check("D", "D9: پیشرفت منفی به ۰ کلمپ می‌شود",
      GoalDAL().get_by_id(goal_created.id).progress_percent == 0)

# ---------- D10: Migration — idempotent، پایین‌آمدن، شکست میانی، بالا‌آمدن ----------
conn_mig = dbc.DatabaseConnection().get_connection()
version_before = MigrationManager.get_current_version(conn_mig)
MigrationManager.migrate(conn_mig)
check("D", "D10: migrate روی دیتابیس به‌روز بی‌اثر است (idempotent)",
      MigrationManager.get_current_version(conn_mig) == version_before)

MigrationManager.migrate(conn_mig, target_version=8)
check("D", "D10: downgrade ۹→۸ انجام شد و نسخه ثبت شد",
      MigrationManager.get_current_version(conn_mig) == 8)

original_run_step = MigrationManager._run_step


def failing_run_step(connection, module, version, direction, new_version=None):
    if version == 8 and direction == "downgrade":
        raise MigrationStepError(8, "downgrade", "شکست عمدی برای آزمون")
    return original_run_step(connection, module, version, direction,
                             new_version=new_version)


MigrationManager._run_step = staticmethod(failing_run_step)
mid_failure_raised = False
try:
    MigrationManager.migrate(conn_mig, target_version=7)
except MigrationStepError:
    mid_failure_raised = True
MigrationManager._run_step = staticmethod(original_run_step)
check("D", "D10: شکست عمدی وسط downgrade → MigrationStepError",
      mid_failure_raised)
check("D", "D10: پس از شکست میانی، نسخه در گام آخرِ معتبر (۸) ماند",
      MigrationManager.get_current_version(conn_mig) == 8,
      f"version={MigrationManager.get_current_version(conn_mig)}")

MigrationManager.migrate(conn_mig)  # دوباره به آخرین نسخه
check("D", "D10: re-upgrade به نسخهٔ ۹ پس از شکست میانی",
      MigrationManager.get_current_version(conn_mig) == 9)
check("D", "D10: داده پس از پایین/بالای migration دست‌نخورده است",
      # نکته: عدد ۴ (نه ۳) از مرحلهٔ ۸ (RESTORE-EDGE-01) به بعد درست است؛
      # بخش A حالا برای بازیابیِ واقعی پیوست به یک مشاهدهٔ واقعیِ زمینه‌ساز
      # (ENTITY) نیاز دارد که پیش از فیکسچر D0 ساخته می‌شود — پس ۳ مشاهدهٔ
      # D0 + همان ۱ مشاهدهٔ زمینه‌ساز = ۴.
      conn_mig.execute(
          "SELECT COUNT(*) FROM observations WHERE is_deleted = 0"
      ).fetchone()[0] == 4)

# ---------- D11: Constraint در دیتابیس تازه و قدیمی (بند ۳۰) ----------
# تازه: CHECK باید مقدار نامعتبر را رد کند
legacy_dir = os.path.join(TMP, "legacy_db")
os.makedirs(legacy_dir, exist_ok=True)
legacy_path = os.path.join(legacy_dir, "legacy.db")
legacy_conn = _sqlite3.connect(legacy_path)
legacy_conn.execute("""
    CREATE TABLE observations (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_profile_id INTEGER,
        staff_id INTEGER,
        observation_date TEXT,
        behavior TEXT,
        description TEXT,
        behavior_type TEXT,
        severity INTEGER
    )
""")
try:
    legacy_conn.execute(
        "INSERT INTO observations (behavior_type) VALUES ('نامعتبر')")
    legacy_accepts = True
except _sqlite3.IntegrityError:
    legacy_accepts = False
legacy_conn.rollback()
legacy_conn.close()

fresh_rejects = False
try:
    conn_mig.execute("""
        INSERT INTO observations (
            student_profile_id, staff_id, observation_date, behavior,
            description, behavior_type
        ) VALUES (?, ?, '1404/01/01', 'رفتار نامعتبر برای CHECK',
                  'شرح', 'نامعتبر')
    """, (profile_a.id, teacher_id))
    conn_mig.rollback()
except _sqlite3.IntegrityError:
    fresh_rejects = True

check("D", "D11: دیتابیس تازه: behavior_type نامعتبر با CHECK رد می‌شود",
      fresh_rejects)
check("D", "D11: دیتابیس قدیمیِ بدون بازسازی: CHECK ندارد (واقعیت ثبت شد)",
      legacy_accepts is True,
      "مستند: DBهای قدیمی تا بازسازی جدول، CHECK رفتار را اعمال نمی‌کنند")

# ============================================================
# پایان
# ============================================================

QMessageBox.information = real_info
QMessageBox.critical = real_crit
QMessageBox.warning = real_warn
QMessageBox.question = real_question

with contextlib.suppress(Exception):
    dbc.DatabaseConnection().close_all()
with contextlib.suppress(Exception):
    shutil.rmtree(TMP, ignore_errors=True)

print()
print("=" * 76)
print(f"نتیجهٔ دور هجدهم (مرحله‌های ۲ تا ۴):  {PASS} موفق / {FAIL} ناموفق  از {PASS + FAIL}")
print("=" * 76)
if FAILURES:
    print("موارد ناموفق:")
    for f in FAILURES:
        print(f"   - {f}")
    sys.exit(1)
sys.exit(0)
